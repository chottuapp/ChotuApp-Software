"""Two isolated store processes. The parent alone writes shared saved-run queue state."""
import json,os,queue,re,subprocess,sys,threading,time
from pathlib import Path
from urllib.parse import urlsplit


def source_key(site):
    host=(urlsplit(site['url']).hostname or '').lower().removeprefix('www.')
    if host in ('shop.fardafoods.com','fardafoods.com'):host='fardafoods.com'
    return host


def output_key(site):return (re.sub(r'[^\w .-]','',site['name']).strip() or 'Shopify').casefold()


class StoreScheduler:
    def __init__(self,engine,store,state,output,pause,send,max_processes=2,retry_limit=2,wait_limit=120,skip_terminal=False):
        self.engine=Path(engine);self.store=store;self.state=state;self.output=Path(output);self.pause=pause;self.send=send
        self.max_processes=max(1,min(2,max_processes));self.retry_limit=retry_limit;self.wait_limit=wait_limit
        if len({output_key(site) for site in state['sites']})!=len(state['sites']):raise RuntimeError('Selected store names produce the same Excel filename; give each store a distinct name.')
        self.pending=[i for i in store.pending(state) if not skip_terminal or state['sites'][i].get('status') not in ('failed','blocked')];self.attempts={};self.active={};self.messages=queue.Queue();self.stop=threading.Event()
        self.control_lock=threading.Lock();self.paused=False;self.completed=0;self.deferred=0;self.failed=0
        self.ready=[];self.waiting={};self.before={};self.initial_deferred=[]
        for index in self.pending:
            site=state['sites'][index];retry=site.get('retry_after_epoch',0)
            if site.get('status')=='source_paused' and retry>time.time():
                if retry-time.time()<=self.wait_limit:self.waiting[index]=retry
                else:self.initial_deferred.append(index)
            else:self.ready.append(index)
    def raw(self,item):self.send(item)
    def mark(self,index,**values):
        def change(state):state['sites'][index].update(values);state['status']='incomplete'
        self.store.update(change)
    def control(self,job,command):
        with self.control_lock:
            try:job['process'].stdin.write(json.dumps({'command':command})+'\n');job['process'].stdin.flush()
            except (OSError,ValueError):pass
    def read(self,index,process):
        try:
            for line in process.stdout:
                try:message=json.loads(line)
                except ValueError:message={'type':'debug','text':line.rstrip()[:2000]}
                if isinstance(message,dict):self.messages.put((index,message))
        finally:self.messages.put((index,{'type':'eof'}))
    def start(self,index):
        current=self.store.load();site=current['sites'][index];folder=self.store.root/self.state['id']/str(index);folder.mkdir(parents=True,exist_ok=True)
        name=re.sub(r'[^\w .-]','',site['name']).strip() or 'Shopify';path=self.output/(name+'.xlsx')
        self.before[index]=(path.stat().st_size,path.stat().st_mtime_ns) if path.is_file() else None
        self.attempts[index]=self.attempts.get(index,0)+1
        env=dict(os.environ,PYTHONIOENCODING='utf-8',PYTHONUNBUFFERED='1',CHOTUAPP_DATA_DIR=str(self.output.parent))
        process=subprocess.Popen([sys.executable,'-u',str(Path(__file__).with_name('store_process_worker.py'))],stdin=subprocess.PIPE,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,encoding='utf-8',errors='replace',env=env,bufsize=1)
        job={'process':process,'site':site,'result':None,'eof':False,'reader':None,'floor':site.get('progress',0)}
        self.active[index]=job
        request={'engine':str(self.engine),'folder':str(folder),'site':site,'template':self.state['template'],'output_root':str(self.output),'resuming':bool(self.state.get('_resuming')) or self.attempts[index]>1}
        process.stdin.write(json.dumps(request)+'\n');process.stdin.flush()
        job['reader']=threading.Thread(target=self.read,args=(index,process),daemon=True);job['reader'].start()
        self.mark(index,status='running',detail='Fetching in separate process')
        self.raw(('CONTINUE SAVED RUN: ' if request['resuming'] else 'NEW FRESH RUN: ')+site['name'])
    def progress(self,index,text):
        _,pct,detail=text.split('|',2);pct=max(self.active[index]['floor'],min(100,int(pct)))
        self.mark(index,progress=pct,detail=detail)
        state=self.store.load();overall=sum(site.get('progress',0) for site in state['sites'])//max(1,len(state['sites']))
        names=', '.join(job['site']['name'] for job in self.active.values())
        self.raw('__CHOTU_PROGRESS__|'+str(overall)+'|Active: '+names)
    def event(self,index,message):
        job=self.active.get(index)
        if job is None:return
        kind=message.get('type');text=str(message.get('text',''));prefix='['+job['site']['name']+'] '
        if kind=='version':
            self.mark(index,software_revision=message.get('revision'),previous_software_revision=message.get('previous'));return
        if kind=='result':job['result']=message;return
        if kind=='eof':job['eof']=True;return
        if kind=='debug':self.raw(('__DEBUG__',prefix+text));return
        if kind!='log':return
        if text.startswith('__CHOTU_PROGRESS__|'):
            self.progress(index,text);return
        if 'REVIEW REQUIRED:' in text or 'SAVED REVIEW EXCEL:' in text:self.mark(index,reviewed=True)
        self.raw(prefix+text)
    def finish(self,index):
        job=self.active[index];process=job['process'];result=job['result'] or {'ok':False,'detail':'Store process exited before reporting its result'}
        site=job['site'];process.wait();process.stdin.close();process.stdout.close();job['reader'].join(timeout=1)
        del self.active[index]
        if result.get('paused'):
            retry=max(time.time()+1,float(result.get('retry_after_epoch',time.time()+60)))
            self.mark(index,status='source_paused',retry_after_epoch=retry,detail='Website cooldown — '+str(result.get('detail','')))
            if self.attempts[index]<=self.retry_limit and retry-time.time()<=self.wait_limit:self.waiting[index]=retry+1
            else:
                self.deferred+=1;self.raw(('__JOB_RESULT__',False,site['name']))
            self.raw('WEBSITE COOLDOWN: '+site['name']+' — saved requests retained; other stores continue.')
            return
        path=self.output/((re.sub(r'[^\w .-]','',site['name']).strip() or 'Shopify')+'.xlsx')
        if result.get('ok') and path.is_file():
            stat=path.stat();unchanged=(stat.st_size,stat.st_mtime_ns)==self.before[index]
            if unchanged:result=dict(result,ok=False,detail='Existing file was unchanged; no new Excel confirmed')
            else:
                from openpyxl import load_workbook
                try:
                    wb=load_workbook(path,read_only=True)
                    try:valid='Products' in wb.sheetnames and wb['Products'].max_row>1
                    finally:wb.close()
                    if not valid:raise ValueError('Export contains no products')
                except Exception as exc:result=dict(result,ok=False,detail='Export validation failed: '+str(exc))
        else:result=dict(result,ok=False,detail=result.get('detail') or 'Output workbook missing')
        if result.get('ok'):
            stat=path.stat();self.mark(index,status='done',progress=100,detail='File saved',reviewed=bool(result.get('reviewed')),output={'path':str(path.resolve()),'size':stat.st_size,'mtime_ns':stat.st_mtime_ns})
            self.completed+=1;self.raw('FILE SAVED NOW: '+str(path)+' — other store processes continue.')
        else:
            self.failed+=1;self.mark(index,status='blocked' if result.get('blocked') else 'failed',detail=str(result.get('detail')));self.raw('FAILED: '+site['name']+': '+str(result.get('detail')))
        self.raw(('__JOB_RESULT__',bool(result.get('ok')),site['name']))
    def shutdown(self):
        self.stop.set()
        for job in list(self.active.values()):
            self.control(job,'stop')
            try:job['process'].stdin.close()
            except (OSError,ValueError):pass
    def run(self):
        self.raw('PARALLEL STORES: up to 2 different websites together; source pacing, cooldowns and checks preserved.')
        for index in self.initial_deferred:
            self.deferred+=1;site=self.state['sites'][index]
            self.raw('WEBSITE COOLDOWN: '+site['name']+' — server delay is too long; Continue Saved Run later.')
            self.raw(('__JOB_RESULT__',False,site['name']))
        try:
            while self.ready or self.waiting or self.active:
                if self.stop.is_set():break
                paused=not self.pause.is_set()
                if paused!=self.paused:
                    self.paused=paused
                    for job in self.active.values():self.control(job,'pause' if paused else 'resume')
                if not paused:
                    for index,retry in list(self.waiting.items()):
                        if retry<=time.time():self.ready.append(index);del self.waiting[index]
                    for index in list(self.ready):
                        if len(self.active)>=self.max_processes:break
                        site=self.state['sites'][index]
                        if any(source_key(job['site'])==source_key(site) or output_key(job['site'])==output_key(site) for job in self.active.values()):continue
                        self.ready.remove(index)
                        try:self.start(index)
                        except Exception as exc:
                            broken=self.active.pop(index,None)
                            if broken:
                                process=broken['process']
                                if process.poll() is None:process.terminate()
                                process.wait(timeout=5)
                                for stream in (process.stdin,process.stdout):
                                    try:stream.close()
                                    except (OSError,ValueError):pass
                            self.failed+=1;self.mark(index,status='failed',detail='Process could not start: '+str(exc))
                            self.raw('FAILED: '+site['name']+': process could not start: '+str(exc));self.raw(('__JOB_RESULT__',False,site['name']))
                try:index,message=self.messages.get(timeout=.1);self.event(index,message)
                except queue.Empty:pass
                for index,job in list(self.active.items()):
                    if job['eof'] and job['process'].poll() is not None:self.finish(index)
            if not self.stop.is_set():
                def complete(state):state['status']='complete' if not self.store.pending(state) else 'incomplete'
                self.store.update(complete)
        finally:
            self.shutdown()
            deadline=time.monotonic()+2
            for index,job in list(self.active.items()):
                process=job['process']
                try:process.wait(timeout=max(.01,deadline-time.monotonic()))
                except subprocess.TimeoutExpired:process.terminate();process.wait(timeout=5)
                try:process.stdout.close()
                except (OSError,ValueError):pass
                self.mark(index,status='pending',detail='Stopped; successful requests retained')
            self.active.clear()


def enable_parallel_stores(g):
    if g.get('_chotu_parallel_stores'):return
    from resume_support import RunLease
    App=g['App'];previous_init=App.__init__;previous_close=App.close
    def init(self,root):
        previous_init(self,root);self._parallel_scheduler=None
        root.title(root.title()+' | Parallel Stores (2)')
    def launch(self,chosen,template,saved_state=None):
        if self.running or not chosen:return
        lease=RunLease(self.saved_runs.root)
        try:
            lease.acquire()
            if len({output_key(site) for site in chosen})!=len(chosen):raise RuntimeError('Give each selected store a distinct name; duplicate names would overwrite the same Excel file.')
            state=saved_state or self.saved_runs.fresh(chosen,template)
            pending=self.saved_runs.pending(state)
            if not pending:lease.release();self.refresh_saved_controls();return
            if not Path(state['template']).is_file():raise RuntimeError('Saved template snapshot missing. Choose a valid template for a new fresh run.')
        except Exception as exc:
            lease.release();g['messagebox'].showerror('Saved run',str(exc));return
        self._source_attempt=None;self._source_final_paused=False;self._non_source_failed=False
        self._saved_store_index=None;self._saved_progress_floor=0;self._saved_lease=lease
        self._saved_run_active=True;self._resuming_saved=bool(saved_state);self.running=True
        self.job_reviewed=any(site.get('reviewed',False) for site in state['sites']);self.job_failed=False;self.job_completed=0;self.job_expected=len(pending)
        self.refresh.configure(state='disabled');self.refresh_selected.configure(state='disabled');self.refresh_saved_controls()
        self.logbox.delete('1.0','end');self.progress_value.set(0);self.progress_label.configure(text='Starting parallel selected stores')
        g['_chotu_pause_gate'].set();self.pause_button.configure(state='normal',text='⏸ Pause')
        self._ui_started=time.monotonic();self._ui_products=0;self._ui_collections=0;self._ui_request_count=0;self._ui_429_count=0
        self.set_live_status('working',str(len(pending))+' stores queued; up to 2 different websites together')
        state=dict(state,_resuming=bool(saved_state))
        raw=getattr(self.events,'queue',self.events)
        scheduler=StoreScheduler(g['__file__'],self.saved_runs,state,g['OUT'],g['_chotu_pause_gate'],raw.put,retry_limit=self._source_auto_retry_limit,wait_limit=self._source_auto_wait_limit)
        self._parallel_scheduler=scheduler
        def work():
            try:scheduler.run()
            except Exception as exc:
                self._non_source_failed=True;raw.put('FAILED: parallel store scheduler: '+str(exc));raw.put(('__JOB_RESULT__',False,'Scheduler'))
            finally:
                self._source_final_paused=scheduler.deferred>0;self._non_source_failed=self._non_source_failed or scheduler.failed>0
                self._saved_run_active=False;self._saved_store_index=None;self._parallel_scheduler=None
                lease.release();self._saved_lease=None;raw.put('__DONE__')
        threading.Thread(target=work,name='chotu-store-scheduler',daemon=True).start()
    def close(self):
        scheduler=getattr(self,'_parallel_scheduler',None)
        if scheduler is None:return previous_close(self)
        if self.running and not g['messagebox'].askyesno('Extraction running','Close now? Successful requests remain saved. Continue Saved Run will retry unfinished store requests.'):
            return
        scheduler.shutdown();g['_chotu_pause_gate'].set();self.root.destroy()
    App.__init__=init;App._launch_jobs=launch;App.close=close;g['_chotu_parallel_stores']=True
