"""Bounded public-source product/shipping audit. No checkout or guessed shipping rates."""
import argparse,csv,hashlib,ipaddress,json,re,socket,time,urllib.error,urllib.parse,urllib.request,urllib.robotparser
from collections import deque
from datetime import datetime,timezone
from html.parser import HTMLParser
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
AGENT='ChotuAppShippingAudit/1.0'
MAX_BYTES=3*1024*1024
PRODUCT=re.compile(r'/products/[^/?#]+|-p\.\d+\.aspx',re.I)
POLICY=re.compile(r'shipping|delivery|faq|contact|terms',re.I)
CURRENCIES=r'(?:PKR|Rs\.?|₨)'
MONEY=re.compile(CURRENCIES+r'\s*([\d,]+(?:\.\d{1,2})?)|([\d,]+(?:\.\d{1,2})?)\s*'+CURRENCIES,re.I)
NOW=lambda:datetime.now(timezone.utc).isoformat()

def public_url(url):
 p=urllib.parse.urlsplit(url)
 if p.scheme!='https' or not p.hostname or p.username or p.password or p.port not in (None,443):raise ValueError('Only public HTTPS URLs without credentials are allowed')
 for item in socket.getaddrinfo(p.hostname,443,type=socket.SOCK_STREAM):
  if not ipaddress.ip_address(item[4][0]).is_global:raise ValueError('Private/internal network addresses are not allowed')
 return p

def canonical(url):
 p=urllib.parse.urlsplit(url);return urllib.parse.urlunsplit((p.scheme,p.netloc,p.path or '/',p.query,''))

def same_host(a,b):return urllib.parse.urlsplit(a).hostname.removeprefix('www.')==urllib.parse.urlsplit(b).hostname.removeprefix('www.')

class Page(HTMLParser):
 def __init__(self):
  super().__init__();self.links=[];self.parts=[];self.structured=[];self.title=[];self.hidden=0;self.in_title=False;self.anchor=None;self.json_script=None
 def handle_starttag(self,tag,attrs):
  a=dict(attrs)
  if tag=='script':
   if a.get('type')=='application/ld+json':self.json_script=[]
   self.hidden+=1
  if tag=='style':self.hidden+=1
  if tag=='title':self.in_title=True
  if tag=='a':self.anchor=[a.get('href',''),[]]
  if tag in ('p','div','li','br','h1','h2','h3','tr'):self.parts.append('\n')
 def handle_endtag(self,tag):
  if tag=='script' and self.json_script is not None:
   try:self.structured.append(json.loads(''.join(self.json_script)))
   except (ValueError,TypeError):pass
   self.json_script=None
  if tag in ('script','style'):self.hidden=max(0,self.hidden-1)
  if tag=='title':self.in_title=False
  if tag=='a' and self.anchor:
   self.links.append((self.anchor[0],' '.join(self.anchor[1]).strip()));self.anchor=None
  if tag in ('p','div','li','h1','h2','h3','tr'):self.parts.append('\n')
 def handle_data(self,data):
  if self.json_script is not None:self.json_script.append(data)
  if self.in_title:self.title.append(data)
  if not self.hidden:
   self.parts.append(data)
   if self.anchor:self.anchor[1].append(data)
 @property
 def lines(self):return [re.sub(r'\s+',' ',s).strip() for s in ''.join(self.parts).splitlines() if s.strip()]

def products_from_json(value):
 if isinstance(value,list):
  for v in value:yield from products_from_json(v)
 elif isinstance(value,dict):
  types=value.get('@type',[]);types=types if isinstance(types,list) else [types]
  if 'Product' in types:yield value
  for v in value.values():
   if isinstance(v,(list,dict)):yield from products_from_json(v)

def structured_rates(product,cities):
 """Only explicit PK shippingDestination + PKR shippingRate; retain conditions."""
 offers=product.get('offers',[]);offers=offers if isinstance(offers,list) else [offers];result=[]
 for offer in offers:
  if not isinstance(offer,dict):continue
  details=offer.get('shippingDetails',[]);details=details if isinstance(details,list) else [details]
  for detail in details:
   if not isinstance(detail,dict):continue
   dest=detail.get('shippingDestination',{});dest=dest if isinstance(dest,list) else [dest];rate=detail.get('shippingRate',{})
   if not isinstance(rate,dict) or rate.get('currency')!='PKR':continue
   amount=rate.get('value')
   if isinstance(amount,bool) or not re.fullmatch(r'\d+(?:\.\d+)?',str(amount)):continue
   for d in dest:
    if not isinstance(d,dict) or d.get('addressCountry') not in ('PK','Pakistan'):continue
    region=d.get('addressRegion') or d.get('addressLocality')
    # Postal-code/province and conditional offers cannot establish a city-wide fixed fee.
    conditional=any(k in d for k in ('postalCode','postalCodeRange')) or any(k in detail for k in ('eligibleTransactionVolume','shippingConditions')) or any(k in offer for k in ('eligibleTransactionVolume','eligibleQuantity'))
    targets=[c for c in cities if region and str(region).casefold()==c.casefold()] if region else list(cities)
    if conditional:targets=[]
    for city in targets:result.append((city,str(amount)+' PKR'))
 return result

def snippets(lines,cities):
 """Candidate statements only: never convert loose prose into verified numeric fees."""
 hits=[]
 for line in lines:
  if re.search(r'shipping|delivery|freight|courier',line,re.I) and (MONEY.search(line) or re.search(r'free|checkout|weight|city|cities|nationwide',line,re.I)):
   hits.append(line[:1800])
 return list(dict.fromkeys(hits))[:30]

class Client:
 def __init__(self,start,log,delay=1):
  self.start=start;self.log=log;self.delay=delay;self.last=0;self.policy=None
  owner=self
  class Redirect(urllib.request.HTTPRedirectHandler):
   def redirect_request(self,req,fp,code,msg,headers,newurl):
    newurl=urllib.parse.urljoin(req.full_url,newurl);public_url(newurl)
    if not same_host(owner.start,newurl):raise ValueError('External-domain redirect requires review: '+newurl)
    if owner.policy is not None and not owner.policy.can_fetch(AGENT,newurl):raise ValueError('Redirect disallowed by robots.txt')
    time.sleep(owner.delay)
    return super().redirect_request(req,fp,code,msg,headers,newurl)
  self.opener=urllib.request.build_opener(Redirect)
 def get(self,url,robots=False):
  public_url(url)
  if not same_host(self.start,url):raise ValueError('External host not approved')
  if not robots and (self.policy is None or not self.policy.can_fetch(AGENT,url)):raise ValueError('robots.txt does not authorize this URL')
  time.sleep(max(0,self.delay-(time.monotonic()-self.last)));self.last=time.monotonic();self.log('request',url=url)
  try:
   with self.opener.open(urllib.request.Request(url,headers={'User-Agent':AGENT,'Accept':'text/html,application/ld+json,text/plain'}),timeout=25) as response:
    raw=response.read(MAX_BYTES+1)
    if len(raw)>MAX_BYTES:raise ValueError('Response exceeds 3 MB')
    charset=response.headers.get_content_charset() or 'utf-8';text=raw.decode(charset,errors='replace');self.log('received',url=url,bytes=len(raw),sha256=hashlib.sha256(raw).hexdigest());return text,hashlib.sha256(raw).hexdigest()
  except Exception as error:self.log('fetch_error',url=url,error=str(error));raise
 def robots(self):
  p=urllib.parse.urlsplit(self.start);url=urllib.parse.urlunsplit((p.scheme,p.netloc,'/robots.txt','',''))
  try:text,_=self.get(url,robots=True)
  except urllib.error.HTTPError as error:
   if error.code!=404:raise
   text='User-agent: *\nAllow: /'
  self.policy=urllib.robotparser.RobotFileParser();self.policy.parse(text.splitlines());self.delay=max(self.delay,self.policy.crawl_delay(AGENT) or self.policy.crawl_delay('*') or 1)
  rate=self.policy.request_rate(AGENT) or self.policy.request_rate('*')
  if rate:self.delay=max(self.delay,rate.seconds/rate.requests)
  if self.delay>30:raise ValueError('Robots crawl delay exceeds this job limit')

def scan(start,cities,max_pages,log,deadline):
 client=Client(start,log);rows={};evidence=[];errors=[];seen=set();queue=deque([start]);discovered=set();policy_lines=[];complete=True
 try:client.robots()
 except Exception as error:return [],[],[{'url':start,'error':'Robots policy unavailable: '+str(error)}],False
 while queue and len(seen)<max_pages and time.monotonic()<deadline:
  url=queue.popleft()
  if url in seen:continue
  seen.add(url)
  try:
   raw,sha=client.get(url);page=Page();page.feed(raw)
   entities=list(products_from_json(page.structured));is_product=bool(PRODUCT.search(urllib.parse.urlsplit(url).path) or entities)
   if is_product:
    product=entities[0] if entities else {};name=product.get('name') or ' '.join(page.title).strip() or url
    row={'Store':urllib.parse.urlsplit(start).hostname,'Product Name':name,'Product Link':url,'Fixed Delivery Charges':'Not Verified',**{c:'Not Verified' for c in cities},'Verification':'Not Verified','Source URL':url,'Checked UTC':NOW(),'Evidence':'','Notes':''}
    for city,rate in structured_rates(product,cities):
     if row[city]=='Not Verified':row[city]=rate
     elif row[city]!=rate:row[city]='Conflicting rates - review'
    if any(row[c]!='Not Verified' for c in cities):row['Verification']='Structured source rate - review conditions';row['Notes']='Explicit PKR structured shipping data; checkout not tested. No order placed.'
    if len({row[c] for c in cities})==1 and row[cities[0]].endswith(' PKR'):row['Fixed Delivery Charges']=row[cities[0]]
    row['Evidence']=' | '.join(snippets(page.lines,cities))[:12000];rows[url]=row
   elif POLICY.search(url):policy_lines.extend((url,line) for line in snippets(page.lines,cities))
   evidence.append({'url':url,'sha256':sha,'checked_at_utc':NOW(),'product':is_product,'shipping_statements':snippets(page.lines,cities),'structured_products':entities[:1]})
   # Product links first; discover same-host public listings and policies, without cart/login requests.
   links=[]
   for href,name in page.links:
    target=canonical(urllib.parse.urljoin(url,href));p=urllib.parse.urlsplit(target)
    if p.scheme!='https' or not p.hostname or not same_host(start,target) or re.search(r'cart|checkout|account|login|logout|\.(?:jpg|png|pdf|zip)$',p.path,re.I):continue
    if PRODUCT.search(p.path):
     discovered.add(target);links.append((0,target))
     if target not in rows:rows[target]={'Store':urllib.parse.urlsplit(start).hostname,'Product Name':name or 'Name not captured','Product Link':target,'Fixed Delivery Charges':'Not Verified',**{c:'Not Verified' for c in cities},'Verification':'Product page not fetched','Source URL':url,'Checked UTC':NOW(),'Evidence':'','Notes':'Discovered link only; product/shipping not verified.'}
    elif POLICY.search(p.path):links.append((1,target))
    elif re.search(r'/collections|/products|category|categories|\.aspx',p.path,re.I) or p.query:links.append((2,target))
   queued=set(queue)
   for _,target in sorted(links):
    if target not in seen and target not in queued and len(queue)<10000:queue.append(target);queued.add(target)
  except Exception as error:errors.append({'url':url,'error':str(error)})
 if queue:complete=False
 for row in rows.values():
  if policy_lines:row['Notes']+=' Store policy statements available in Evidence sheet; not assumed to apply to this product.'
  if not complete:row['Notes']+=' Scan limited; full catalog coverage not established.'
 for url,line in policy_lines:evidence.append({'url':url,'shipping_statement':line,'scope':'Store policy; applicability needs review'})
 return list(rows.values()),evidence,errors,complete

def safe(value):
 value=str(value) if value is not None else ''
 return "'"+value if value.startswith(('=','+','-','@')) else value

def export(out,rows,evidence,errors,summary,cities):
 out.mkdir(parents=True,exist_ok=True);headers=['Store','Product Name','Product Link','Fixed Delivery Charges',*cities,'Verification','Source URL','Checked UTC','Evidence','Notes']
 with (out/'Products_Delivery_Audit.csv').open('w',encoding='utf-8-sig',newline='') as f:
  writer=csv.DictWriter(f,fieldnames=headers);writer.writeheader();writer.writerows({k:safe(row.get(k,'')) for k in headers} for row in rows)
 w=Workbook();s=w.active;s.title='Products';s.append(headers)
 for row in rows:s.append([safe(row.get(h,''))[:32700] for h in headers])
 for name,data,keys in [('Evidence',evidence,['url','shipping_statement','shipping_statements','scope','sha256']),('Errors',errors,['url','error'])]:
  sheet=w.create_sheet(name);sheet.append(keys)
  for row in data:sheet.append([safe(json.dumps(row.get(k,''),ensure_ascii=False) if isinstance(row.get(k),list) else row.get(k,''))[:32700] for k in keys])
 sheet=w.create_sheet('Run Summary');sheet.append(['Setting','Value'])
 for key,value in summary.items():sheet.append([key,safe(value)])
 for sheet in w:
  sheet.freeze_panes='C2';sheet.auto_filter.ref=sheet.dimensions
  for cell in sheet[1]:cell.fill=PatternFill('solid',fgColor='123F35');cell.font=Font(color='FFFFFF',bold=True)
  for col in sheet.columns:
   sheet.column_dimensions[col[0].column_letter].width=35
  sheet.row_dimensions[1].height=28
 w.save(out/'Products_Delivery_Audit.xlsx')
 for name,value in [('evidence.json',evidence),('errors.json',errors),('summary.json',summary)]: (out/name).write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')

def main():
 ap=argparse.ArgumentParser();ap.add_argument('--links',required=True);ap.add_argument('--cities',default='Karachi,Lahore,Islamabad,Gujranwala');ap.add_argument('--max-pages',type=int,default=100);ap.add_argument('--out',type=Path,default=Path('shipping-output'));args=ap.parse_args()
 cities=list(dict.fromkeys(c.strip() for c in args.cities.split(',') if c.strip()))
 if not cities or len(cities)>12 or any(len(c)>80 for c in cities):raise ValueError('Use 1-12 city names, each up to 80 characters')
 urls=list(dict.fromkeys(canonical(u) for u in re.split(r'[\s,]+',args.links) if u))
 if not urls or len(urls)>5 or not 1<=args.max_pages<=500:raise ValueError('Use 1-5 HTTPS links and 1-500 pages per store')
 args.out.mkdir(parents=True,exist_ok=True)
 def log(event,**fields):
  line=json.dumps({'time':NOW(),'event':event,**fields},ensure_ascii=False);print(line,flush=True)
  with (args.out/'activity.log').open('a',encoding='utf-8') as f:f.write(line+'\n')
 rows=[];evidence=[];errors=[];coverage=[];deadline=time.monotonic()+45*60
 for url in urls:
  log('store.started',url=url)
  try:public_url(url);r,e,x,complete=scan(url,cities,args.max_pages,log,deadline);rows.extend(r);evidence.extend(e);errors.extend(x);coverage.append({'url':url,'queue_exhausted':complete,'products':len(r)});log('store.finished',url=url,products=len(r),errors=len(x))
  except Exception as error:errors.append({'url':url,'error':str(error)});coverage.append({'url':url,'queue_exhausted':False,'products':0})
 # Multiple submitted links can discover the same product. Prefer successfully fetched rows.
 unique={}
 for row in rows:
  if row['Product Link'] not in unique or row['Verification']!='Product page not fetched':unique[row['Product Link']]=row
 rows=list(unique.values());summary={'checked_at_utc':NOW(),'products':len(rows),'errors':len(errors),'coverage':coverage,'max_pages_per_store':args.max_pages,'cities':cities,'notice':'Not Verified never means free. Queue exhaustion is not proof of complete site inventory. Website blocks are reported, not bypassed. No checkout or purchase performed.'}
 export(args.out,rows,evidence,errors,summary,cities);log('export.finished',products=len(rows),errors=len(errors),path=str(args.out/'Products_Delivery_Audit.xlsx'))
 report=Path(__import__('os').environ.get('GITHUB_STEP_SUMMARY',str(args.out/'run-summary.md')))
 with report.open('a',encoding='utf-8') as f:f.write(f'## ChotuApp delivery audit\n\nProducts discovered: **{len(rows)}** · Errors: **{len(errors)}**\n\nDownload **ChotuApp-Shipping-Results** from this run’s Artifacts section. Excel, CSV, source evidence and copyable logs are included.\n\nUnconfirmed rates remain **Not Verified**. A completed job does not guarantee a complete catalog or verified shipping.\n')
if __name__=='__main__':main()
