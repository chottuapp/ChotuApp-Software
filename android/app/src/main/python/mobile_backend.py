"""Embedded Android worker. No Termux, subprocess, desktop UI or shell required."""
import base64,hashlib,importlib.util,json,sys,time,threading,contextlib,datetime
from pathlib import Path
import requests
REPO='chottuapp/ChotuApp-Software'
class Paused(KeyboardInterrupt):pass

def load(name,path):
 spec=importlib.util.spec_from_file_location(name,path);module=importlib.util.module_from_spec(spec);sys.modules[name]=module;spec.loader.exec_module(module);return module

def run(root,data,token,bridge=None):
 if bridge is None:
  from java import jclass
  bridge=jclass("com.chotuapp.mobile.FetchService")
 with reporting(Path(data),bridge,token):return _run(root,data,token,bridge)

def _run(root,data,token,bridge=None):
 if bridge is None:
  from java import jclass
  bridge=jclass("com.chotuapp.mobile.FetchService")
 root=Path(root);data=Path(data);sys.path.insert(0,str(root))
 def check():
  if bridge.isPaused():raise Paused()
 class PublicSession:
  def __init__(self):self.session=requests.Session();self.session.headers['User-Agent']='Mozilla/5.0 ChotuApp-PublicCatalog/1.0'
  def get(self,url,**kwargs):
   check();response=self.session.get(url,timeout=(15,45));check();return response
 state=data/'local-status.json'
 try:previous=json.loads(state.read_text())
 except (OSError,ValueError):previous={}
 if previous.get('status')!='exported_review_required':
  worker=load('android_approved_worker',root/'approved_worker.py')
  code=worker.fetch(root,data,PublicSession())
  if code:return json.dumps({'status':'paused' if code==130 else 'attention','message':'Saved work retained. See activity for the source error.'})
 check()
 if hasattr(bridge,'githubToken'):token=str(bridge.githubToken() or '')
 if not token:return json.dumps({'status':'ready_to_upload','message':'Excel saved on this phone. Connect GitHub to queue Rule 804 and Drive upload.'})
 sha=upload(data,token,check)
 return json.dumps({'status':'cloud_queued','source_sha256':sha,'message':'Workbook queued for Rule 804 and Drive verification. Drive completion is shown separately.'})

def upload(data,token,check=lambda:None):
 state=json.loads((data/'local-status.json').read_text());files=[Path(p) for p in state.get('files',[])]
 if state.get('status')!='exported_review_required' or not files:raise ValueError('No completed workbook')
 file=max(files,key=lambda p:p.stat().st_mtime)
 if not file.resolve().is_relative_to((data/'files').resolve()) or file.suffix!='.xlsx':raise ValueError('Unexpected output path')
 content=file.read_bytes()
 if len(content)>32*1024*1024:raise ValueError('File exceeds safe repository upload limit; retained on phone')
 sha=hashlib.sha256(content).hexdigest();path='catalog-inputs/khaadi-mobile/'+sha[:16]+'/Khaadi.xlsx'
 session=requests.Session();session.headers.update(Authorization='Bearer '+token,Accept='application/vnd.github+json')
 def put(path,body):
  check();url='https://api.github.com/repos/'+REPO+'/contents/'+path
  response=session.get(url,params={'ref':'main'},timeout=45)
  if response.status_code not in (200,404):raise RuntimeError('GitHub access unavailable ('+str(response.status_code)+'). File retained on phone.')
  old=response.json() if response.status_code==200 else {}
  blob=hashlib.sha1(b'blob '+str(len(body)).encode()+b'\0'+body).hexdigest()
  if old.get('sha')==blob:return
  payload={'message':'Supply mobile Khaadi workbook for checked Rule 804 delivery','branch':'main','content':base64.b64encode(body).decode()}
  if old.get('sha'):payload['sha']=old['sha']
  check();response=session.put(url,json=payload,timeout=180)
  if response.status_code not in (200,201):raise RuntimeError('GitHub upload failed ('+str(response.status_code)+'). File retained on phone.')
 put(path,content)
 try:session_meta=json.loads((data/'session-meta.json').read_text())
 except (OSError,ValueError):session_meta={}
 marker=json.dumps({'store':'Khaadi','fresh_campaign_id':session_meta.get('fresh_campaign_id'),'input_path':path,'sha256':sha,'coverage':'Public default-page discovery; full catalog coverage unverified'},indent=2).encode()
 put('catalog-inputs/khaadi-mobile/latest.json',marker)
 return sha

def dashboard():
 result={}
 for key,branch in (('runners','runner-status-reports'),('drive','drive-sync-reports')):
  try:
   response=requests.get('https://raw.githubusercontent.com/'+REPO+'/'+branch+'/latest.json',params={'check':time.time_ns()},timeout=25);response.raise_for_status();result[key]=response.json()
  except requests.RequestException:result[key]={'unavailable':True}
 result.get('runners',{})['active_runners']=sum(run.get('status')=='in_progress' for run in result.get('runners',{}).get('runs',[]))
 try:
  response=requests.get('https://raw.githubusercontent.com/'+REPO+'/android-builds/studio-release.json',params={'check':time.time_ns()},timeout=15);response.raise_for_status();result['app_update']=response.json()
 except (requests.RequestException,ValueError):result['app_update']={'unavailable':True}
 try:
  response=requests.get('https://raw.githubusercontent.com/'+REPO+'/main/cloud/audit804-release.json',params={'check':time.time_ns()},timeout=15);response.raise_for_status();result['fresh_release']=response.json()
 except (requests.RequestException,ValueError):pass
 from store_progress import rows
 result['store_progress']=rows(result.get('runners',{}))
 return json.dumps(result)

def verify_token(token):
 response=requests.get('https://api.github.com/repos/'+REPO,headers={'Authorization':'Bearer '+token,'Accept':'application/vnd.github+json'},timeout=30)
 if response.status_code!=200:raise ValueError('GitHub connection failed. Check token access to ChotuApp-Software.')
 if not response.json().get('permissions',{}).get('push'):raise ValueError('This GitHub account needs write access to ChotuApp-Software.')
 return 'Connected'

def self_test(root,data):
 """On-device fixture export: embedded Python, SQLite, XLSX and headless source."""
 root=Path(root);sys.path.insert(0,str(root))
 import khaadi_source_support as source
 from bs4 import BeautifulSoup
 from openpyxl import load_workbook
 fixture=root/'test-fixtures';soup=BeautifulSoup((fixture/'master.html').read_text(),'html.parser')
 for node in soup.select('.set-item')[0].select('[data-attr="size"] input'):
  if node['data-attr-value']=='00S':node.parent.decompose()
 from types import SimpleNamespace
 product=source.normalize_product(str(soup),source.BASE+'/black-cambric-kurta/1-26-320-A-C1-VG_MULTI.html',lambda url:SimpleNamespace(text=(fixture/('xs.html' if '=0XS' in url else 'l.html')).read_text()))
 original=source.PublicCatalog.catalog
 try:
  source.PublicCatalog.catalog=lambda reader:([product],False)
  helper=load('android_fixture_worker',root/'approved_worker.py')
  if helper.fetch(root,Path(data),object())!=0:raise ValueError('On-device fixture export failed')
 finally:source.PublicCatalog.catalog=original
 state=json.loads((Path(data)/'local-status.json').read_text());book=load_workbook(state['files'][0])
 try:
  if [c.value for c in next(book['Products'].iter_rows())][:3]!=['shop_code','product_id','Category']:raise ValueError('Template schema changed')
  if book['Variant Combinations'].max_row!=3:raise ValueError('Selected fixture variants missing')
 finally:book.close()
 return 'Embedded export verified'

def read_state(data):
 try:return json.loads((Path(data)/'local-status.json').read_text())
 except (OSError,ValueError):return {}

def safe_log(value,token=''):
 import re
 value=str(value)
 if token:value=value.replace(token,'[redacted]')
 value=re.sub(r'(?:github_pat_|gh[pousr]_)[A-Za-z0-9_]+','[redacted]',value)
 if re.search(r'authorization|access_token|refresh_token|client_secret|password|set-cookie|cookie:',value,re.I):return '[credential-bearing line omitted]'
 return re.sub(r'/data/(?:user/\d+|data)/com\.chotuapp\.(?:mobile|studio)/[^ \n]*','[app-data]',value)[:900]

def log_tail(data,token=''):
 path=Path(data)/'fetch-log.txt'
 try:
  with path.open('rb') as stream:
   stream.seek(max(0,path.stat().st_size-18000));raw=stream.read(18000).decode('utf-8',errors='replace')
  return [safe_log(line,token) for line in raw.splitlines()[-60:]]
 except OSError:return []

def mobile_report(data,token,active,paused=False):
 state=read_state(data);progress=state.get('last_progress_epoch');status=state.get('status','ready')
 try:campaign=json.loads((Path(data)/'session-meta.json').read_text()).get('fresh_campaign_id')
 except (OSError,ValueError):campaign=None
 if paused and active:status='pausing'
 elif status=='running' and not active:status='stopped'
 return {'store':'Khaadi','execution':'Android app','app_version':'1.2-studio','fresh_campaign_id':campaign,'observed_utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),'worker_active':bool(active),'status':status,'worker_updated':state.get('updated'),'last_event':safe_log(state.get('last_event',''),token),'error':safe_log(state.get('error',''),token),'progress_age_seconds':max(0,round(time.time()-progress)) if isinstance(progress,(int,float)) else None,'log_tail':log_tail(data,token)}

def telemetry_status(data,**record):
 data=Path(data);data.mkdir(parents=True,exist_ok=True);path=data/'telemetry-status.json'
 # Unique temporary files avoid collisions with manual Send now and periodic sync.
 import tempfile,os
 with tempfile.NamedTemporaryFile(mode='w',dir=data,delete=False,encoding='utf-8') as stream:
  json.dump(record,stream);temporary=stream.name
 os.replace(temporary,path)

def sync_logs(data,token,active=False,paused=False):
 data=Path(data);now=datetime.datetime.now(datetime.timezone.utc).isoformat()
 if not token:
  record={'status':'connection_required','message':'Connect GitHub once to enable remote logs. Fetching can continue locally.','attempted_utc':now};telemetry_status(data,**record);return json.dumps(record)
 try:
  session=requests.Session();session.headers.update(Authorization='Bearer '+token,Accept='application/vnd.github+json');base='https://api.github.com/repos/'+REPO;branch='local-run-status'
  ref=session.get(base+'/git/ref/heads/'+branch,timeout=15)
  if ref.status_code==404:
   main=session.get(base+'/git/ref/heads/main',timeout=15)
   if main.status_code!=200:raise RuntimeError('GitHub HTTP '+str(main.status_code))
   created=session.post(base+'/git/refs',json={'ref':'refs/heads/'+branch,'sha':main.json()['object']['sha']},timeout=15)
   if created.status_code not in (201,422):raise RuntimeError('GitHub HTTP '+str(created.status_code))
  elif ref.status_code!=200:raise RuntimeError('GitHub HTTP '+str(ref.status_code))
  report=mobile_report(data,token,active,paused)
  # Separate Android path prevents an old PC heartbeat overwriting phone logs.
  url=base+'/contents/local/khaadi-android.json';old=session.get(url,params={'ref':branch},timeout=15)
  if old.status_code not in (200,404):raise RuntimeError('GitHub HTTP '+str(old.status_code))
  payload={'branch':branch,'message':'Update redacted Android Khaadi progress and logs','content':base64.b64encode(json.dumps(report).encode()).decode()}
  if old.status_code==200:payload['sha']=old.json()['sha']
  response=session.put(url,json=payload,timeout=20)
  if response.status_code not in (200,201):raise RuntimeError('GitHub HTTP '+str(response.status_code))
  record={'status':'synced','last_success_utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),'message':'Remote logs synced. Automatic sync every 30 seconds while fetching.'}
 except (requests.RequestException,ValueError,KeyError,OSError,RuntimeError) as exc:
  # Do not expose HTTP response bodies or exceptions which might include credentials.
  reason=str(exc) if isinstance(exc,RuntimeError) and str(exc).startswith('GitHub HTTP ') else 'Network or repository access unavailable'
  if reason=='GitHub HTTP 403':reason='GitHub denied access. Allow Contents: Read and write for ChotuApp-Software'
  elif reason=='GitHub HTTP 401':reason='GitHub token expired or invalid. Reconnect in Settings'
  record={'status':'retry_pending','message':reason+'; local fetching is unchanged.','attempted_utc':now}
 telemetry_status(data,**record);return json.dumps(record)

@contextlib.contextmanager
def reporting(data,bridge,token):
 stop=threading.Event()
 def current_token():
  method=getattr(bridge,'githubToken',None)
  return str(method() or '') if method else token
 def send(terminal=False):
  active=False if terminal else bool(bridge.isActive()) if hasattr(bridge,'isActive') else not bridge.isPaused()
  sync_logs(data,current_token(),active,bridge.isPaused())
 def loop():
  while not stop.is_set():send();stop.wait(30)
  send(terminal=True)
 thread=threading.Thread(target=loop,daemon=True);thread.start()
 try:yield
 finally:stop.set();thread.join(timeout=90)
