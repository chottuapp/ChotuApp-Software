"""Optional Windows login startup and one fresh selected-store run per local day."""
import copy
import datetime as dt
import json
import os
from pathlib import Path
import sys
import tempfile

STARTUP_MARKER = 'REM CHOTUAPP_DAILY_AUTOMATION_V1'


def parse_time(value):
    parts=value.strip().split(':')
    if len(parts)!=2 or any(not p.isdigit() for p in parts):raise ValueError('Enter daily time as HH:MM, for example 09:00.')
    hour,minute=map(int,parts)
    if not 0<=hour<24 or not 0<=minute<60:raise ValueError('Daily time must be between 00:00 and 23:59.')
    return f'{hour:02d}:{minute:02d}'


def save_json(path, value):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    temporary=None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent,delete=False,mode='w',encoding='utf-8') as stream:
            temporary=Path(stream.name);json.dump(value,stream,indent=2);stream.flush();os.fsync(stream.fileno())
        os.replace(temporary,path)
    finally:
        if temporary:temporary.unlink(missing_ok=True)


def startup_file():
    if os.name!='nt':raise RuntimeError('Laptop login startup is supported on Windows. Daily runs while this app is open work on other systems.')
    return Path(os.environ['APPDATA'])/'Microsoft/Windows/Start Menu/Programs/Startup/ChotuApp Daily Auto Fetch.cmd'


def configure_startup(enabled, app_file, destination=None, executable=None):
    destination=Path(destination) if destination else startup_file()
    if destination.exists() and STARTUP_MARKER not in destination.read_text(encoding='utf-8'):
        raise RuntimeError('A different startup file uses this name; it was preserved.')
    if not enabled:
        destination.unlink(missing_ok=True);return
    executable=Path(executable or sys.executable)
    pythonw=executable.with_name('pythonw.exe')
    if pythonw.is_file():executable=pythonw
    def quoted(path):
        text=str(Path(path).resolve())
        if any(c in text for c in '\r\n"'):raise ValueError('Unsupported startup path.')
        return '"'+text.replace('%','%%')+'"'
    payload='@echo off\r\n'+STARTUP_MARKER+'\r\nstart "" /min '+quoted(executable)+' '+quoted(app_file)+' --chotu-auto\r\n'
    destination.parent.mkdir(parents=True,exist_ok=True)
    destination.write_text(payload,encoding='utf-8')


def choose_action(settings, now, login_pending, saved_state):
    """Resume today's own interrupted run; other days always start fresh."""
    if not settings.get('enabled'):return None
    today=now.date().isoformat();inflight=settings.get('inflight') or {}
    if inflight.get('date')==today:
        if saved_state and saved_state.get('id')==inflight.get('run_id') and saved_state.get('status')!='complete':
            return 'resume'
    if settings.get('last_attempt_date')==today:return None
    if login_pending and settings.get('run_on_open',True):return 'fresh'
    if now.strftime('%H:%M')>=parse_time(settings.get('time','09:00')):return 'fresh'
    return None


def enable_automation(g):
    App=g['App'];tk=g['tk'];messagebox=g['messagebox']
    if getattr(App,'_daily_automation_installed',False):return
    if not hasattr(App,'continue_saved_run'):raise RuntimeError('Saved Runs support is required before Daily Auto Fetch.')
    import customtkinter as ctk
    from resume_support import RunLease
    original_init=App.__init__;original_poll=App.poll

    def init(self,root):
        original_init(self,root)
        self._auto_path=Path(g['__file__']).resolve().parent/'daily_automation.json'
        try:
            settings=json.loads(self._auto_path.read_text(encoding='utf-8'))
            if not isinstance(settings,dict):raise ValueError('Invalid automation settings')
            parse_time(settings.get('time','09:00'))
        except (OSError,ValueError):settings={}
        self._auto_settings=settings;self._auto_active=False;self._auto_login_pending=True;self._auto_lease=None
        self.auto_enabled=tk.BooleanVar(value=bool(settings.get('enabled',False)))
        self.auto_login=tk.BooleanVar(value=bool(settings.get('startup',False)))
        self.auto_on_open=tk.BooleanVar(value=bool(settings.get('run_on_open',True)))
        self.auto_time=tk.StringVar(value=settings.get('time','09:00'))
        self.auto_panel=ctk.CTkFrame(root,fg_color='#10271b',corner_radius=8)
        self.auto_panel.pack(side='top',fill='x',padx=18,pady=(2,4))
        first=ctk.CTkFrame(self.auto_panel,fg_color='transparent');first.pack(fill='x',padx=12,pady=(7,2))
        ctk.CTkLabel(first,text='Daily Auto Fetch',font=ctk.CTkFont(size=14,weight='bold')).pack(side='left',padx=(0,16))
        ctk.CTkCheckBox(first,text='Enable daily',variable=self.auto_enabled,width=125).pack(side='left',padx=7)
        ctk.CTkCheckBox(first,text='Open at laptop login',variable=self.auto_login,width=180).pack(side='left',padx=7)
        self.auto_save_button=ctk.CTkButton(first,text='Save selected stores',width=155,command=self.save_daily_automation)
        self.auto_save_button.pack(side='right',padx=(10,0))
        second=ctk.CTkFrame(self.auto_panel,fg_color='transparent');second.pack(fill='x',padx=12,pady=(2,7))
        ctk.CTkCheckBox(second,text='Run on opening if today is pending',variable=self.auto_on_open,width=285).pack(side='left',padx=(0,12))
        ctk.CTkLabel(second,text='Time').pack(side='left',padx=(0,5))
        ctk.CTkEntry(second,textvariable=self.auto_time,width=65).pack(side='left',padx=(0,12))
        self.auto_status=ctk.CTkLabel(second,text='',anchor='w',wraplength=380);self.auto_status.pack(side='left',fill='x',expand=True)
        self._wheel_workspace.pack_forget();self._wheel_workspace.pack(fill='both',expand=True)
        if hasattr(self,'_tag_wheel_widgets'):self._tag_wheel_widgets(self.auto_panel)
        self.update_daily_status()
        root.after(15000,self._daily_auto_tick)

    def status(self,detail=None):
        if detail:self.auto_status.configure(text=detail);return
        s=self._auto_settings
        if not s.get('enabled'):text='Off — select stores and template, then Save selected stores'
        elif self._auto_active:text='Automatic run active — each completed file saves to ANAS_DATA_FETCH'
        else:text=f"{len(s.get('sites',[]))} stores • {s.get('time','09:00')} daily • {s.get('last_result','Ready')}"
        self.auto_status.configure(text=text)

    def save(self):
        if self.running:
            messagebox.showinfo('Daily automation','Finish or close the current run before changing the daily schedule.');return
        try:
            enabled=self.auto_enabled.get();startup=self.auto_login.get();schedule=parse_time(self.auto_time.get())
            chosen=[dict(self.sites[i]) for i in self._selected_indices()]
            template=Path(self.template.get())
            if enabled and (not chosen or not template.is_file()):raise ValueError('Choose a valid template and tick the stores to run daily.')
            if enabled:
                from openpyxl import load_workbook
                book=load_workbook(template,read_only=True)
                try:
                    if not {'Products','Product Images','Variant Combinations'}.issubset(book.sheetnames):raise ValueError('Select your ChotuApp product template.')
                finally:book.close()
            previous=copy.deepcopy(self._auto_settings)
            updated=copy.deepcopy(previous)
            updated.update(enabled=enabled,startup=startup,time=schedule,run_on_open=self.auto_on_open.get(),sites=chosen,template=str(template.resolve()))
            if os.name=='nt' or startup:configure_startup(startup,Path(g['__file__']))
            try:save_json(self._auto_path,updated)
            except Exception:
                if os.name=='nt':configure_startup(bool(previous.get('startup')),Path(g['__file__']))
                raise
            self._auto_settings=updated;self._auto_login_pending=True
            if not enabled and self._auto_lease:self._auto_lease.release();self._auto_lease=None
            self.update_daily_status('Saved — daily runs use these stores and template' if enabled else 'Daily automatic fetching disabled')
        except Exception as exc:messagebox.showerror('Daily automation',str(exc))

    def tick(self):
        try:
            if not self._auto_settings.get('enabled') or self.running:return
            if self._auto_lease is None:
                lease=RunLease(self._auto_path.parent/'daily_instance')
                lease.path.parent.mkdir(exist_ok=True)
                try:lease.acquire()
                except (OSError,RuntimeError):
                    self.update_daily_status('Daily automation is already owned by another open window');return
                self._auto_lease=lease
            state=copy.deepcopy(self.saved_runs.ui_snapshot)
            now=dt.datetime.now()
            action=choose_action(self._auto_settings,now,self._auto_login_pending,state)
            if not action:return
            self._auto_login_pending=False
            template=Path(self._auto_settings.get('template',''))
            if not template.is_file():
                self.update_daily_status('Template missing — Browse file and Save selected stores');return
            # An unfinished manual run is preserved; never replace its saved queue automatically.
            if action=='fresh' and state and self.saved_runs.pending(state):
                inflight=self._auto_settings.get('inflight') or {}
                if state.get('id')!=inflight.get('run_id'):
                    self.update_daily_status('Manual saved run pending — Continue it first, or explicitly start a New Fresh Run');return
            settings=copy.deepcopy(self._auto_settings)
            settings.update(last_attempt_date=now.date().isoformat(),last_result='Starting')
            save_json(self._auto_path,settings);self._auto_settings=settings
            self._auto_active=True
            if action=='resume':
                self._launch_jobs(state['sites'],state['template'],saved_state=state)
            else:self._launch_jobs(copy.deepcopy(settings['sites']),str(template))
            if self.running:
                snapshot=self.saved_runs.ui_snapshot or {}
                settings['inflight']={'date':now.date().isoformat(),'run_id':snapshot.get('id')}
                save_json(self._auto_path,settings);self._auto_settings=settings
                self.update_daily_status()
            else:
                self._auto_active=False;settings['last_result']='Could not start — check template / saved run';save_json(self._auto_path,settings)
                self.update_daily_status()
        except Exception as exc:self.update_daily_status('Daily automation needs attention: '+str(exc))
        finally:self.root.after(15000,self._daily_auto_tick)

    def poll(self):
        original_poll(self)
        if self._auto_active and not self.running:
            settings=copy.deepcopy(self._auto_settings)
            ok=not self.job_failed and self.job_completed==self.job_expected
            settings['last_result']='Saved — review audit reports' if ok and self.job_reviewed else 'Saved successfully' if ok else 'Run incomplete — use Continue Saved Run'
            if ok:settings['last_completed_date']=dt.date.today().isoformat();settings.pop('inflight',None)
            # Avoid an automatic failure retry loop; user can explicitly Continue.
            elif settings.get('inflight'):settings['inflight']['date']='failed'
            self._auto_active=False;self._auto_settings=settings
            try:save_json(self._auto_path,settings)
            except OSError as exc:settings['last_result']='Settings save failed: '+str(exc)
            self.update_daily_status()

    def legacy_tick(self):
        # The old unconditional 24-hour refresh conflicts with explicit schedule.
        self.root.after(60000,self._top20_auto_tick)

    App.__init__=init;App.poll=poll;App._top20_auto_tick=legacy_tick
    App.save_daily_automation=save;App._daily_auto_tick=tick;App.update_daily_status=status
    App._daily_automation_installed=True
