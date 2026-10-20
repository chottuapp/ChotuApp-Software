"""Upgrade the observed original export layout without rewriting workbook data/style parts."""
import copy,io,re,zipfile,xml.etree.ElementTree as ET
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils import column_index_from_string,get_column_letter
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main';REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
Q=lambda name:'{'+NS+'}'+name

def formula(text,products=False):
 if not text:return text
 tokens=Tokenizer('='+text);changed=False
 for token in tokens.items:
  if token.type!='OPERAND' or token.subtype!='RANGE':continue
  value=token.value;prefix='';address=value;target=products
  if '!' in value:
   sheet,address=value.rsplit('!',1);target=sheet.strip("'").casefold()=='products';prefix=sheet+'!'
  if not target:continue
  # Structured references and deleted category references need explicit review.
  if '[' in address:raise ValueError('Legacy structured formula requires review')
  match=re.fullmatch(r'(\$?[A-Z]+)(\$?\d+)?(?::(\$?[A-Z]+)(\$?\d+)?)?',address,re.I)
  if not match:continue
  left=column_index_from_string(match[1].replace('$',''));right=column_index_from_string((match[3] or match[1]).replace('$',''))
  if min(left,right)<=2<=max(left,right):raise ValueError('Formula references removed legacy category column')
  def col(value):return ('$' if value.startswith('$') else '')+('B' if value.replace('$','').upper()=='A' else value.replace('$',''))
  rewritten=prefix+col(match[1])+(match[2] or '')+(':'+col(match[3])+(match[4] or '') if match[3] else '')
  if rewritten!=token.value:token.value=rewritten;changed=True
 return tokens.render()[1:] if changed else text

def refs(value):
 output=[]
 for part in value.split():
  match=re.fullmatch(r'(\$?[A-Z]+)(\$?\d+)?(?::(\$?[A-Z]+)(\$?\d+)?)?',part,re.I)
  if not match:raise ValueError('Unsupported legacy worksheet reference')
  left=column_index_from_string(match[1].replace('$',''));right=column_index_from_string((match[3] or match[1]).replace('$',''))
  cols=[2 if i==1 else i for i in range(left,right+1) if i!=2]
  if not cols:continue
  if max(cols)-min(cols)+1!=len(set(cols)):raise ValueError('Non-contiguous legacy reference')
  first=get_column_letter(min(cols))+(match[2] or '');last=get_column_letter(max(cols))+(match[4] or match[2] or '')
  output.append(first+(':'+last if match[3] else ''))
 return ' '.join(output)

def inline(cell,value):
 for child in list(cell):cell.remove(child)
 if value=='':cell.attrib.pop('t',None);return
 cell.set('t','inlineStr');ET.SubElement(ET.SubElement(cell,Q('is')),Q('t')).text=value

def serialize(tree,original):
 from xml.sax.saxutils import quoteattr
 namespaces={}
 for _,(prefix,uri) in ET.iterparse(io.BytesIO(original),events=('start-ns',)):
  if prefix in namespaces and namespaces[prefix]!=uri:raise ValueError('Scoped namespace requires reviewed legacy migration')
  namespaces[prefix]=uri
  ET.register_namespace(prefix,uri)
 payload=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
 start=payload.find(b'<',payload.find(b'?>')+2);end=payload.find(b'>',start);opening=payload[start:end]
 extra=[]
 for prefix,uri in namespaces.items():
  key='xmlns'+(':'+prefix if prefix else '')
  if not re.search(rb'\b'+re.escape(key.encode())+rb'=',opening):extra.append(' '+key+'='+quoteattr(uri))
 return payload[:end]+''.join(extra).encode()+payload[end:]

def normalize(source,out,code=None):
 book=load_workbook(source,read_only=True,data_only=False)
 try:headers=list(next(book['Products'].iter_rows(values_only=True)))
 finally:book.close()
 if headers[:3]==['shop_code','product_id','Category']:return source,False
 if headers[:3]!=['product_id','category','subcategory']:raise ValueError('Unrecognized original legacy schema')
 with zipfile.ZipFile(source) as incoming:
  workbook=ET.fromstring(incoming.read('xl/workbook.xml'));rels=ET.fromstring(incoming.read('xl/_rels/workbook.xml.rels'));sheet=next(s for s in workbook.find(Q('sheets')) if s.get('name')=='Products');target=next(r.get('Target') for r in rels if r.get('Id')==sheet.get('{'+REL+'}id'));part=target.lstrip('/') if target.startswith('/') else 'xl/'+target
  tree=ET.fromstring(incoming.read(part))
  if tree.find(Q('tableParts')) is not None:raise ValueError('Legacy Products table requires reviewed migration')
  for row in tree.find(Q('sheetData')):
   cells=[c for c in row if c.tag==Q('c')];old_a=next((c for c in cells if re.fullmatch('A[0-9]+',c.get('r',''))),None);old_b=next((c for c in cells if re.fullmatch('B[0-9]+',c.get('r',''))),None)
   number=row.get('r');new=copy.deepcopy(old_b if old_b is not None else old_a) if old_b is not None or old_a is not None else ET.Element(Q('c'))
   new.set('r','A'+number);inline(new,'shop_code' if number=='1' else code or '')
   if old_b is not None:row.remove(old_b)
   if old_a is not None:old_a.set('r','B'+number)
   row.insert(0,new)
   if number=='1':
    category=next(c for c in row if c.get('r')=='C1');inline(category,'Category')
  cols=tree.find(Q('cols'))
  if cols is not None:
   rewritten=[]
   for col in cols:
    mapped=sorted(2 if n==1 else 1 if n==2 else n for n in range(int(col.get('min')),int(col.get('max'))+1));groups=[]
    for n in mapped:
     if groups and n==groups[-1][-1]+1:groups[-1].append(n)
     else:groups.append([n])
    for group in groups:
     item=copy.deepcopy(col);item.set('min',str(group[0]));item.set('max',str(group[-1]));rewritten.append(item)
   cols[:]=sorted(rewritten,key=lambda c:int(c.get('min')))
  for parent in tree.iter():
   for node in list(parent):
    tag=node.tag.rsplit('}',1)[-1]
    if tag in ('dataValidation','conditionalFormatting','hyperlink','mergeCell'):
     key='sqref' if 'sqref' in node.attrib else 'ref'
     if key not in node.attrib:continue
     mapped=refs(node.get(key))
     if mapped:node.set(key,mapped)
     else:parent.remove(node)
  for validations in tree.iter(Q('dataValidations')):validations.set('count',str(len(validations)))
  auto=tree.find(Q('autoFilter'))
  if auto is not None:
   for column in list(auto):
    if column.tag==Q('filterColumn'):
     if column.get('colId')=='1':auto.remove(column)
     elif column.get('colId')=='0':column.set('colId','1')
  changes={part:tree}
  for info in incoming.infolist():
   if info.filename.startswith('xl/worksheets/') and info.filename.endswith('.xml'):
    xml=tree if info.filename==part else ET.fromstring(incoming.read(info.filename));changed=info.filename==part
    for node in xml.iter():
     if node.tag.rsplit('}',1)[-1] in ('f','formula','formula1','formula2') and node.text:
      rewritten=formula(node.text,info.filename==part)
      if rewritten!=node.text:node.text=rewritten;changed=True
    if changed:changes[info.filename]=xml
  for node in workbook.iter(Q('definedName')):
   if node.get('name') in ('_xlnm._FilterDatabase','_xlnm.Print_Area'):continue
   if node.text:
    rewritten=formula(node.text)
    if rewritten!=node.text:node.text=rewritten;changes['xl/workbook.xml']=workbook
  out=Path(out);out.parent.mkdir(parents=True,exist_ok=True)
  with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as destination:
   for info in incoming.infolist():destination.writestr(info,serialize(changes[info.filename],incoming.read(info.filename)) if info.filename in changes else incoming.read(info.filename))
 # Read-back verifies every original data/formula value except the authorized category removal.
 old=load_workbook(source,read_only=True,data_only=False);new=load_workbook(out,read_only=True,data_only=False)
 try:
  if old.sheetnames!=new.sheetnames:raise ValueError('Legacy sheet identity changed')
  for name in old.sheetnames:
   if old[name].max_row!=new[name].max_row or old[name].max_column!=new[name].max_column:raise ValueError('Legacy worksheet dimensions changed')
   for number,(before,after) in enumerate(zip(old[name].iter_rows(),new[name].iter_rows()),1):
    if name=='Products' and after[0].value!=('shop_code' if number==1 else code or None):raise ValueError('Shop code migration verification failed')
    for index,cell in enumerate(before,1):
     if name=='Products' and index==2:continue
     column=2 if name=='Products' and index==1 else index
     expected='Category' if name=='Products' and number==1 and index==3 else '='+formula(cell.value[1:],name=='Products') if cell.data_type=='f' else cell.value
     other=after[column-1]
     if other.value!=expected or getattr(other,'_style_id',0)!=getattr(cell,'_style_id',0):raise ValueError('Legacy template preservation verification failed')
 finally:old.close();new.close()
 return out,True
