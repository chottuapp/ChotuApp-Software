"""804 Audit: independent all-row description/category and store identity gate.
Bounded automatic checks, not a claim of human semantic review. Uncertain files HOLD.
"""
import hashlib,html,json,re
from datetime import datetime,timezone
from pathlib import Path
from urllib.parse import urlsplit
from openpyxl import load_workbook
VERSION='804-audit-v2'
# Independent reviewer vocabulary; do not call the classification implementation.
TYPES=[
 ('laptop',r'\b(?:laptops?|notebook|macbook|chromebook|ideapad)\b',r'\blaptops?\b'),
 ('printer toner',r'\b(?:printer\s+toner|toner\s+(?:reload|cartridge))\b',r'\b(?:printer|toner|cartridge)\b'),
 ('keyboard',r'\bkeyboards?\b',r'\bkeyboards?\b'),
 ('earrings',r'\b(?:earrings?|hoops?|studs)\b',r'\b(?:earrings?|hoops?|studs)\b'),
 ('necklace',r'\bnecklaces?\b',r'\bnecklaces?\b'),
 ('body wash',r'\b(?:body\s+wash|shower\s+gel)\b',r'\b(?:body\s+wash|shower\s+gel)\b'),
 ('hand wash',r'\bhand\s+wash\b',r'\bhand\s+wash\b'),
 ('face wash',r'\b(?:face|facial)\s+wash\b',r'\b(?:face|facial)\s+(?:wash|cleanser)\b'),
 ('hair shampoo',r'\bshampoo\b',r'\bshampoo\b'),
 ('hair conditioner',r'\b(?:hair\s+)?conditioner\b',r'\bconditioner\b'),
 ('hair oil',r'\bhair\s+oil\b',r'\bhair\s+oil\b'),
 ('fragrance',r'\b(?:perfume|eau\s+de\s+parfum|fragrance\s+mist)\b',r'\b(?:fragrance|perfume)\b'),
 ('sunscreen',r'\b(?:sunscreen|sunblock)\b',r'\b(?:sunscreen|sunblock|sun\s+protection)\b'),
 ('lipstick',r'\b(?:lipstick|lip\s+paint)\b',r'\b(?:lipstick|lip\s+colour|lip\s+color)\b'),
 ('mascara',r'\bmascara\b',r'\bmascara\b'),
 ('foundation',r'\bfoundation\b',r'\bfoundation\b'),
 ('lip gloss',r'\blip\s*gloss\b',r'\blip\s*gloss\b'),
 ('lip liner',r'\blip\s+liners?\b',r'\blip\s+liners?\b'),
 ('lip balm',r'\blip\s+balms?\b',r'\blip\s+balms?\b'),
 ('eyeshadow',r'\beye\s*shadow\b',r'\beye\s*shadow\b'),
 ('eyeliner',r'\b(?:eye\s*liner|kajal|kohl)\b',r'\b(?:eye\s*liner|kajal|kohl)\b'),
 ('concealer',r'\bconcealers?\b',r'\bconcealers?\b'),
 ('makeup brush',r'\bmakeup\s+brush(?:es)?\b',r'\bmakeup\s+brush(?:es)?\b'),
 ('makeup sponge',r'\b(?:makeup|beauty)\s+sponges?\b',r'\b(?:makeup|beauty)\s+sponges?\b'),
 ('hair serum',r'\bhair\s+(?:growth\s+)?serum\b',r'\bhair\s+serum\b'),
 ('face serum',r'\b(?:facial|face)\s+serum\b',r'\b(?:facial|face)\s+serum\b'),
 ('body lotion',r'\bbody\s+(?:lotion|milk)\b',r'\bbody\s+(?:lotion|milk)\b'),
 ('face toner',r'\b(?:face|facial)\s+toner\b',r'\b(?:face|facial)\s+toner\b'),
 ('garam masala',r'\bgaram\s+masala\b',r'\bgaram\s+masala\b'),
 ('tea bags',r'\btea\s+bags?\b',r'\btea\s+bags?\b'),
 ('pizza cheese',r'\bpizza\s+cheese\b',r'\bpizza\s+cheese\b'),
 ('loose leaf tea',r'\btea\s+leaf\s+blend\b|\bloose[ -]leaf\s+tea\b',r'\bloose[ -]leaf\s+tea\b'),
 ('blush',r'\bblush(?:er)?s?\b',r'\bblush(?:er)?s?\b'),
 ('mozzarella',r'\bmozzarella\b',r'\bmozzarella\b'),
 ('unstitched',r'\bun[ -]?stitched\b',r'\bun[ -]?stitched\b'),
 ('jeans',r'\bjeans?\b|\bdenim\s+(?:pants?|trousers?)\b',r'\b(?:jeans?|denim|trousers?)\b'),
 ('hoodie',r'\bhood(?:ie|ies)\b',r'\bhood(?:ie|ies)\b'),
 ('kurta',r'\bkurt[ai]s?\b',r'\b(?:kurt[ai]s?|ready.to.wear|suits?|outfits?)\b'),
 ('trousers',r'\b(?:trousers?|pants?|culottes?|shalwars?)\b',r'\b(?:trousers?|pants?|culottes?|shalwars?|joggers?)\b'),
]
# Independent evidence checks for reviewed primary objects and equivalent labels.
REVIEW_TYPES=[
 ('tea leaves',r'\b(?:curled|tea)\s+leaves\b|\b(?:leaf blend|super dana)\b',r'loose.leaf tea'),
 ('mouse pad',r'\bmouse\s*pad\b',r'mouse pads?'),
 ('usb hub',r'\b(?:usb hub|multiport hub|docking station)\b',r'usb hubs?.*docks?'),
 ('laundry cleansing wipes',r'\b(?:clothes|fabric)\b.*\bwipes\b',r'laundry stain.removal wipes'),

 ('lassi',r'\blassi\b',r'lassi'),
 ('ssd enclosure',r'\b(?:ssd|hdd|drive) enclosure\b',r'ssd enclosures?'),
 ('external storage',r'\b(?:external|portable) (?:hard|ssd|hdd|drive)\b',r'external storage drives?'),
 ('usb flash drive',r'\b(?:flash drive|pen drive)\b',r'usb flash drives?'),
 ('phone stand',r'\b(?:phone|mobile) (?:stand|holder)\b',r'phone holders?|mounts?'),
 ('power supply',r'\bpower supply\b',r'power supplies'),
 ('computer case',r'\b(?:tower|pc|atx|computer) (?:case|chassis)\b',r'computer cases?'),
 ('ram',r'\b(?:ddr[345]|ram|memory module)\b',r'computer memory|ddr[345] memory'),
 ('wet wipes',r'\b(?:travelling|travel|wet|refreshing|fresh) wipes\b',r'wet wipes'),
 ('makeup removal wipe',r'\b(?:cleansing|make ?up remov\w*)\b.*\bwipes\b',r'makeup remover wipes'),
 ('laundry wipe',r'\b(?:clothes|fabric)\b.*\bwipes\b',r'laundry stain.removal wipes'),
 ('pet mist',r'\bmist\b.*\b(?:dog|cat|pet)\b',r'pet deodori[sz]ers.*mists'),
 ('toilet paper',r'\btoilet (?:roll|paper)\b',r'toilet paper'),
 ('kitchen towel',r'\bkitchen (?:towel|roll)\b',r'kitchen paper towels?'),
 ('paper roll',r'\bmaxi roll\b',r'paper towels?.*rolls?'),
 ('body lotion',r'\b(?:body|fragrance) lotion\b',r'body lotions?'),

 ('processor cooler',r'\b(?:cpu|processor|liquid)\s+cooler\b',r'cpu coolers'),
 ('monitor arm',r'\bmonitor\s+(?:arm|mount|stand)\b',r'monitor arms?'),
 ('stylus',r'\bstylus\b',r'stylus pens?'),
 ('webcam',r'\bweb\s*cam\b',r'webcams?'),
 ('game controller',r'\b(?:controller|gamepad|dualsense)\b',r'game controllers?'),
 ('computer mouse',r'\bmouse\b',r'computer mice'),
 ('headset',r'\b(?:headsets?|headphones?)\b',r'headsets?|headphones?'),
 ('earbuds',r'\b(?:earbuds?|galaxy buds)\b',r'earphones?|earbuds?'),
 ('backpack',r'\b(?:backpack|laptop bag|top load laptop bag)\b',r'backpacks?|laptop bags?'),
 ('keyboard combo',r'\b(?:desktop set|wireless combo|bluetooth wireless combo|keys s combo)\b',r'keyboard.*mouse.*sets?'),
 ('solid state drive',r'\b(?:ssd|solid.state drive)\b',r'solid.state drives?'),
 ('power bank',r'\bpower bank\b',r'power banks?'),
 ('wall charger',r'\b(?:wall charger|gan.*charger)\b',r'wall chargers?'),
 ('usb cable',r'\b(?:coiled|usb|type.c|hdmi|displayport|ethernet)\s+cable\b',r'cables?'),
 ('adapter',r'\b(?:adapter|converter)\b',r'adapters?|converters?|hubs?'),
 ('yoghurt',r'\byogh?urt\b',r'yogh?urt'),
 ('fabric detergent',r'\b(?:white|fabric|abaya)\s+(?:cleansing\s+)?shampoo\b',r'fabric detergents?'),
 ('baby lotion',r'\bbaby(?: milk)? (?:lotion|cream)\b',r'baby lotions?|creams?'),
 ('baby wash',r'\bbaby (?:washing gel|bath|body wash)\b',r'baby bath|body wash'),
 ('sanitiser',r'\bhand saniti[sz]er\b',r'hand saniti[sz]ers?'),
 ('tissues',r'\b(?:mini|compact|facial)\s+(?:mini\s+)?tissues?\b',r'facial tissues?'),
 ('body powder',r'\b(?:prickly heat|body) powder\b',r'body powders?'),
 ('lip oil',r'\blip oil\b',r'lip oils?'),
 ('lash serum',r'\b(?:eyelash|lash) serum\b',r'eyelash serums?'),
 ('false eyelashes',r'\b(?:mink lashes|false lashes|false eyelashes)\b',r'false eyelashes'),
 ('face powder',r'\b(?:loose|compact|setting) powder\b',r'face powders?'),
 ('bb cream',r'\b(?:bb|cc) cream\b',r'bb.*cc creams?'),
 ('contour',r'\bcontour(?:ing)?\b',r'contour|bronzer'),
 ('highlighter',r'\bhighlighter\b',r'highlighters?'),
 ('makeup fixer',r'\b(?:makeup fixer|make up fixer|setting spray)\b',r'makeup setting sprays?'),
 ('makeup brush',r'\b(?:face|makeup|paint) brush\b',r'makeup brushes?'),
 ('multipurpose pencil',r'\blip liner (?:plus|and|&) eyebrow\b',r'multipurpose makeup pencils?'),
 ('dual lip colour',r'\bdual.ended\b.*lipstick.*lip gloss',r'lip colour sets?'),
 ('lip cheek colour',r'\blip\s*(?:&|and)\s*cheek\b',r'lip.*cheek colour'),
 ('eye palette',r'\b(?:eye|eyeshadow) palette\b',r'eyeshadow palettes?'),
 ('press on nails',r'\b(?:press.on|stick.on) nails\b',r'press.on nails'),
]
TYPES=REVIEW_TYPES+TYPES
TYPES=[(name,re.compile(a,re.I),re.compile(b,re.I)) for name,a,b in TYPES]
CARE=re.compile(r'\b(?:(?:jewellery|jewelry)\s+care|care instructions|how to use|directions|usage|precautions|avoid contact|apply to|wash with)\b',re.I)
def plain(value):return re.sub(r'\s+',' ',html.unescape(re.sub('<[^>]*>',' ',str(value or '')))).strip()
def tokens(value):return set(re.findall(r'\d+(?:[.,]\d+)?',plain(value)))
def issues_for(row,source=None):
 name=plain(row.get('name'));description=plain(row.get('description'));category=plain(row.get('Category'));issues=[]
 if not name:issues.append('Missing product name')
 if not description or len(re.findall(r'\b\w+\b',description))<5 or description.rstrip('.')==name.rstrip('.'):
  issues.append('Description lacks enough product-specific content for review')
 if re.search(r'\b(?:lorem ipsum|description coming soon|no description available)\b',description,re.I):issues.append('Placeholder description')
 if not category:issues.append('Category blank or unresolved')
 if category.lower() in {'accessories','clothing','beauty','skin care','hair care','food','sale','new arrivals','collection','women','men','kids','general','products'}:issues.append('Category is too broad or promotional for product-purpose classification')
 if re.search(r'Product Clarification:|description refers to .*while this listing',description,re.I):issues.append('Conflicting source description requires attention')
 from reviewed804_repairs import product_content
 if source and product_content(source).strip() not in ('0','None','') and not tokens(product_content(source))<=tokens(description):issues.append('Description lost source numeric facts')
 body=CARE.split(description,1)[0][:1600]
 category_review=category
 for plural,singular in {'glosses':'gloss','washes':'wash','shampoos':'shampoo','conditioners':'conditioner','mascaras':'mascara','foundations':'foundation','serums':'serum','toners':'toner','lotions':'lotion','concealers':'concealer','eyeliners':'eyeliner','eyeshadows':'eyeshadow','lipsticks':'lipstick','perfumes':'perfume','fragrances':'fragrance','sunscreens':'sunscreen'}.items():category_review=re.sub(r'\b'+plural+r'\b',singular,category_review,flags=re.I)
 primary=name.split(',')[0] or body
 promotion=re.match(r'(?is)^buy\s+(.+?)\s*(?:&|and)\s*get\s+.+?\s+free$',primary)
 if promotion:primary=promotion.group(1)

 found=[(kind,allowed) for kind,pattern,allowed in TYPES if pattern.search(primary)]
 # Electronics specs mention other components; the earliest named object wins.
 electronic={'processor cooler','monitor arm','stylus','webcam','game controller','computer mouse','headset','earbuds','backpack','keyboard combo','solid state drive','power bank','wall charger','usb cable','adapter','ssd enclosure','external storage','usb flash drive','phone stand','power supply','computer case','ram','mouse pad','usb hub','laptop','keyboard','printer toner'}
 candidates=[(pattern.search(primary).start(),index,kind,allowed) for index,(kind,pattern,allowed) in enumerate(TYPES) if kind in electronic and pattern.search(primary)]
 if candidates:
  _,_,kind,allowed=min(candidates);found=[(kind,allowed)]

 if not found:found=[(kind,allowed) for kind,pattern,allowed in TYPES if pattern.search(body[:260])]
 bundle=bool(re.search(r'\b(?:bundle|combo|kit|deal|set of)\b',name,re.I))
 if re.search(r'printer|toner|reload',name,re.I):bundle=False
 from audit804_semantics import review
 review_description=str(row.get('description') or '')
 if category=='Unstitched Suits' and row.get('website_category_status')=='verified_product_breadcrumb' and re.search(r'(?i)\bunstitched\b',str(row.get('website_category_path') or '')) and re.search(r'(?i)\bshirt\s*&\s*trouser\s+\d+(?:\.\d+)?\s*m\b',review_description):
  review_description='Unstitched suit. '+review_description
 verified_path=''
 if row.get('website_category_status') in ('verified_navigation_membership','verified_product_breadcrumb') and row.get('website_category_checked_at') and row.get('website_category_source_url'):
  if (urlsplit(str(row['website_category_source_url'])).hostname or '').removeprefix('www.')==(urlsplit(str(row.get('source_url') or '')).hostname or '').removeprefix('www.'):
   verified_path=str(row.get('website_category_path') or '')
 semantic=review(name,review_description,category,str(row.get("sku") or ""),source_type=str(row.get("brand") or ""),verified_path=verified_path)
 if category and plain(row.get('brand')) and category.casefold()==plain(row.get('brand')).casefold() and semantic!=[]:issues.append('Brand name used as product Category')
 if semantic is not None:issues.extend(semantic)
 else:
  if bundle and category:
   if not re.search(r'\b(?:bundle|combo|kit|set|pack|outfit|suit)s?\b',category,re.I):issues.append('Bundle appears classified as a single component')
   families={'dairy':r'cheese|milk|butter|ghee|dairy|yogh?urt|raita|lassi','makeup':r'lipstick|lip\s*gloss|eye\s*shadow|mascara|foundation|concealer|makeup','hair':r'hair|shampoo|conditioner','skin':r'skin|face|facial|serum|cream|lotion','spice':r'spice|masala|cumin|coriander|turmeric'}
   for family,pattern in families.items():
    if re.search(r'\b'+family+r'\b',category,re.I) and not re.search(pattern,name+' '+body,re.I):issues.append('Bundle family and product contents disagree')
   if category.lower().startswith('hair') and found and not found[0][0].startswith('hair'):issues.append('Hair bundle category conflicts with primary product purpose')
   if re.search(r'\b(?:skin|face)',category,re.I) and found and found[0][0] in ('lipstick','lip gloss','lip liner','eyeshadow','eyeliner','mascara','foundation','concealer'):issues.append('Cosmetics bundle incorrectly classified as skincare')
  elif found and category:
   if not found[0][1].search(category_review):issues.append('Primary product type and Category disagree: '+found[0][0])
  elif category:
   anchors=[w if w.endswith('ss') else w[:-2] if w.endswith(('ches','shes','xes','sses','zes')) else w[:-1] if w.endswith('s') else w for w in re.findall('[a-z]{3,}',category.lower()) if w not in {'and','for','with','care','sets','set'}]
   evidence=(name+' '+body).lower()
   if not anchors or not all(re.search(r'\b'+re.escape(w)+r'(?:s|es)?\b',evidence) for w in anchors):issues.append('Category lacks clear product-purpose evidence; manual review required')
 if semantic is None and re.search(r'\b[23]\s*[- ]?(?:pc|piece)',name,re.I) and re.search(r'\b(?:kurtas?|shirts?|trousers?|dupattas?)\b',category,re.I) and not re.search(r'suit|set|outfit|ready',category,re.I):issues.append('Multi-piece outfit may be classified as one component')
 return list(dict.fromkeys(issues))

def rows(path):
 book=load_workbook(path,read_only=True,data_only=False)
 try:
  sheet=book['Products'];iterator=sheet.iter_rows(values_only=True);headers=next(iterator)
  for values in iterator:
   if any(value is not None for value in values):yield dict(zip(headers,values))
 finally:book.close()

def run(path,store,codes,hosts,*,source=None):
 path=Path(path);findings=[];expected=codes.get(store);seen=set();count=0;finding_count=0
 originals={str(row.get('product_id')):row for row in rows(source)} if source else {}
 from reviewed804_repairs import roots,conflicting_name
 catalogue=roots(originals.values()) if originals else set()
 for row in rows(path):
  count+=1;pid=str(row.get('product_id') or '');issues=issues_for(row,originals.get(pid))
  if store=='Seran' and conflicting_name(row,catalogue):issues.append('Description names a different known catalog product')
  if not pid or pid in seen:issues.append('Missing or duplicate product ID')
  seen.add(pid)
  if expected and str(row.get('shop_code') or '')!=expected:issues.append('Shop code missing, mixed or not confirmed for this store')
  if not expected and str(row.get('shop_code') or '').strip():issues.append('Unconfirmed shop code must remain blank')
  if store=='LAAM' and not re.search(r'cloth|suit|unstitched|stitched|wear|kurta|kurti|dress|trouser|shirt|jean|hood|jacket|shawl|dupatta|abaya|kaftan|saree|lehenga|sweater|cardigan|jogger|short|polo|legging|coat|outfit|gown|skirt|shalwar|pant|frock|top|sweatshirt',str(row.get('Category') or ''),re.I):issues.append('Rule LAAM: clothing purpose not confirmed')
  url=urlsplit(str(row.get('source_url') or ''))
  if url.scheme!='https' or (url.hostname or '').lower().removeprefix('www.') not in hosts:issues.append('Store source host mismatch')
  if plain(row.get('Category')).casefold()==store.casefold():issues.append('Store name used as product Category')
  if issues:
   finding_count+=1
   if len(findings)<100:findings.append({'product_id':pid,'issues':issues})
 if expected and list(codes.values()).count(expected)>1:finding_count+=1;findings.append({'product_id':None,'issues':['Shop code assigned to multiple stores; identity review required']})
 if not count:finding_count+=1;findings.append({'product_id':None,'issues':['Empty workbook']})
 if originals and set(originals)!=seen:finding_count+=1;findings.append({'product_id':None,'issues':['Source/output product identity mismatch']})
 return {'name':'804 Audit','version':VERSION,'store':store,'shop_code':expected,'identity_policy':'confirmed_code' if expected else 'owner_authorized_blank_code_exact_store_filename','status':'pass' if not finding_count else 'hold','products_checked':count,'findings_count':finding_count,'findings':findings,'details_truncated':finding_count>len(findings),'workbook_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'checked_utc':datetime.now(timezone.utc).isoformat(),'review_scope':'Every product description, independent product-type/category checks, source facts and confirmed shop-code identity. Ambiguity is held; this is not human semantic review or proof of full website coverage.'}
