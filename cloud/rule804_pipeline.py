"""Automated Rule 804 preservation/classification checks for the fresh 26-store batch."""
import csv,hashlib,importlib.util,json,re,shutil,sys
from datetime import datetime,timezone
from pathlib import Path
from decimal import Decimal,InvalidOperation
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from subcategory_inference_support import infer
EXCLUDED={'LAAM','Khaadi','Sapphire'}
def toolkit():
    spec=importlib.util.spec_from_file_location('rule804_xml_tools',ROOT/'catalog-tools/complete_rule804_batch.py');module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module);return module

def process(source,out,store,codes):
    if store in EXCLUDED:raise ValueError('Excluded store cannot enter fresh Rule 804 pipeline')
    tools=toolkit();book=load_workbook(source,read_only=True,data_only=False);patches={};warnings=[];ids=set();count=0;stamp=datetime.now(timezone.utc).isoformat()
    try:
        if not {'Products','Product Images','Option Groups','Options','Variant Combinations'}<=set(book.sheetnames):raise ValueError('Required template sheets missing')
        sheet=book['Products'];rows=sheet.iter_rows(values_only=True);headers=next(rows);indices={h:i+1 for i,h in enumerate(headers)}
        if list(headers[:3])!=['shop_code','product_id','Category']:raise ValueError('Locked Rule 804 schema mismatch')
        required={'name','description','base_price (PKR)','source_url','subcategory_source','subcategory_evidence','subcategory_checked_at'}
        if not required<=set(headers):raise ValueError('Required description/category evidence fields missing')
        for number,values in enumerate(rows,2):
            if not any(v is not None for v in values):continue
            row=dict(zip(headers,values));pid=row['product_id'];count+=1
            if not pid or str(pid) in ids:raise ValueError('Missing or duplicate product ID')
            ids.add(str(pid))
            try:price=Decimal(str(row['base_price (PKR)']))
            except (InvalidOperation,TypeError):raise ValueError('Invalid positive product price')
            if not price.is_finite() or price<=0:raise ValueError('Invalid positive product price')
            description=str(row.get('description') or '')
            if len(description)>=32767:raise ValueError('Description reaches Excel limit; manual source review needed')
            changes={};code=codes.get(store)
            if code and row.get('shop_code')!=code:changes['shop_code']=code
            if not code and row.get('shop_code'):raise ValueError('Unconfirmed shop code; held for review')
            old=str(row.get('Category') or '');label=old
            if str(row.get('subcategory_source') or '').lower() not in ('manual','owner','manual_override'):
                result=infer(row.get('name'),description)
                if result:
                    label,origin,quote=result
                    if label!=old:changes.update(Category=label,subcategory_source=origin,subcategory_evidence=quote,subcategory_checked_at=stamp)
            if not label:warnings.append({'product_id':pid,'issue':'Ambiguous product category retained blank'})
            polished=tools.tidy_description(description)
            if re.findall(r'\d+(?:[.,]\d+)?',description)!=re.findall(r'\d+(?:[.,]\d+)?',polished):raise ValueError('Description numeric facts changed')
            overview=tools.product_overview(row,label)
            lead=str(row.get('name') or '')+'.'
            if overview and polished.startswith(lead):polished=overview+polished[len(lead):]
            if not description.strip():warnings.append({'product_id':pid,'issue':'Missing description; source review required'})
            if polished!=description:changes['description']=polished
            for key,value in changes.items():patches[get_column_letter(indices[key])+str(number)]=value
        if not count:raise ValueError('Empty catalog')
        for name in ('Product Images','Option Groups','Options','Variant Combinations'):
            iterator=book[name].iter_rows(values_only=True);labels=next(iterator)
            if 'product_id' not in labels:continue
            col=labels.index('product_id')
            for values in iterator:
                if values[col] is not None and str(values[col]) not in ids:raise ValueError('Orphan product reference in '+name)
    finally:book.close()
    out.parent.mkdir(parents=True,exist_ok=True);temporary=out.with_name(out.stem+'.pending.xlsx')
    if patches:
        part=tools.patch_sheet(source,temporary,patches);tools.verify(source,temporary,part,patches)
    else:shutil.copy2(source,temporary)
    temporary.replace(out)
    return {'store':store,'name':out.name,'products':count,'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'bytes':out.stat().st_size,'patched_cells':len(patches),'warnings':warnings,'rule804':'automated evidence-bound pass; unresolved evidence retained for review','preservation':'Unaffected XLSX parts and Products cells verified unchanged; formulas, prices, stock, variants, image links and source order preserved'}

def run(source_dir,out_dir):
    source_dir=Path(source_dir);out_dir=Path(out_dir);out_dir.mkdir(parents=True,exist_ok=True)
    codes=json.loads((ROOT/'cloud/confirmed-shop-codes.json').read_text());registry=json.loads((ROOT/'cloud/fetch-update-release.json').read_text())['fresh_stores'];allowed={s['name'] for s in registry}-EXCLUDED
    manifest_path=out_dir/'manifest.json'
    try:previous={x['name']:x for x in json.loads(manifest_path.read_text())['files']}
    except (OSError,ValueError,KeyError):previous={}
    files=[];held=[]
    revision=hashlib.sha256(Path(__file__).read_bytes()+(ROOT/'subcategory_inference_support.py').read_bytes()+(ROOT/'catalog-tools/complete_rule804_batch.py').read_bytes()+(ROOT/'cloud/confirmed-shop-codes.json').read_bytes()).hexdigest()
    for source in sorted(source_dir.glob('*.xlsx')):
        if source.stem not in allowed:continue
        sha=hashlib.sha256(source.read_bytes()).hexdigest();old=previous.get(source.name);out=out_dir/source.name
        if old and old.get('source_sha256')==sha and old.get('pipeline_revision')==revision and out.is_file() and hashlib.sha256(out.read_bytes()).hexdigest()==old['sha256']:files.append(old);continue
        try:record=process(source,out,source.stem,codes);record['pipeline_revision']=revision;files.append(record)
        except Exception as exc:
            out.unlink(missing_ok=True);held.append({'store':source.stem,'error':str(exc)});print('RULE804 HELD:',source.stem,str(exc),flush=True)
    # Manifest is the upload allowlist; reports remain in GitHub artifacts, never Drive.
    manifest={'files':files,'held':held,'checked_utc':datetime.now(timezone.utc).isoformat()};tmp=out_dir/'manifest.pending';tmp.write_text(json.dumps(manifest,indent=2),encoding='utf-8');tmp.replace(manifest_path)
    print('RULE804 AUTOMATED:',len(files),'verified outputs;',len(held),'held',flush=True);return manifest
if __name__=='__main__':run(sys.argv[1],sys.argv[2])
