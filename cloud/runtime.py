"""Headless adapter around the existing desktop extraction/checkpoint engine."""
import collections,hashlib,ipaddress,json,logging,logging.handlers,os,re,shutil,socket,threading,time
from pathlib import Path
from urllib.parse import urlsplit
import sys
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from resume_support import RunLease,RunStore
from parallel_store_support import StoreScheduler,output_key
from diagnostics import Collector
from continuation_policy import integer

MODULES=('subcategory_inference_support.py','website_category_support.py','ranking_support.py','resume_support.py','fast_fetch_support.py','farda_source_support.py','complete_description_support.py','czone_source_support.py','software_revision_support.py','seran_source_support.py','laam_source_support.py','source_compatibility_support.py','website_order_support.py')
DEFAULT_TEMPLATE=ROOT/'importer-files/Chotu_Products_Import_Template_Discount_Tags.xlsx'


def prioritize_stores(sites):
    """Stable queue by prior observed count, never invent counts for new stores."""
    counts=json.loads((ROOT/'cloud/store-priority.json').read_text())['product_counts']
    return sorted(sites,key=lambda site:counts.get(site['name'],float('inf')))


def prepare_engine(data):
    """Generate a separate runtime; never patch the Windows application in place."""
    import install_fast_fetch,install_farda_source,install_czone_source,install_speed_patch
    runtime=Path(data)/'runtime';runtime.mkdir(parents=True,exist_ok=True)
    source=install_fast_fetch.transform((ROOT/'main.py').read_text(encoding='utf-8-sig'))
    source=install_farda_source.transform(source);source=install_czone_source.transform(source)
    hook='\n# CHOTUAPP_SERAN_PUBLIC_SOURCE_V1\nfrom seran_source_support import enable_seran_source\nenable_seran_source(globals())\nfrom laam_source_support import enable_laam_source\nenable_laam_source(globals())\nfrom source_compatibility_support import enable_source_compatibility\nenable_source_compatibility(globals())\nfrom website_order_support import enable_website_order\nenable_website_order(globals())\n\n'
    guard="if __name__=='__main__':"
    if source.count(guard)!=1:raise ValueError('Unexpected engine startup entry point')
    source=source.replace(guard,hook+guard,1)
    content,_=install_speed_patch.transform(source.encode());compile(content,str(runtime/'main.py'),'exec')
    for name in MODULES:
        target=runtime/name;payload=(ROOT/name).read_bytes()
        if not target.exists() or target.read_bytes()!=payload:target.write_bytes(payload)
    if not (runtime/'main.py').exists() or (runtime/'main.py').read_bytes()!=content:(runtime/'main.py').write_bytes(content)
    return runtime/'main.py'


def validate_site(site,resolve=True):
    if not isinstance(site,dict):raise ValueError('Each store needs a name and URL')
    name=str(site.get('name','')).strip();url=str(site.get('url','')).strip();parsed=urlsplit(url)
    if not name or len(name)>100 or not name.isprintable() or output_key({'name':name})!=name.casefold():raise ValueError('Use a store name with letters, numbers, spaces, dots or hyphens')
    if parsed.scheme not in ('http','https') or not parsed.hostname or parsed.username or parsed.password or parsed.query or parsed.fragment or parsed.path not in ('','/'):
        raise ValueError('Use the public store homepage URL, without tracking text or product paths')
    if parsed.port not in (None,80,443):raise ValueError('Only standard website ports are supported')
    host=parsed.hostname.lower()
    if host=='localhost' or host.endswith(('.local','.localhost','.internal')):raise ValueError('A public store hostname is required')
    if resolve:
        try:addresses={item[4][0] for item in socket.getaddrinfo(host,parsed.port or (443 if parsed.scheme=='https' else 80),type=socket.SOCK_STREAM)}
        except socket.gaierror as exc:raise ValueError('Store hostname could not be resolved') from exc
        if not addresses or any(not ipaddress.ip_address(x).is_global for x in addresses):raise ValueError('Private or local store addresses are not supported')
    return {'name':name,'url':parsed.scheme+'://'+parsed.netloc.rstrip('/')}


class CloudService:
    def __init__(self,data,engine=None,resolve_sites=True):
        self.data=Path(data).resolve();self.data.mkdir(parents=True,exist_ok=True)
        self.store=RunStore(self.data/'saved_runs');self.lease=RunLease(self.store.root);self.lease.acquire()
        try:self.engine=Path(engine) if engine else prepare_engine(self.data)
        except Exception:self.lease.release();raise
        self.output=self.data/'ANAS_DATA_FETCH';self.output.mkdir(exist_ok=True)
        self.pause=threading.Event();self.pause.set();self.lock=threading.RLock();self.thread=None;self.scheduler=None;self.resolve_sites=resolve_sites;self.error=None
        self.diagnostics=Collector(self.data/'telemetry.json')
        self.lines=collections.deque(maxlen=500);self.counter=0
        self.logger=logging.Logger('chotu-cloud-'+str(id(self)))
        handler=logging.handlers.RotatingFileHandler(self.data/'fetch.log',maxBytes=2*1024*1024,backupCount=2,encoding='utf-8');self.logger.addHandler(handler)
        self.logger.addHandler(logging.FileHandler(self.data/'fetch-full.log',encoding='utf-8'))
        self.settings=self.data/'controls.json'
        try:
            if json.loads(self.settings.read_text()).get('paused'):self.pause.clear()
        except (OSError,ValueError):pass
        # After a crash, absolute snapshot/output paths belong to this persistent volume.
        def relocate(state):
            folder=self.store.root/state['id'];state['template']=str(folder/'template.xlsx')
            for site in state['sites']:
                if site.get('output'):
                    file=self.output/Path(site['output']['path']).name
                    if file.is_file():
                        site['output']=dict(site['output'],path=str(file))
        self.store.update(relocate)
    @property
    def running(self):return self.thread is not None and self.thread.is_alive()
    def event(self,item):
        text=item[1] if isinstance(item,tuple) and len(item)>1 else str(item)
        with self.lock:
            self.counter+=1;self.lines.append({'id':self.counter,'text':str(text)});self.logger.info(str(text));self.diagnostics.observe(str(text))
    def start(self,sites=None,resume=False):
        from openpyxl import load_workbook
        with self.lock:
            if self.running:raise ValueError('A run is already active. Pause or stop it before starting another.')
            if resume:
                state=self.store.load()
                if not state or not self.store.pending(state):raise ValueError('No unfinished saved run')
                state=dict(state,_resuming=True)
            else:
                if not isinstance(sites,list) or not 1<=len(sites)<=30:raise ValueError('Select 1 to 30 stores')
                chosen=[validate_site(x,self.resolve_sites) for x in sites]
                if len({output_key(x) for x in chosen})!=len(chosen):raise ValueError('Every store needs a distinct name')
                # Historical exported counts prioritize smaller known catalogs; unknowns remain stable last.
                chosen=prioritize_stores(chosen)
                self.event('QUEUE: smaller historical catalogs first; unknown counts last; fresh source data fetched again')
                template=self.data/'template.xlsx'
                if not template.exists():shutil.copy2(DEFAULT_TEMPLATE,template)
                book=load_workbook(template,read_only=True)
                try:
                    if 'Products' not in book.sheetnames:raise ValueError('Template must contain Products')
                finally:book.close()
                state=self.store.fresh(chosen,template)
            if not Path(state['template']).is_file():raise ValueError('Saved template is missing')
            # Add the new field to a resumed template without replacing its custom columns.
            book=load_workbook(state['template'])
            try:
                sheet=book['Products'];field='stock_quantity (kitni quantity bechni hai)'
                from ranking_support import ensure_columns
                changed=ensure_columns(sheet)
                if field not in [cell.value for cell in sheet[1]]:
                    from copy import copy
                    previous=sheet.cell(1,sheet.max_column)
                    cell=sheet.cell(1,sheet.max_column+1,field);cell._style=copy(previous._style)
                    changed=True
                if changed:book.save(state['template'])
            finally:book.close()
            self.diagnostics=Collector(self.data/'telemetry.json')
            self.pause.set();self.settings.write_text(json.dumps({'paused':False}));self.error=None
            scheduler=StoreScheduler(self.engine,self.store,state,self.output,self.pause,self.event,retry_limit=5,wait_limit=900,skip_terminal=integer(os.environ.get('CLOUD_CONTINUATION_INDEX','0'))>0);self.scheduler=scheduler
            def work():
                try:scheduler.run()
                except Exception as exc:self.error=str(exc);self.event('FAILED: '+str(exc))
            self.thread=threading.Thread(target=work,name='cloud-store-scheduler',daemon=True);self.thread.start()
    def control(self,command):
        with self.lock:
            if command=='pause':self.pause.clear()
            elif command=='resume':
                if not self.running:raise ValueError('Use Continue Saved Run to restart stopped work')
                self.pause.set()
            elif command=='stop':
                if self.scheduler:self.scheduler.shutdown()
            else:raise ValueError('Unknown control')
            self.settings.write_text(json.dumps({'paused':not self.pause.is_set() or command=='stop'}))
    def snapshot(self):
        with self.lock:
            state=self.store.load();sites=[]
            for site in (state or {}).get('sites',[]):sites.append({k:site.get(k) for k in ('name','url','status','progress','detail','reviewed','software_revision')})
            files=[{'name':p.name,'bytes':p.stat().st_size,'modified':p.stat().st_mtime} for p in self.output.iterdir() if p.is_file() and p.suffix.lower() in ('.xlsx','.csv') and not p.name.startswith(('.','~')) and 'BUILDING' not in p.name]
            return {'running':self.running,'paused':not self.pause.is_set(),'error':self.error,'run_id':(state or {}).get('id'),'status':(state or {}).get('status','idle'),'can_continue':bool(state and self.store.pending(state)),'sites':sites,'files':sorted(files,key=lambda x:x['modified'],reverse=True),'logs':list(self.lines),'disk_free_bytes':shutil.disk_usage(self.data).free}
    def close(self):
        # Never hold the event lock while joining: the scheduler may still write a log.
        if self.scheduler:self.scheduler.shutdown()
        if self.thread:self.thread.join(timeout=20)
        if self.running:raise RuntimeError('Worker has not stopped; lease retained')
        self.diagnostics.flush()
        self.lease.release()
        for handler in self.logger.handlers:handler.close()
