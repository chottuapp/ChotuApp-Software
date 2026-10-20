"""Owner-approved solo original exports -> existing Rule 804 -> checked same-folder delivery."""
import base64,datetime,hashlib,json,tempfile
from pathlib import Path
from rule804_pipeline import process
ROOT=Path(__file__).resolve().parents[1]
REPO='chottuapp/ChotuApp-Software'

def candidates(runs,policy,gh):
 allowed={row['group']:row['store'] for row in policy};lanes={};result=[];seen=set()
 for run in reversed(runs):
  rid=str(run['databaseId']);title=run.get('displayTitle','')
  if title.startswith('Cloud group '):lanes[rid]=title.removeprefix('Cloud group ')
  elif title.startswith('Automatic continuation of '):lanes[rid]=lanes.get(title.split()[-1],'')
  if lanes.get(rid) in allowed:result.append((allowed[lanes[rid]],run));seen.add(rid)
 for row in policy:
  for rid in row.get('completed_run_ids',[]):
   if str(rid) in seen:continue
   run=json.loads(gh('run','view',str(rid),'--repo',REPO,'--json','databaseId,displayTitle,status,conclusion,headSha'))
   result.append((row['store'],run))
 # Only the newest completed export for a store may be delivered in each scan.
 return sorted(result,key=lambda pair:int(pair[1]['databaseId']),reverse=True)

def completed(stage,store):
 summaries=list(stage.rglob('run-summary.json'))
 if len(summaries)==1:
  data=json.loads(summaries[0].read_text())
  return any(row.get('name')==store and row.get('status')=='done' for row in data.get('stores',[]))
 return False

def upload(state,folder,rc,runs,gh,policy,*,root=ROOT):
 if not policy:return []
 revision=hashlib.sha256(b''.join((root/p).read_bytes() for p in ('cloud/rule804_pipeline.py','subcategory_inference_support.py','catalog-tools/complete_rule804_batch.py','cloud/confirmed-shop-codes.json','cloud/original_store_ingest.py','cloud/legacy_original_schema.py'))).hexdigest()
 codes=json.loads((root/'cloud/confirmed-shop-codes.json').read_text());sent=[];handled=set()
 for store,run in candidates(runs,policy,gh):
  rid=str(run['databaseId'])
  if store in handled:continue
  previous=state.get('original_uploads',{}).get(store,{})
  if str(previous.get('source_run','')).isdigit() and int(previous['source_run'])>int(rid):continue
  if run.get('status')!='completed':continue
  input_headers=[]
  try:
   artifacts=json.loads(gh('api','repos/'+REPO+'/actions/runs/'+rid+'/artifacts?per_page=100'))['artifacts']
   found=[a for a in artifacts if a['name']=='chotuapp-excel-files' and not a.get('expired')]
   if not found:continue
   artifact=max(found,key=lambda a:a['id']);identity=rid+':'+str(artifact['id'])+':'+revision+':'+folder
   if state.setdefault('original_processed',{}).get(store)==identity:handled.add(store);continue
   with tempfile.TemporaryDirectory() as tmp:
    stage=Path(tmp)/'input';stage.mkdir();gh('run','download',rid,'--repo',REPO,'--name','chotuapp-excel-files','--dir',str(stage))
    if not completed(stage,store):continue
    handled.add(store)
    files=list(stage.rglob(store+'.xlsx'))
    if len(files)!=1 or files[0].is_symlink():raise ValueError('Unique completed workbook required')
    if store=='LAAM':
     if not run.get('headSha'):
      run.update(json.loads(gh('run','view',rid,'--repo',REPO,'--json','headSha')))
     raw=json.loads(gh('api','repos/'+REPO+'/contents/laam_source_support.py?ref='+run['headSha']))
     actual=hashlib.sha256(base64.b64decode(raw['content'])).hexdigest()
     if actual!=hashlib.sha256((root/'laam_source_support.py').read_bytes()).hexdigest():raise ValueError('LAAM clothing-only source revision requires review')
    from openpyxl import load_workbook
    book=load_workbook(files[0],read_only=True,data_only=False)
    try:input_headers=list(next(book['Products'].iter_rows(values_only=True)))
    finally:book.close()
    output=Path(tmp)/'checked'/files[0].name
    from legacy_original_schema import normalize
    normalized,migrated=normalize(files[0],Path(tmp)/'normalized'/files[0].name,codes.get(store))
    record=process(normalized,output,store,codes,approved_original=True)
    if migrated:record.update(legacy_schema_migrated=True,normalized_source_sha256=record['source_sha256'],source_sha256=hashlib.sha256(files[0].read_bytes()).hexdigest())
    rc('copy',str(output.parent),'chotudrive:'+folder,'--include',output.name,'--checksum')
    rc('check',str(output.parent),'chotudrive:'+folder,'--include',output.name,'--one-way')
    record.update(source_run=rid,source_artifact=artifact['id'],pipeline_revision=revision,folder=folder,verified_utc=datetime.datetime.now(datetime.timezone.utc).isoformat(),coverage='Completed exported workbook; full website coverage remains unverified')
    state.setdefault('original_uploads',{})[store]=record;state['original_processed'][store]=identity
    state.setdefault('original_holds',{}).pop(store,None);sent.append(store)
  except (RuntimeError,ValueError,KeyError,OSError) as exc:
   state.setdefault('original_holds',{})[store]={'source_run':rid,'error':str(exc),'input_headers':input_headers,'observed_utc':datetime.datetime.now(datetime.timezone.utc).isoformat()}
   print('ORIGINAL RULE804 HELD:',store,str(exc),flush=True)
 return sent
