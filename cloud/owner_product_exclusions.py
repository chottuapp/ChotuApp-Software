"""Exact owner-authorized export exclusions; never a general audit bypass."""
import zipfile
import xml.etree.ElementTree as ET
from openpyxl import load_workbook

EXCLUSIONS={'Astore Bags': {'CPb7a9fb37ac6f':'https://astorebags.pk/products/macbook-sleeve-brown-13-inches'}, 'Elite Cosmetics': {'CPaeabf1fd66f3': 'https://elitecosmeticspk.com/products/comfor-beauty-aloe-vera-essence-oil-40ml'}}
NS='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'

def apply(source,out,store):
 rules=EXCLUSIONS.get(store,{})
 if not rules:return source,[]
 book=load_workbook(source,read_only=True,data_only=False)
 try:
  rows={s.title:[dict(zip(next(s.values),v)) for v in s.iter_rows(min_row=2,values_only=True)] for s in book}
  removed=set()
  for r in rows['Products']:
   pid=str(r.get('product_id') or '')
   if pid in rules:
    if r.get('source_url')!=rules[pid]:raise ValueError('Owner exclusion source identity mismatch')
    removed.add(pid)
  if not removed:return source,[]
  deleted={}
  used=set();candidates=set()
  for name,items in rows.items():
   if name not in ('Products','Variant Combinations'):continue
   for r in items:
    groups={str(v) for k,v in r.items() if v and (str(k).endswith('group_id') or str(k).endswith('group_ids'))}
    (candidates if str(r.get('product_id')) in removed else used).update(groups)
  exclusive=candidates-used
  for name,items in rows.items():
   deleted[name]={i for i,r in enumerate(items,2) if str(r.get('product_id')) in removed or name in ('Options','Option Groups') and str(r.get('group_id')) in exclusive}
  paths={s.title:s._worksheet_path.lstrip('/') for s in book}
 finally:book.close()
 out.parent.mkdir(parents=True,exist_ok=True)
 with zipfile.ZipFile(source) as src,zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as dst:
  bypath={paths[k]:v for k,v in deleted.items() if v}
  for info in src.infolist():
   data=src.read(info.filename)
   if info.filename in bypath:
    tree=ET.fromstring(data);sheet=tree.find(NS+'sheetData')
    for row in list(sheet):
     if int(row.attrib['r']) in bypath[info.filename]:sheet.remove(row)
    data=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
   dst.writestr(info,data)
 return out,[{'product_id':pid,'source_url':rules[pid],'reason':('Outside owner-locked Astore bags/wallets/mini-pouches scope: laptop sleeve' if store=='Astore Bags' else 'Owner explicitly requested deletion of conflicting product on 25 October 2026')} for pid in sorted(removed)]
