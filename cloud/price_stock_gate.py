"""Strict monetary and availability consistency; never fills unknown stock or prices."""
from decimal import Decimal,InvalidOperation
from collections import defaultdict
from openpyxl import load_workbook

def value(number):
 if number is None or str(number).strip()=='':return None
 if isinstance(number,bool):raise ValueError('Boolean is not a price/quantity')
 try:result=Decimal(str(number))
 except InvalidOperation:raise ValueError('Invalid numeric value')
 if not result.is_finite():raise ValueError('Non-finite numeric value')
 return result

def check(row):
 issues=[]
 try:
  base=value(row.get('base_price (PKR)'))
  if base is None or base<=0:issues.append('Price/stock gate: missing or nonpositive PKR base price')
  for key in ('sale_price (PKR, agar discount ho)','sale_price (PKR, only if discounted)'):
   sale=value(row.get(key))
   if sale is not None and (sale<=0 or base is None or sale>=base):issues.append('Price/stock gate: invalid sale price; must be positive and below base')
  sales=[value(row[k]) for k in ('sale_price (PKR, agar discount ho)','sale_price (PKR, only if discounted)') if row.get(k) not in (None,'')]
  if len(set(sales))>1:issues.append('Price/stock gate: conflicting sale-price columns')
 except ValueError as exc:issues.append('Price/stock gate: '+str(exc))
 available=str(row.get('available (yes/no)') or '').strip().lower()
 if available not in ('yes','no'):issues.append('Price/stock gate: availability is unknown')
 for key in ('stock_quantity (kitni quantity bechni hai)','stock_quantity (agar public ho)'):
  try:
   quantity=value(row.get(key))
   if quantity is not None:
    if quantity<0 or quantity!=quantity.to_integral_value():issues.append('Price/stock gate: invalid stock quantity')
    elif quantity==0 and available=='yes':issues.append('Price/stock gate: available product has zero stock')
  except ValueError as exc:issues.append('Price/stock gate: '+str(exc))
 return list(dict.fromkeys(issues))

def workbook_issues(path):
 book=load_workbook(path,read_only=True,data_only=False);issues=defaultdict(list);products={};variants=defaultdict(list)
 try:
  for title in ('Products','Variant Combinations'):
   if title not in book:continue
   iterator=book[title].iter_rows(values_only=True);headers=next(iterator)
   for values in iterator:
    if not any(v is not None for v in values):continue
    row=dict(zip(headers,values));pid=str(row.get('product_id') or '')
    issues[pid].extend(check(row))
    if title=='Products':
     products[pid]=row
     if str(row.get('available (yes/no)') or '').strip().lower()=='no':issues[pid].append('Stock policy: fully out-of-stock products must be excluded; retain unavailable variants only under an available product')
    else:variants[pid].append(row)
  for pid,rows in variants.items():
   if pid not in products:issues[pid].append('Price/stock gate: orphan variant')
   else:
    expected='yes' if any(str(r.get('available (yes/no)')).strip().lower()=='yes' for r in rows) else 'no'
    if str(products[pid].get('available (yes/no)') or '').strip().lower()!=expected:issues[pid].append('Price/stock gate: product and variant availability disagree')
 finally:book.close()
 return {pid:list(dict.fromkeys(found)) for pid,found in issues.items() if found}
