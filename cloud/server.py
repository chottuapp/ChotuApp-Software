"""Single-owner authenticated cloud dashboard. Put behind HTTPS before public use."""
import base64,hmac,json,os,signal,sys,threading
from http.server import BaseHTTPRequestHandler,ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlsplit,parse_qs
sys.path.insert(0,str(Path(__file__).resolve().parent))
from runtime import CloudService
STATIC=Path(__file__).resolve().parent


def make_server(service,password,host='127.0.0.1',port=8080):
    if len(password)<16:raise ValueError('CHOTUAPP_WEB_PASSWORD must have at least 16 characters')
    expected=('admin:'+password).encode()
    class Handler(BaseHTTPRequestHandler):
        def setup(self):
            super().setup();self.connection.settimeout(30)
        def log_message(self,*args):pass  # Do not print credentials or private URLs to hosting logs.
        def send(self,status,body=b'',content_type='application/json',headers=None):
            if isinstance(body,str):body=body.encode()
            self.send_response(status);self.send_header('Content-Type',content_type);self.send_header('Content-Length',str(len(body)))
            self.send_header('Cache-Control','no-store');self.send_header('X-Content-Type-Options','nosniff')
            self.send_header('Content-Security-Policy',"default-src 'self'; script-src 'self'; style-src 'self'; frame-ancestors 'none'; base-uri 'none'; form-action 'self'")
            for k,v in (headers or {}).items():self.send_header(k,v)
            self.end_headers();self.wfile.write(body)
        def allowed(self):
            header=self.headers.get('Authorization','')
            try:provided=base64.b64decode(header[6:],validate=True) if header.startswith('Basic ') else b''
            except ValueError:provided=b''
            if hmac.compare_digest(provided,expected):return True
            self.send(401,b'Authentication required','text/plain',{'WWW-Authenticate':'Basic realm="ChotuApp", charset="UTF-8"'});return False
        def do_GET(self):
            if not self.allowed():return
            path=urlsplit(self.path)
            if path.path=='/api/status':self.send(200,json.dumps(service.snapshot()));return
            if path.path=='/download':
                name=parse_qs(path.query).get('name',[''])[0];file=service.output/name
                if file.name!=name or file.resolve().parent!=service.output.resolve() or name not in {x['name'] for x in service.snapshot()['files']}:
                    self.send(404,b'File not found','text/plain');return
                self.send(200,file.read_bytes(),'application/octet-stream',{'Content-Disposition':"attachment; filename*=UTF-8''"+__import__('urllib.parse',fromlist=['quote']).quote(name)});return
            files={'/':'index.html','/app.js':'app.js','/style.css':'style.css'}
            if path.path in files:
                file=STATIC/files[path.path];kind={'.html':'text/html; charset=utf-8','.js':'text/javascript; charset=utf-8','.css':'text/css; charset=utf-8'}[file.suffix]
                self.send(200,file.read_bytes(),kind);return
            self.send(404,b'Not found','text/plain')
        def do_POST(self):
            if not self.allowed():return
            origin=self.headers.get('Origin')
            if self.headers.get('X-ChotuApp-Request')!='1' or (origin and urlsplit(origin).netloc!=self.headers.get('Host')):
                self.send(403,b'Invalid request origin','text/plain');return
            try:
                size=int(self.headers.get('Content-Length','0'))
                if not 0<=size<=2*1024*1024:raise ValueError('Request too large')
                body=self.rfile.read(size)
                if self.path=='/api/template':
                    if service.running:raise ValueError('Stop the active run before changing the template')
                    import io,zipfile
                    with zipfile.ZipFile(io.BytesIO(body)) as archive:
                        if sum(item.file_size for item in archive.infolist())>50*1024*1024:raise ValueError('Expanded template exceeds 50 MB')
                    from openpyxl import load_workbook
                    book=load_workbook(io.BytesIO(body),read_only=True)
                    try:
                        if 'Products' not in book.sheetnames:raise ValueError('Template must contain Products')
                    finally:book.close()
                    with service.lock:
                        if service.running:raise ValueError('Run already active')
                        temporary=service.data/'template.pending.xlsx';temporary.write_bytes(body);os.replace(temporary,service.data/'template.xlsx')
                else:
                    data=json.loads(body or b'{}')
                    if self.path=='/api/start':service.start(data.get('sites'))
                    elif self.path=='/api/continue':service.start(resume=True)
                    elif self.path=='/api/control':service.control(data.get('command'))
                    else:self.send(404,b'Not found','text/plain');return
                self.send(200,'{"ok":true}')
            except (ValueError,KeyError,TypeError) as exc:self.send(400,json.dumps({'error':str(exc)}))
            except Exception:self.send(500,'{"error":"Request failed; check server logs or template format"}')
    server=ThreadingHTTPServer((host,port),Handler);server.daemon_threads=True
    return server


def main():
    password=os.environ.get('CHOTUAPP_WEB_PASSWORD','')
    if len(password)<16:raise SystemExit('Set CHOTUAPP_WEB_PASSWORD to a unique password of at least 16 characters')
    service=CloudService(os.environ.get('CHOTUAPP_CLOUD_DATA','/data'))
    server=make_server(service,password,os.environ.get('HOST','127.0.0.1'),int(os.environ.get('PORT','8080')))
    def stop(*args):threading.Thread(target=server.shutdown,daemon=True).start()
    signal.signal(signal.SIGTERM,stop);signal.signal(signal.SIGINT,stop)
    state=service.store.load()
    if os.environ.get('CHOTUAPP_AUTO_CONTINUE','1')=='1' and service.pause.is_set() and state and service.store.pending(state):service.start(resume=True)
    print('ChotuApp cloud dashboard started; username: admin',flush=True)
    try:server.serve_forever()
    finally:server.server_close();service.close()

if __name__=='__main__':main()
