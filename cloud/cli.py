"""Run fetching on a cloud runner without a web server or desktop window."""
import argparse,json,signal,sys,time,zipfile,os,sqlite3
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parent))
from runtime import CloudService


def parse_stores(text):
    if text.lstrip().startswith('['):return json.loads(text)
    result=[]
    for line in text.replace(';','\n').splitlines():
        if not line.strip():continue
        name,separator,url=line.partition('|')
        if not separator:raise ValueError('Use Name | https://website.com, separated by semicolons')
        result.append({'name':name.strip(),'url':url.strip()})
    return result


def checkpoint(data,destination,max_bytes=250*1024*1024):
    data=Path(data);destination=Path(destination);temporary=destination.with_suffix('.pending.zip')
    try:
        with zipfile.ZipFile(temporary,'w',zipfile.ZIP_DEFLATED) as archive:
            for folder in (data/'saved_runs',data/'runtime'/'source_pacing',data/'ANAS_DATA_FETCH'):
                for path in folder.rglob('*'):
                    if path.is_file() and not path.is_symlink() and path.name!='active.lock':archive.write(path,path.relative_to(data).as_posix())
            for name in ('template.xlsx','controls.json'):
                if (data/name).is_file():archive.write(data/name,name)
        if temporary.stat().st_size>max_bytes:
            from checkpoint_parts import split_archive
            parts=destination.with_suffix('.parts');split_archive(temporary,parts,part_bytes=max_bytes)
            destination.unlink(missing_ok=True)
            print('Large checkpoint saved in checksummed parts:',parts,flush=True)
        else:temporary.replace(destination)
    finally:temporary.unlink(missing_ok=True)


def _restore_unlocked(archive,data):
    data=Path(data).resolve();data.mkdir(parents=True,exist_ok=True)
    with zipfile.ZipFile(archive) as zipped:
        import shutil
        expanded=sum(item.file_size for item in zipped.infolist())
        if expanded>8*1024**3:raise ValueError('Checkpoint exceeds 8 GB expanded safety limit')
        if shutil.disk_usage(data).free<expanded+256*1024**2:raise ValueError('Insufficient disk space to restore checkpoint')
        for item in zipped.infolist():
            path=Path(item.filename)
            allowed=(path.parts[:1] in [('saved_runs',),('ANAS_DATA_FETCH',)] or path.parts[:2]==('runtime','source_pacing') or item.filename in ('template.xlsx','controls.json'))
            if not allowed or path.is_absolute() or '..' in path.parts or '\\' in item.filename or (item.external_attr>>16)&0o170000==0o120000:raise ValueError('Unsafe checkpoint entry')
            target=(data/path).resolve()
            if data not in target.parents:raise ValueError('Checkpoint path escaped data folder')
        for item in zipped.infolist():
            target=data/item.filename
            if item.is_dir():target.mkdir(parents=True,exist_ok=True)
            else:
                target.parent.mkdir(parents=True,exist_ok=True)
                with zipped.open(item) as source,target.open('wb') as destination:shutil.copyfileobj(source,destination)

    database=data/'saved_runs'/'run_state.sqlite3'
    if database.exists():
        with sqlite3.connect(database) as db:
            row=db.execute('SELECT payload FROM state WHERE id=1').fetchone()
        state=json.loads(row[0]) if row else {}
        for site in state.get('sites',[]):
            output=site.get('output') or {};file=data/'ANAS_DATA_FETCH'/Path(output.get('path','')).name
            if file.is_file() and file.stat().st_size==output.get('size'):
                os.utime(file,ns=(file.stat().st_atime_ns,output['mtime_ns']))


def restore(archive,data):
    from resume_support import RunLease
    folder=Path(data)/'saved_runs';folder.mkdir(parents=True,exist_ok=True);lease=RunLease(folder);lease.acquire()
    try:
        from checkpoint_parts import archive_path
        with archive_path(archive) as source:_restore_unlocked(source,data)
    finally:lease.release()


def print_logs(service,seen):
    for message in service.snapshot()['logs']:
        if message['id']>seen:print(message['text'],flush=True);seen=message['id']
    return seen


def write_summary(service,checkpoint_status):
    from openpyxl import load_workbook
    state=service.store.load() or {};stores=[]
    for site in state.get('sites',[]):
        item={key:site.get(key) for key in ('name','status','detail','reviewed','software_revision')}
        output=site.get('output') or {};file=Path(output.get('path',''))
        if site.get('status')=='done' and file.is_file():
            book=load_workbook(file,read_only=True)
            try:
                item.update(file=file.name,products=book['Products'].max_row-1,variants=book['Variant Combinations'].max_row-1 if 'Variant Combinations' in book.sheetnames else 0)
            finally:book.close()
        stores.append(item)
    summary={'run_id':state.get('id'),'status':state.get('status','incomplete'),'checkpoint':checkpoint_status,'stores':stores}
    (service.output/'run-summary.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
    lines=['RUN SUMMARY:']
    for item in stores:
        if item.get('file'):lines.append(item['name']+': DONE — '+str(item['products'])+' products, '+str(item['variants'])+' variants; '+item['file']+(' (REVIEW)' if item.get('reviewed') else ''))
        else:lines.append(item['name']+': '+str(item.get('status','incomplete')).upper()+' — '+str(item.get('detail','')))
    lines.append('Checkpoint: '+checkpoint_status);print('\n'.join(lines),flush=True)
    if os.environ.get('GITHUB_STEP_SUMMARY'):
        with Path(os.environ['GITHUB_STEP_SUMMARY']).open('a',encoding='utf-8') as stream:
            stream.write('## ChotuApp results\n\n')
            for line in lines[1:]:stream.write('- '+line.replace('<','&lt;').replace('>','&gt;').replace('\n',' ')+'\n')
    return summary


def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--data',type=Path,default=Path('.cloud-data'));parser.add_argument('--stores',default='');parser.add_argument('--continue-saved',action='store_true');parser.add_argument('--minutes',type=int,default=285);parser.add_argument('--checkpoint',type=Path);parser.add_argument('--restore',type=Path)
    args=parser.parse_args()
    if not 1<=args.minutes<=300:parser.error('Minutes must be between 1 and 300')
    if args.restore:restore(args.restore,args.data)
    service=CloudService(args.data)
    def stop(*unused):service.control('stop')
    signal.signal(signal.SIGTERM,stop);signal.signal(signal.SIGINT,stop)
    result=1
    try:
        service.start(parse_stores(args.stores),resume=args.continue_saved);deadline=time.monotonic()+args.minutes*60;seen=0
        while service.running:
            snapshot=service.snapshot()
            for message in snapshot['logs']:
                if message['id']>seen:print(message['text'],flush=True);seen=message['id']
            if time.monotonic()>=deadline:print('Cloud time budget reached; saving unfinished work.',flush=True);service.control('stop');break
            time.sleep(.5)
        service.close();seen=print_logs(service,seen);state=service.store.load();result=0 if state and state['status']=='complete' else 2
        checkpoint_status='not requested'
        if not args.checkpoint and (args.data/'stop-request').exists():args.checkpoint=Path('checkpoint.zip')
        if args.checkpoint:
            try:
                checkpoint(args.data,args.checkpoint);checkpoint_status='saved';print('Checkpoint saved:',args.checkpoint if args.checkpoint.exists() else args.checkpoint.with_suffix('.parts'),flush=True)
            except ValueError as exc:checkpoint_status='not saved: '+str(exc);print('CHECKPOINT NOT SAVED:',exc,flush=True);result=2
        write_summary(service,checkpoint_status)
        print('Run finished. Excel files:',service.output,flush=True)
    finally:
        if service.running:service.close()
        elif service.lease.file:service.close()
    return result

if __name__=='__main__':sys.exit(main())
