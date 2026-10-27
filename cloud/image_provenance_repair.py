"""Backfill missing metadata in reviewed Czone single-SKU exports; never alter image/stock data."""
import hashlib
from pathlib import Path
from urllib.parse import urlsplit
from collections import Counter
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

def repair(source,out,store):
 if store!='Computer Zone':return source,0
 book=load_workbook(source,read_only=True,data_only=False);patches={};count=0
 try:
  def records(name):
   iterator=book[name].iter_rows(values_only=True);headers=next(iterator)
   return headers,[dict(zip(headers,row)) for row in iterator if any(x is not None for x in row)]
  _,products=records('Products');headers,variants=records('Variant Combinations');products={str(p['product_id']):p for p in products}
  needed={'variant_image_status','fallback_image_url','variant_image_note'}
  if not needed<=set(headers):return source,0
  totals=Counter(str(v.get('product_id')) for v in variants)
  # Keep original Excel row coordinates including empty rows.
  iterator=book['Variant Combinations'].iter_rows(values_only=True);next(iterator)
  for number,values in enumerate(iterator,2):
   v=dict(zip(headers,values));pid=str(v.get('product_id') or '');p=products.get(pid)
   if not p or totals[pid]!=1 or v.get('variant_image_status') or v.get('fallback_image_url'):continue
   if p.get('option_groups') or any(v.get('option_'+str(i)+'_value') or v.get('option_'+str(i)+'_group_id') for i in range(1,4)):continue
   url=str(v.get('source_url') or '');parsed=urlsplit(url)
   if parsed.scheme!='https' or parsed.hostname!='www.czone.com.pk' or url!=p.get('source_url'):continue
   vid=str(v.get('source_variant_id') or '')
   if not vid or pid!='CP'+hashlib.sha256(('https://www.czone.com.pk|'+vid).encode()).hexdigest()[:12]:continue
   image=v.get('variant_image_url');covers={p.get(k) for k in ['final_image_url (main photo)','final_image_url (cover/thumbnail — hosted link)','scraped_image_url (auto)','scraped_image_url (cover, auto)']}
   if not isinstance(image,str) or not image.startswith(('https://','http://')) or image not in covers:continue
   changes={'variant_image_status':'source_linked','variant_image_note':'Identity-verified source product image for this single-SKU product; no separate size/colour variants.'}
   for key,value in changes.items():patches[get_column_letter(headers.index(key)+1)+str(number)]=value
   count+=1
 finally:book.close()
 if not patches:return source,0
 from rule804_pipeline import toolkit
 tools=toolkit();out=Path(out);out.parent.mkdir(parents=True,exist_ok=True)
 part=tools.patch_sheet(source,out,patches,sheet_name='Variant Combinations');tools.verify(source,out,part,patches,sheet_name='Variant Combinations')
 return out,count
