"""Fail-closed 804 Audit delivery, matching existing Drive workbooks by confirmed shop code."""
import hashlib,json,shutil,tempfile
from pathlib import Path
from datetime import datetime,timezone
from urllib.parse import urlsplit
from openpyxl import load_workbook
from audit804 import run as audit
from rule804_pipeline import process
from legacy_original_schema import normalize
from audit804_queue import save as save_held,restore as restore_held
ROOT=Path(__file__).resolve().parents[1]

def codes_in(path):
 book=load_workbook(path,read_only=True,data_only=False);found=set();blank=False
 try:
  iterator=book['Products'].iter_rows(values_only=True);headers=list(next(iterator))
  if 'shop_code' not in headers:return set()
  index=headers.index('shop_code')
  for row in iterator:
   if not any(value is not None for value in row):continue
   value=str(row[index] or '').strip()
   if value:found.add(value)
   else:blank=True
 finally:book.close()
 return found if not blank else set()

class DriveIdentity:
 def __init__(self,rc,folder,backup=None):self.backup=backup;self.rc=rc;self.remote='chotudrive:'+folder;self.tmp=None;self.names={};self.index={}
 def __enter__(self):
  self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
  listing=json.loads(self.rc('lsjson',self.remote,'--files-only','--max-depth','1') or '[]')
  for item in listing:
   name=item.get('Name','')
   if not name.lower().endswith('.xlsx'):continue
   if Path(name).name!=name:raise ValueError('Unexpected Drive file path')
   if name in self.names:raise ValueError('Duplicate Drive filenames require identity review')
   self.names[name]=item
  if self.names:
   self.rc('copy',self.remote,str(self.root),'--include','*.xlsx','--max-depth','1')
   for name in self.names:
    try:codes=codes_in(self.root/name)
    except Exception:codes=set()
    if len(codes)==1:self.index.setdefault(next(iter(codes)),[]).append(name)
  return self
 def __exit__(self,*args):
  if self.tmp:self.tmp.cleanup()
 def replace(self,file,code,canonical):
  if not code or codes_in(file)!={code}:raise ValueError('New workbook shop code is not a single confirmed identity')
  matches=self.index.get(code,[])
  if len(matches)>1:raise ValueError('Multiple existing files use this shop code; replacement held')
  old=matches[0] if matches else None
  destination=old or canonical
  if destination in self.names and destination!=old:raise ValueError('Target filename belongs to an unconfirmed/different shop code; old file retained')
  if old and self.backup:self.backup(self.root/old,code)
  if old:self.rc('check',str(self.root),self.remote,'--include',old,'--one-way')
  with tempfile.TemporaryDirectory() as tmp:
   stage=Path(tmp);target=stage/destination;shutil.copy2(file,target)
   try:
    self.rc('copy',str(stage),self.remote,'--include',destination,'--checksum')
    self.rc('check',str(stage),self.remote,'--include',destination,'--one-way')
   except Exception:
    if old and (self.root/old).exists():self.rc('copyto',str(self.root/old),self.remote+'/'+old,'--checksum')
    elif destination not in self.names:
     try:self.rc('deletefile',self.remote+'/'+destination)
     except Exception:pass
    raise
   shutil.copy2(file,self.root/destination);self.names[destination]={'Name':destination};self.index[code]=[destination]
  return {'replaced_file':old,'file_name':destination}

def completed_files(stage):
 manifests=list(stage.rglob('completion-manifest.json'))
 if len(manifests)==1:return [row for row in json.loads(manifests[0].read_text()).get('files',[]) if row.get('complete') is True]
 summaries=list(stage.rglob('run-summary.json'))
 if len(summaries)!=1:return []
 return [{'store':row['name'],'name':row.get('file'),'complete':True} for row in json.loads(summaries[0].read_text()).get('stores',[]) if row.get('status')=='done' and row.get('file')]

def run_campaign(state,config,all_runs,gh,rc,*,root=ROOT):
 from importlib.util import spec_from_file_location,module_from_spec
 spec=spec_from_file_location('audit_groups',root/'catalog-tools/start_fresh26_rule804_drive_windows.py');helper=module_from_spec(spec);spec.loader.exec_module(helper)
 campaign=config['dispatcher'];folder=config['destination_folder'];prefix='Cloud group audit804-'+campaign+'-'
 groups={run['displayTitle'].removeprefix('Cloud group ') for run in all_runs if run.get('displayTitle','').startswith(prefix)}
 runs=helper.related(all_runs,groups);registry=json.loads((root/'cloud/fetch-update-release.json').read_text())['fresh_stores'];sites={s['name']:s for s in registry}
 codes=json.loads((root/'cloud/confirmed-shop-codes.json').read_text())
 revision=hashlib.sha256(b''.join((root/name).read_bytes() for name in ('cloud/audit804.py','cloud/audit804_semantics.py','cloud/reviewed804_repairs.py','cloud/audit804_delivery.py','cloud/rule804_pipeline.py','cloud/legacy_original_schema.py','cloud/confirmed-shop-codes.json','subcategory_inference_support.py','catalog-tools/complete_rule804_batch.py'))).hexdigest()
 state['audit804_campaign']=campaign;state.setdefault('audit804_results',{});state.setdefault('audit804_processed',{});identity=None
 def handle(source,store,source_run,artifact_id,cache_key):
  nonlocal identity
  prior=state['audit804_results'].get(store,{});source_sha=hashlib.sha256(source.read_bytes()).hexdigest();marker=prior.get('held_input') if prior.get('source_sha256')==source_sha else None;code=codes.get(store)
  if store=='LAAM':
   import base64
   details=next((r for r in all_runs if str(r['databaseId'])==str(source_run)),{})
   head=details.get('headSha') or json.loads(gh('run','view',str(source_run),'--repo','chottuapp/ChotuApp-Software','--json','headSha'))['headSha']
   evidence=json.loads(gh('api','repos/chottuapp/ChotuApp-Software/contents/laam_source_support.py?ref='+head))
   if hashlib.sha256(base64.b64decode(evidence['content'])).digest()!=hashlib.sha256((root/'laam_source_support.py').read_bytes()).digest():raise ValueError('Rule LAAM clothing-only adapter revision not confirmed')
  with tempfile.TemporaryDirectory() as tmp:
   stage=Path(tmp);normalized,migrated=normalize(source,stage/'normalized'/source.name,code);output=stage/'checked'/(store+'.xlsx')
   processing=process(normalized,output,store,codes,approved_original=store in {'LAAM','Sapphire','Khaadi'})
   host=(urlsplit(sites[store]['url']).hostname or '').removeprefix('www.');hosts={host}
   if store=='Sapphire':hosts={'pk.sapphireonline.pk'}
   if store=='Farda Foods':
    from farda_source_support import SHOP
    hosts.add(urlsplit(SHOP).hostname)
   report=audit(output,store,codes,hosts,source=normalized)
   repairs=[r for r in processing.get('warnings',[]) if r.get('issue','').startswith(('Copied','Removed'))]
   report.update(description_repairs_count=len(repairs),description_repairs_examples=repairs[:10])
   report.update(campaign_id=campaign,source_run=str(source_run),source_artifact=artifact_id,source_sha256=source_sha)
   if report['status']!='pass':
    if not marker:marker=save_held(source,store,campaign,str(source_run))
    report['held_input']=marker
   else:
    if identity is None:
     def backup_old(path,shop_code):
      marker=save_held(path,'rollback-'+shop_code,campaign,'drive-before-replace');state.setdefault('audit804_replacement_backups',{})[shop_code]=marker
     identity=DriveIdentity(rc,folder,backup_old);identity.__enter__()
    replaced=identity.replace(output,code,store+'.xlsx')
    record={'shop_code':code,'source_run':str(source_run),'sha256':report['workbook_sha256'],'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'folder':folder,'verified_utc':datetime.now(timezone.utc).isoformat(),'audit804':{'name':'804 Audit','status':'pass','checked_utc':report['checked_utc'],'products_checked':report['products_checked'],'workbook_sha256':report['workbook_sha256']},'rule804':'applied; 804 Audit passed','campaign_id':campaign,**replaced}
    state.setdefault('uploaded',{})[store]=record
    state.setdefault('original_uploads',{}).pop(store,None)
    state.setdefault('daily_uploads',{}).setdefault(folder,{})[store]=record
   state['audit804_results'][store]=report;state['audit804_processed'][store]=cache_key
 def safe_handle(source,store,rid,artifact_id,key):
  try:handle(source,store,rid,artifact_id,key)
  except Exception as exc:
   prior=state['audit804_results'].get(store,{})
   marker=prior.get('held_input') if prior.get('source_sha256')==hashlib.sha256(source.read_bytes()).hexdigest() else None
   if not marker:marker=save_held(source,store,campaign,str(rid))
   state['audit804_results'][store]={'name':'804 Audit','status':'hold','store':store,'campaign_id':campaign,'source_run':str(rid),'source_artifact':artifact_id,'error':str(exc),'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'held_input':marker,'checked_utc':datetime.now(timezone.utc).isoformat()}
   state['audit804_processed'][store]=key
 try:
  # Recheck durable held inputs after code/policy updates, even if Actions artifacts expired.
  for store,report in list(state['audit804_results'].items()):
   key='source:'+report.get('source_sha256','')+':'+store+':'+revision
   if report.get('status')!='hold' or not report.get('held_input') or state['audit804_processed'].get(store,'').endswith(revision):continue
   with tempfile.TemporaryDirectory() as tmp:
    source=restore_held(report['held_input'],Path(tmp)/(store+'.xlsx'),store,campaign);safe_handle(source,store,report.get('source_run','phone'),report.get('source_artifact'),key)
  for run in reversed(runs):
   rid=str(run['databaseId']);artifacts=json.loads(gh('api','repos/chottuapp/ChotuApp-Software/actions/runs/'+rid+'/artifacts?per_page=100'))['artifacts'];found=[a for a in artifacts if a.get('name')=='chotuapp-excel-files' and not a.get('expired')]
   if not found:continue
   item=max(found,key=lambda a:a['id']);base=rid+':'+str(item['id'])
   if state.setdefault('audit804_artifacts',{}).get(rid)==base+':'+revision:continue
   with tempfile.TemporaryDirectory() as tmp:
    stage=Path(tmp);gh('run','download',rid,'--repo','chottuapp/ChotuApp-Software','--name','chotuapp-excel-files','--dir',str(stage))
    for entry in completed_files(stage):
     store=entry.get('store');name=entry.get('name','')
     if store not in config['stores']:raise ValueError('Unapproved store in completed source')
     key=base+':'+store+':'+revision
     if store not in config['stores'] or Path(name).name!=name or name!=store+'.xlsx':raise ValueError('Completed source outside approved batch')
     matches=list(stage.rglob(name))
     if len(matches)!=1:raise ValueError('Unique completed source required')
     actual_sha=hashlib.sha256(matches[0].read_bytes()).hexdigest();key='source:'+actual_sha+':'+store+':'+revision
     if state['audit804_processed'].get(store)==key:continue
     if entry.get('sha256') and actual_sha!=entry['sha256']:raise ValueError('Completed source checksum mismatch')
     safe_handle(matches[0],store,rid,item['id'],key)
   state['audit804_artifacts'][rid]=base+':'+revision
  marker=root/'catalog-inputs/khaadi-mobile/latest.json'
  if marker.exists():
   metadata=json.loads(marker.read_text())
   if metadata.get('fresh_campaign_id')==campaign:
    from mobile_khaadi_ingest import input_file
    supplied=input_file(root)
    if supplied:
     source,sha=supplied;key='source:'+sha+':Khaadi:'+revision
     if state['audit804_processed'].get('Khaadi')!=key:safe_handle(source,'Khaadi','phone',None,key)
 finally:
  if identity:identity.__exit__(None,None,None)
 state['sync_status']='finished';state['audit804_checked_utc']=datetime.now(timezone.utc).isoformat()
 return state
