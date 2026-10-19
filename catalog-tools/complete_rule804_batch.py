"""Evidence-bound final catalog repairs with exact preservation of untouched XLSX parts."""
import argparse,collections,hashlib,json,re,zipfile,xml.etree.ElementTree as ET
from pathlib import Path
from datetime import datetime,timezone
from openpyxl import load_workbook
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from subcategory_inference_support import infer
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main';REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
ET.register_namespace('',NS)
FOLDERS=['rule804-20261019-19-reviewed','rule804-20261019-next2-reviewed','rule804-20261020-junaid-zellbury-reviewed']
BROAD={'ZARIYA','LUXE','BOTTOMWEAR','BOTTOM WEAR','BAGS','JEWELRY','EYES','HAIR','ESSENTIALS','READY TO WEAR','INFANT GIRLS','ACCESSORIES','LIP PENCIL','FACE POWDER','EYE LINER & MASCARA','ARTISANAL JEWELLERY'}
EXTRA=[('Hair Clips & Pins',r'\bclaw\s*clips?\b'),('Hair Ties & Accessories',r'\b(?:pearl beaded|organza)\s+pony\b'),('Headbands',r'\bhairbands?\b'),('Hair Combs',r'\bhair trimming comb\b'),('Hair Jewellery',r'\bhair vine\b'),('Forehead Jewellery',r'\bbindiya\b|\bmathapatti\b'),('Armlets',r'\barmlet\b|\bbazu band\b'),('Hand Harnesses',r'\bpanjangla\b'),('Earrings',r'\bbaliyan\b'),('Commemorative Coin Cards',r'\b(?:gold )?coin card\b')]

def extra_category(row,source=None):
 name=str(row.get('name') or '');body=str(row.get('description') or '');proof=[]
 for label,pattern in EXTRA:
  if re.search(pattern,name,re.I):return label,'inferred_product_name','Explicit named product type: '+name
 if re.search(r'false nail art|(?:flower|heart) nails',name,re.I) and (re.search(r'false nail',name,re.I) or re.search(r'nail art\s*\(\d+\s*pcs',body,re.I)):
  return 'Artificial Nails','inferred_product_name','Named nail product and supplied pack contents: '+name+'; '+body[:300]
 type_text=str((source or {}).get('product_type') or '')
 supplied_type=' '.join(str(v or '') for v in [type_text,(source or {}).get('vendor'),row.get('brand')])
 field='official product_type'
 if not type_text:type_text=str(row.get('brand') or '');field='supplied brand/type field'
 tags=[str(t) for t in ((source or {}).get('tags') or [])]
 if source and any(re.search(r'\bshawls?\b',t,re.I) for t in tags) and re.search(r'\bvelvet\b',body,re.I):return 'Shawls','inferred_product_type','Official product tag: '+', '.join(t for t in tags if re.search('shawl',t,re.I))+'; supplied velvet description'
 if re.search(r'\b(?:deal|bundle|set)\b',name,re.I) and re.search(r'\bfragrances?\b|\bperfumes?\b',type_text,re.I):return 'Fragrance Sets','inferred_product_type',field+': '+type_text+'; bundle name: '+name
 if re.search(r'\bset\b',name,re.I) and re.search(r'\b(?:necklace|choker)\b',type_text,re.I):return 'Jewellery Sets','inferred_product_type',field+': '+type_text+'; set name: '+name
 if re.search(r'\b(?:jeans|denim trousers|denim pants)\b',type_text,re.I):return 'Jeans','inferred_product_type',field+': '+type_text
 if re.search(r'\b(?:[23]\s*piece|[23]-piece)\b',supplied_type,re.I) and re.search(r'\bstitched\b',supplied_type,re.I) and not re.search('unstitched',supplied_type,re.I):return 'Stitched Outfit Sets','inferred_product_type','Explicit supplied product type/vendor: '+supplied_type
 if re.search(r'\bshalwar\b',type_text,re.I) and not re.search(r'\bkameez\b',type_text,re.I):return 'Shalwars','inferred_product_type',field+': '+type_text
 if re.search(r'\bkurta\b',type_text,re.I):return 'Kurtas & Kurtis','inferred_product_type',field+': '+type_text
 if re.search(r'terracotta powder',name,re.I) and re.search('powder|cosmetics',supplied_type,re.I):return 'Face Powders','inferred_product_type',field+': '+type_text+'; name: '+name
 if re.search(r'\bstitch(?:ed)?\b',body,re.I) and re.search(r'\b(?:three|two|[23])[- ]piece\b',body,re.I) and re.search(r'\b(?:outfit|set|ensemble)\b',body,re.I) and not re.search('unstitched',body,re.I):return 'Stitched Outfit Sets','inferred_description','Explicit stitched multi-piece outfit: '+body[:600]
 # Only explicit source type labels, never price-band collections or marketing tags.
 if type_text and field=='official product_type':
  result=infer(type_text,'')
  if result:return result[0],'inferred_product_type','Official product_type: '+type_text
 return None

def tidy_description(text):
 if not isinstance(text,str):return text
 # Correct prose only. Keep every numeric value, warning, detail, option and product name.
 fixes=[(r'\bwill be Deliverd\b','will be delivered'),(r'\bMix Designs will be delivered\b','Mixed designs will be delivered'),(r'Colour May vary','Colour may vary'),(r'\bClean it with dry and soft cloth\b','Clean it with a dry, soft cloth')]
 result=re.sub(r'(?im)^\s*dsasd\s*$','',text)
 for pattern,replacement in fixes:result=re.sub(pattern,replacement,result,flags=re.I)
 return result

def product_overview(row,label,source=None):
 name=str(row.get('name') or '').strip();kind=str((source or {}).get('product_type') or '')
 phrases={'Shawls':'a shawl','Fragrance Sets':'a fragrance bundle','Hair Clips & Pins':'a hair claw clip' if 'claw' in name.lower() else 'a hair clip','Hair Ties & Accessories':'a hair tie','Headbands':'a headband','Hair Combs':'a hair-styling comb','Hair Jewellery':'a hair-jewellery design','Forehead Jewellery':'a forehead-jewellery design','Armlets':'an armlet','Hand Harnesses':'a hand-jewellery harness','Earrings':'an earring style','Artificial Nails':'an artificial nail-art set','Commemorative Coin Cards':'a commemorative coin card','Shalwars':'a shalwar','Kurtas & Kurtis':'a kurta','Stitched Outfit Sets':'a stitched multi-piece outfit','Face Powders':'a face powder','Jeans':'a pair of jeans','Gemstones':'a stone listed as aqeeq'}
 if label=='Jewellery Sets':
  phrases[label]='a choker-style jewellery set' if re.search('choker',kind,re.I) else 'a jewellery set'
 if label=='Shawls' and re.search('velvet',str(row.get('description')),re.I):phrases[label]='a velvet shawl'
 return name+' is '+phrases[label]+'.' if label in phrases else None

def patch_sheet(source,out,patches):
 with zipfile.ZipFile(source) as z:
  book=ET.fromstring(z.read('xl/workbook.xml'));sheets=book.find('{'+NS+'}sheets');rid=next(s.attrib['{'+REL+'}id'] for s in sheets if s.attrib['name']=='Products')
  rels=ET.fromstring(z.read('xl/_rels/workbook.xml.rels'));target=next(r.attrib['Target'] for r in rels if r.attrib['Id']==rid);part=target.lstrip('/') if target.startswith('/') else 'xl/'+target
  tree=ET.fromstring(z.read(part));data=tree.find('{'+NS+'}sheetData');byrow={int(r.attrib['r']):r for r in data}
  for address,value in patches.items():
   rnum=int(re.search(r'\d+',address)[0]);row=byrow[rnum];cell=next((c for c in row if c.attrib.get('r')==address),None)
   if cell is None:
    cell=ET.Element('{'+NS+'}c',{'r':address})
    def number(addr):
     n=0
     for char in re.match(r'[A-Z]+',addr)[0]:n=n*26+ord(char)-64
     return n
    n=number(address);pos=next((i for i,c in enumerate(row) if number(c.attrib['r'])>n),len(row));row.insert(pos,cell)
   if cell.find('{'+NS+'}f') is not None:raise ValueError('Refusing to overwrite formula '+address)
   for child in list(cell):cell.remove(child)
   cell.attrib['t']='inlineStr';inline=ET.SubElement(cell,'{'+NS+'}is');t=ET.SubElement(inline,'{'+NS+'}t');t.attrib['{http://www.w3.org/XML/1998/namespace}space']='preserve';t.text=str(value)
  content=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
  with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as target:
   for info in z.infolist():target.writestr(info,content if info.filename==part else z.read(info.filename))
 return part

def verify(source,out,part,patches):
 with zipfile.ZipFile(source) as a,zipfile.ZipFile(out) as b:
  assert set(a.namelist())==set(b.namelist())
  for name in a.namelist():
   if name!=part:assert a.read(name)==b.read(name),name
  before=ET.fromstring(a.read(part));after=ET.fromstring(b.read(part))
  # Remove only the authorised replacement cells. Every other XML element must match.
  for tree in (before,after):
   for row in tree.find('{'+NS+'}sheetData'):
    for cell in list(row):
     if cell.attrib.get('r') in patches:row.remove(cell)
  assert ET.tostring(before)==ET.tostring(after),'Non-authorised Products XML changed'
 w=load_workbook(out,read_only=True,data_only=False)
 try:
  s=w['Products']
  seen=set()
  for row in s.iter_rows():
   for cell in row:
    address=getattr(cell,'coordinate',None)
    if address in patches:
     assert cell.value==patches[address],address;seen.add(address)
  assert seen==set(patches)
 finally:w.close()

def main():
 ap=argparse.ArgumentParser();ap.add_argument('--evidence',type=Path);ap.add_argument('--visual-decisions',type=Path);ap.add_argument('--out',type=Path,default=Path('catalog-files/rule804-20261020-completed23'));args=ap.parse_args();args.out.mkdir(parents=True,exist_ok=True)
 evidence={}
 if args.evidence:
  evidence={r['source_url']:r for r in json.loads(args.evidence.read_text())['results'] if not r.get('error')}
 visual={}
 if args.visual_decisions:visual={(x['store'],x['product_id']):x for x in json.loads(args.visual_decisions.read_text())}
 manifest={'rule':'Current repository Rule 804, including owner-authorized unknown-code blanks and product-purpose Category override','checked_at_utc':datetime.now(timezone.utc).isoformat(),'files':[]};unresolved=[]
 from openpyxl.utils import get_column_letter
 for folder in FOLDERS:
  for source in sorted((Path('catalog-files')/folder).glob('*.xlsx')):
   w=load_workbook(source,read_only=True,data_only=False);sheet=w['Products'];iterator=sheet.iter_rows(values_only=True);headers=next(iterator)
   assert list(headers[:3])==['shop_code','product_id','Category'],source
   indices={label:i+1 for i,label in enumerate(headers)};patches={};cat_changes=0;desc_changes=0;n=0;ids=set();codes=set();missing=[];changes=[]
   for r,values in enumerate(iterator,2):
    if not any(v is not None for v in values):continue
    row=dict(zip(headers,values));n+=1;pid=row['product_id'];assert pid and pid not in ids,(source,pid);ids.add(pid)
    if row['shop_code']:codes.add(row['shop_code'])
    old=str(row.get('Category') or '');proof=evidence.get(row.get('source_url'))
    inferred=None
    if str(row.get('subcategory_source') or '').casefold() not in ('manual','owner','manual_override') and (not old or old in BROAD):
     inferred=extra_category(row,proof)
     if not inferred and old in BROAD:inferred=infer(row.get('name'),row.get('description'))
     decision=visual.get((source.stem,pid))
     if not inferred and decision:inferred=(decision['category'],'reviewed_product_image',decision['evidence'])
    if inferred and inferred[0]!=old:
     label,origin,quote=inferred;stamp=manifest['checked_at_utc']
     for col,value in {'Category':label,'subcategory_source':origin,'subcategory_evidence':quote+'; product URL: '+str(row.get('source_url')),'subcategory_checked_at':stamp}.items():
      assert col in indices;patches[get_column_letter(indices[col])+str(r)]=value
     cat_changes+=1;changes.append({'product_id':pid,'name':row.get('name'),'old_category':old,'new_category':label,'evidence':quote})
    if not old and not inferred:missing.append({'product_id':pid,'name':row.get('name'),'url':row.get('source_url'),'reason':'Insufficient unambiguous product-type evidence; not invented'})
    before=row.get('description');after=tidy_description(before)
    if inferred and inferred[0]!=old and isinstance(after,str):
     overview=product_overview(row,inferred[0],proof);lead=str(row.get('name') or '')+'.'
     if overview and after.startswith(lead):after=overview+after[len(lead):]
    if after!=before:
     assert re.findall(r'\d+(?:[.,]\d+)?',str(before))==re.findall(r'\d+(?:[.,]\d+)?',str(after));patches[get_column_letter(indices['description'])+str(r)]=after;desc_changes+=1
   w.close();out=args.out/source.name
   if patches:part=patch_sheet(source,out,patches);verify(source,out,part,patches)
   else:out.write_bytes(source.read_bytes())
   record={'file':source.name,'source_path':str(source),'products':n,'known_shop_codes':sorted(codes),'categories_repaired':cat_changes,'descriptions_polished':desc_changes,'unresolved_categories':len(missing),'changed_cells':len(patches),'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'source_md5':hashlib.md5(source.read_bytes()).hexdigest(),'sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'md5':hashlib.md5(out.read_bytes()).hexdigest(),'source_bytes':source.stat().st_size,'bytes':out.stat().st_size,'checks':['Non-Products XLSX parts byte-identical','All unmodified Products XML unchanged','Every patch checked against reopened workbook','Unique product IDs; original order retained','Unknown codes deliberately preserved blank','Prices, variants, availability, images and formulas unchanged'],'changes':changes}
   manifest['files'].append(record);unresolved.extend({'store':source.stem,**row} for row in missing);print(source.name,n,'category fixes',cat_changes,'description fixes',desc_changes,'unresolved',len(missing),flush=True)
   (args.out/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
 (args.out/'unresolved.json').write_text(json.dumps(unresolved,ensure_ascii=False,indent=2));print('TOTAL',len(manifest['files']),'files; remaining categories',len(unresolved))
if __name__=='__main__':main()
