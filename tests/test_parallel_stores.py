import json,sys,tempfile,threading,time,unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from parallel_store_support import StoreScheduler,source_key
from resume_support import RunStore
from install_parallel_stores import transform,HOOK,install
from openpyxl import Workbook

ENGINE='''import os,time,json,threading,re
from pathlib import Path
from openpyxl import Workbook
_chotu_tls=threading.local()
_chotu_pause_gate=threading.Event();_chotu_pause_gate.set()
def _chotu_wait():_chotu_pause_gate.wait()
def _chotu_open_checkpoint(base,name):return _chotu_tls.saved_api_path
OUT=Path('.')
class App:
    @staticmethod
    def _output_snapshot(name):return {}
    def _run_checked(self,site,template):
        folder=_chotu_tls.saved_api_path.parent
        name=site['name']; safe=re.sub(r'[^\\w .-]','',name).strip()
        _chotu_open_checkpoint(site['url'],name)
        meta={'pid':os.getpid(),'start':time.time(),'workers':_chotu_tls.fast_worker_count}
        (folder/'started.json').write_text(json.dumps(meta))
        mode=site.get('mode','ok')
        if mode=='crash':os._exit(3)
        if mode=='cooldown':
            self.events.put('FAILED: '+name+': SOURCE PAUSED: website rate limit; retry this saved store after at least 1s.')
            self.events.put(('__JOB_RESULT__',False,name));return
        countfile=folder/'count.txt';count=int(countfile.read_text()) if countfile.exists() else 0
        target=site.get('steps',12)
        while count<target:
            _chotu_wait();time.sleep(.02);count+=1;countfile.write_text(str(count))
        if mode!='stale':
            wb=Workbook();wb.active.title='Products';wb.active.append(['product_id','name','base_price (PKR)']);wb.active.append(['P1',name,123]);wb.save(OUT/(safe+'.xlsx'))
        meta['end']=time.time();(folder/'finished.json').write_text(json.dumps(meta))
        self.events.put('SAVED REVIEW EXCEL: '+str(OUT/(safe+'.xlsx')))
        self.events.put(('__JOB_RESULT__',True,name))
'''

class ParallelStoreTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.engine=self.root/'main.py';self.engine.write_text(ENGINE)
        self.template=self.root/'template.xlsx';Workbook().save(self.template)
        self.store=RunStore(self.root/'saved_runs');self.output=self.root/'exports';self.output.mkdir()
        self.pause=threading.Event();self.pause.set();self.events=[]
    def tearDown(self):self.temp.cleanup()
    def build(self,sites,**kwargs):
        state=self.store.fresh(sites,self.template)
        return StoreScheduler(self.engine,self.store,state,self.output,self.pause,self.events.append,**kwargs)
    def meta(self,scheduler,index,name):return json.loads((self.store.root/scheduler.state['id']/str(index)/name).read_text())
    def wait_for(self,condition,timeout=8):
        end=time.monotonic()+timeout
        while time.monotonic()<end:
            if condition():return
            time.sleep(.02)
        self.fail('Timed out waiting for process event')
    def test_different_domains_overlap_and_export_to_correct_stores(self):
        scheduler=self.build([{'name':'Seran','url':'https://seran.example','steps':25},{'name':'Elite','url':'https://elite.example','steps':25}])
        scheduler.run();a=self.meta(scheduler,0,'finished.json');b=self.meta(scheduler,1,'finished.json')
        self.assertNotEqual(a['pid'],b['pid']);self.assertLess(max(a['start'],b['start']),min(a['end'],b['end']))
        self.assertTrue((self.output/'Seran.xlsx').is_file());self.assertTrue((self.output/'Elite.xlsx').is_file())
        self.assertEqual(self.store.load()['status'],'complete');self.assertTrue(all(a['workers']<=2 for a in (a,b)))
    def test_duplicate_output_names_rejected_before_workers(self):
        with self.assertRaisesRegex(RuntimeError,'distinct name'):
            self.build([{'name':'Same','url':'https://a.example'},{'name':'same','url':'https://b.example'}])
    def test_same_domain_never_overlaps(self):
        scheduler=self.build([{'name':'A','url':'https://www.one.example'},{'name':'B','url':'https://one.example/collection'}]);scheduler.run()
        self.assertGreaterEqual(self.meta(scheduler,1,'started.json')['start'],self.meta(scheduler,0,'finished.json')['end'])
        self.assertEqual(source_key({'url':'https://shop.fardafoods.com'}),source_key({'url':'https://fardafoods.com'}))
    def test_at_most_two_store_processes(self):
        scheduler=self.build([{'name':name,'url':'https://'+name+'.example','steps':20} for name in ('A','B','C')]);scheduler.run()
        a=self.meta(scheduler,0,'finished.json');b=self.meta(scheduler,1,'finished.json');third=self.meta(scheduler,2,'started.json')
        self.assertGreaterEqual(third['start'],min(a['end'],b['end']))
    def test_existing_long_server_cooldown_is_deferred_without_process_or_wait(self):
        state=self.store.fresh([{'name':'Paused','url':'https://paused.example'}],self.template)
        self.store.update(lambda s:s['sites'][0].update(status='source_paused',retry_after_epoch=time.time()+3600))
        scheduler=StoreScheduler(self.engine,self.store,self.store.load(),self.output,self.pause,self.events.append,wait_limit=120)
        start=time.monotonic();scheduler.run()
        self.assertLess(time.monotonic()-start,1);self.assertEqual(scheduler.deferred,1)
        self.assertFalse((self.store.root/state['id']/'0').exists());self.assertEqual(self.store.load()['sites'][0]['status'],'source_paused')
    def test_cooldown_does_not_block_another_store(self):
        scheduler=self.build([{'name':'Slow','url':'https://slow.example','mode':'cooldown'},{'name':'Good','url':'https://good.example'}],retry_limit=0)
        scheduler.run();state=self.store.load();self.assertEqual(state['sites'][0]['status'],'source_paused')
        self.assertEqual(state['sites'][1]['status'],'done');self.assertEqual(state['status'],'incomplete');self.assertTrue((self.output/'Good.xlsx').is_file())
    def test_crash_is_failed_without_marking_old_export_done(self):
        scheduler=self.build([{'name':'Bad','url':'https://bad.example','mode':'crash'},{'name':'Good','url':'https://good.example'}]);scheduler.run()
        self.assertEqual(self.store.load()['sites'][0]['status'],'failed');self.assertEqual(self.store.load()['sites'][1]['status'],'done')
    def test_source_block_is_saved_separately_without_retry_loop(self):
        altered=ENGINE.replace("        if mode=='crash':os._exit(3)","        if mode=='blocked':\n            self.events.put('FAILED: '+name+': SOURCE BLOCKED: public catalog denied')\n            self.events.put(('__JOB_RESULT__',False,name));return\n        if mode=='crash':os._exit(3)")
        self.engine.write_text(altered)
        scheduler=self.build([{'name':'Blocked','url':'https://blocked.example','mode':'blocked'},{'name':'Good','url':'https://good.example'}]);scheduler.run()
        state=self.store.load();self.assertEqual(state['sites'][0]['status'],'blocked');self.assertEqual(state['sites'][1]['status'],'done')
        self.assertEqual(scheduler.attempts[0],1);self.assertEqual(self.store.pending(state),[0])
    def test_raw_http_limit_becomes_resumable_cooldown(self):
        altered=ENGINE.replace("        if mode=='crash':os._exit(3)","        if mode=='raw429':\n            from fast_fetch_support import RequestGate\n            gate=RequestGate(lambda:None,lambda text:None);gate.limited=True;gate.cooldown=time.monotonic()+3600;_chotu_tls.fast_gate=gate\n            raise RuntimeError('429 Too Many Requests')\n        if mode=='crash':os._exit(3)")
        self.engine.write_text(altered)
        scheduler=self.build([{'name':'Limited','url':'https://limited.example','mode':'raw429'}],retry_limit=0)
        scheduler.run();site=self.store.load()['sites'][0]
        self.assertEqual(site['status'],'source_paused');self.assertGreater(site['retry_after_epoch'],time.time()+3500)
    def test_auto_resume_skips_terminal_failure_manual_resume_retries(self):
        state=self.store.fresh([{'name':'Bad','url':'https://bad.example'}],self.template)
        self.store.update(lambda s:s['sites'][0].update(status='failed'))
        for skip,expected in ((True,[]),(False,[0])):
            scheduler=StoreScheduler(self.engine,self.store,self.store.load(),self.output,self.pause,self.events.append,skip_terminal=skip)
            self.assertEqual(scheduler.pending,expected)
    def test_stale_existing_export_is_not_success(self):
        wb=Workbook();wb.active.title='Products';wb.active.append(['id']);wb.active.append(['old']);wb.save(self.output/'Old.xlsx')
        scheduler=self.build([{'name':'Old','url':'https://old.example','mode':'stale'}]);scheduler.run()
        self.assertEqual(self.store.load()['sites'][0]['status'],'failed')
    def test_completed_store_skipped_on_resume(self):
        scheduler=self.build([{'name':'A','url':'https://a.example'}]);scheduler.run();before=(self.output/'A.xlsx').stat().st_mtime_ns
        resumed=StoreScheduler(self.engine,self.store,self.store.load(),self.output,self.pause,self.events.append);resumed.run()
        self.assertEqual(resumed.pending,[]);self.assertEqual((self.output/'A.xlsx').stat().st_mtime_ns,before)
    def test_pause_and_stop_preserve_checkpoint_for_resume(self):
        scheduler=self.build([{'name':'A','url':'https://a.example','steps':50}]);thread=threading.Thread(target=scheduler.run);thread.start()
        folder=self.store.root/scheduler.state['id']/'0';self.wait_for(lambda:(folder/'count.txt').exists())
        self.pause.clear();time.sleep(.2);count=int((folder/'count.txt').read_text());time.sleep(.2)
        self.assertEqual(int((folder/'count.txt').read_text()),count)
        scheduler.shutdown();thread.join(8);self.assertFalse(thread.is_alive());self.assertEqual(self.store.load()['sites'][0]['status'],'pending')
        self.assertGreater(count,0);self.pause.set()
        old_revision=self.store.load()['sites'][0]['software_revision']
        self.engine.write_text(ENGINE.replace("meta['end']=time.time()","meta['updated_code']=True;meta['end']=time.time()"))
        resumed=StoreScheduler(self.engine,self.store,dict(self.store.load(),_resuming=True),self.output,self.pause,self.events.append);resumed.run()
        self.assertEqual((folder/'count.txt').read_text(),'50');self.assertEqual(self.store.load()['status'],'complete')
        site=self.store.load()['sites'][0]
        self.assertNotEqual(site['software_revision'],old_revision);self.assertEqual(site['previous_software_revision'],old_revision)
        self.assertTrue(self.meta(resumed,0,'finished.json')['updated_code'])
    def test_code_revision_invalidates_derived_parser_cache_only(self):
        from software_revision_support import prepare_revision
        from types import SimpleNamespace
        from unittest.mock import patch
        folder=self.root/'checkpoint';folder.mkdir();raw=folder/'pages.sqlite3';raw.write_bytes(b'saved response')
        fake=SimpleNamespace(PARSER_VERSION='webx-test')
        with patch.dict(sys.modules,{'czone_source_support':fake}):
            first,previous=prepare_revision(self.engine,folder);self.assertIsNone(previous)
            self.assertEqual(fake.PARSER_VERSION,'webx-test@'+first)
            again,previous=prepare_revision(self.engine,folder);self.assertEqual((again,previous),(first,first))
            self.engine.write_text(ENGINE+'\n# updated parser\n')
            changed,previous=prepare_revision(self.engine,folder)
            self.assertNotEqual(changed,first);self.assertEqual(previous,first)
            self.assertEqual(fake.PARSER_VERSION,'webx-test@'+changed)
            self.assertEqual(raw.read_bytes(),b'saved response')
    def test_installer_is_additive_and_backed_up(self):
        source="# CHOTUAPP_DURABLE_RUNS_V1\n# CHOTUAPP_CONTROLLED_FAST_V1\nif __name__=='__main__':\n    pass\n"
        self.engine.write_text(source);before=self.engine.read_bytes();(self.root/'saved_runs'/'keep.txt').write_text('keep')
        install(self.root);after=self.engine.read_text();self.assertEqual(after.replace(HOOK,''),source)
        self.assertEqual(transform(after),after);self.assertEqual(next((self.root/'backups').glob('*/main.py')).read_bytes(),before)
        self.assertEqual((self.root/'saved_runs'/'keep.txt').read_text(),'keep')


@unittest.skipUnless(__import__('os').environ.get('DISPLAY'),'GUI integration requires a virtual display')
class ParallelGuiTests(unittest.TestCase):
    def test_existing_buttons_launch_parallel_children_and_resume_queue(self):
        import importlib.util,os,requests
        from unittest.mock import patch
        import install_fast_fetch,remove_daily_auto
        repo=Path(__file__).resolve().parents[1];factory=requests.Session
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);program=root/'app';program.mkdir();desktop=root/'desktop';desktop.mkdir()
            for name in ('resume_support.py','fast_fetch_support.py','manual_mode_support.py','parallel_store_support.py','store_process_worker.py','complete_description_support.py'):
                (program/name).write_bytes((repo/name).read_bytes())
            source=transform(install_fast_fetch.transform(remove_daily_auto.transform((repo/'main.py').read_text())))
            (program/'main.py').write_text(source)
            fake=root/'fake_main.py';fake.write_text(ENGINE)
            template=root/'template.xlsx';Workbook().save(template)
            spec=importlib.util.spec_from_file_location('parallel_gui_fixture',program/'main.py');module=importlib.util.module_from_spec(spec)
            window=None
            try:
                with patch.dict(os.environ,{'CHOTUAPP_DATA_DIR':str(desktop)}):spec.loader.exec_module(module)
                window=module.tk.Tk();errors=[];window.report_callback_exception=lambda *args:errors.append(args)
                app=module.App(window);window.update()
                chosen=[{'name':'First','url':'https://first.example','steps':20},{'name':'Second','url':'https://second.example','steps':20}]
                import parallel_store_support
                original=parallel_store_support.StoreScheduler
                def scheduler(engine,*args,**kw):return original(fake,*args,**kw)
                with patch.object(parallel_store_support,'StoreScheduler',side_effect=scheduler):
                    app._launch_jobs(chosen,template)
                    deadline=time.monotonic()+8
                    while app.running and time.monotonic()<deadline:window.update();time.sleep(.02)
                self.assertFalse(app.running);self.assertEqual(app.saved_runs.load()['status'],'complete')
                self.assertEqual(app.job_completed,2);self.assertEqual(app.job_expected,2);self.assertFalse(app.job_failed)
                self.assertTrue((desktop/'ANAS_DATA_FETCH'/'First.xlsx').is_file());self.assertTrue((desktop/'ANAS_DATA_FETCH'/'Second.xlsx').is_file())
                deadline=time.monotonic()+2
                while app.fresh_run_button.cget('state')!='normal' and time.monotonic()<deadline:window.update();time.sleep(.02)
                self.assertEqual(app.fresh_run_button.cget('state'),'normal');self.assertEqual(app.continue_run_button.cget('state'),'disabled')
                self.assertFalse(errors)
            finally:
                if window is not None:window.destroy()
                requests.Session=factory

if __name__=='__main__':unittest.main()
