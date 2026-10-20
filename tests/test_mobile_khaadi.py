import hashlib,importlib.util,json,tempfile,unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock,patch
from openpyxl import Workbook,load_workbook
import mobile_khaadi_ingest as ingest
from rule804_pipeline import process
ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('khaadi_mobile',ROOT/'catalog-tools/khaadi_mobile.py');mobile=importlib.util.module_from_spec(spec);spec.loader.exec_module(mobile)
class MobileTests(unittest.TestCase):
 def workbook(self,root,url='https://pk.khaadi.com/kurta/ABC123.html'):
  source=root/'input.xlsx';book=Workbook();sheet=book.active;sheet.title='Products'
  sheet.append(['shop_code','product_id','Category','name','description','base_price (PKR)','source_url','subcategory_source','subcategory_evidence','subcategory_checked_at','formula'])
  sheet.append([None,'P1',None,'Cotton Kurta','Cotton kurta. Size: XS. Price: 6000 PKR.',6000,url,None,None,None,'=100+50'])
  for name in ('Product Images','Option Groups','Options','Variant Combinations'):
   s=book.create_sheet(name);s.append(['product_id','value']);s.append(['P1','original'])
  book.save(source);return source
 def prepare(self,root,url=None):
  for p in ('cloud/rule804_pipeline.py','subcategory_inference_support.py','catalog-tools/complete_rule804_batch.py','cloud/confirmed-shop-codes.json','cloud/mobile_khaadi_ingest.py'):
   dest=root/p;dest.parent.mkdir(parents=True,exist_ok=True);dest.write_bytes((ROOT/p).read_bytes())
  source=self.workbook(root,**({'url':url} if url else {}));sha=hashlib.sha256(source.read_bytes()).hexdigest();path='catalog-inputs/khaadi-mobile/'+sha[:16]+'/Khaadi.xlsx';out=root/path;out.parent.mkdir(parents=True);out.write_bytes(source.read_bytes())
  marker=root/ingest.MARKER;marker.write_text(json.dumps({'store':'Khaadi','input_path':path,'sha256':sha}));return out,marker
 def test_rule804_precedes_upload_preserves_formula_and_separates_count(self):
  with tempfile.TemporaryDirectory() as tmp:
   root=Path(tmp);self.prepare(root);state={'uploaded':{'Other':{'sha256':'old'}}};seen=[]
   def rc(*args):
    seen.append(args[0]);book=load_workbook(Path(args[1])/'Khaadi.xlsx');self.assertEqual(book['Products']['K2'].value,'=100+50');self.assertTrue(book['Products']['C2'].value);self.assertEqual(book['Variant Combinations']['B2'].value,'original');book.close()
   self.assertTrue(ingest.upload(state,'target',rc,root=root));self.assertEqual(seen,['copy','check']);self.assertEqual(list(state['uploaded']),['Other']);self.assertIn('Khaadi',state['original_uploads'])
   self.assertFalse(ingest.upload(state,'target',Mock(side_effect=AssertionError),root=root))
 def test_failed_drive_check_does_not_mark_complete(self):
  with tempfile.TemporaryDirectory() as tmp:
   root=Path(tmp);self.prepare(root);state={}
   def rc(*args):
    if args[0]=='check':raise RuntimeError('Drive unavailable')
   with self.assertRaises(RuntimeError):ingest.upload(state,'target',rc,root=root)
   self.assertNotIn('Khaadi',state['original_uploads'])
 def test_foreign_product_source_never_uploads(self):
  with tempfile.TemporaryDirectory() as tmp:
   root=Path(tmp);self.prepare(root,url='https://us.khaadi.com/kurta/ABC123.html');rc=Mock()
   with self.assertRaises(ValueError):ingest.upload({},'target',rc,root=root)
   rc.assert_not_called()
 def test_marker_cannot_select_other_store_or_modified_workbook(self):
  with tempfile.TemporaryDirectory() as tmp:
   root=Path(tmp);source,marker=self.prepare(root);data=json.loads(marker.read_text());data['input_path']='../../secrets.xlsx';marker.write_text(json.dumps(data))
   with self.assertRaises(ValueError):ingest.input_file(root)
   data['input_path']=str(source.relative_to(root));marker.write_text(json.dumps(data));source.write_bytes(b'changed')
   with self.assertRaises(ValueError):ingest.input_file(root)
 def test_default_fresh_pipeline_still_excludes_originals(self):
  for name in ('Khaadi','LAAM','Sapphire'):
   with self.assertRaises(ValueError):process(Path('unused'),Path('unused'),name,{})
 def test_curl_keeps_tls_verification_and_final_host(self):
  def command(args,**kwargs):
   self.assertNotIn('--insecure',args);Path(args[args.index('--output')+1]).write_bytes(b'public');return SimpleNamespace(returncode=0,stdout='200\nhttps://us.khaadi.com/')
  with patch.object(mobile.subprocess,'run',side_effect=command):
   response=mobile.CurlSession().get('https://pk.khaadi.com/');self.assertEqual(response.url,'https://us.khaadi.com/');self.assertEqual(response.content,b'public')
 def test_phone_completion_uploads_content_then_marker_without_leaking_auth(self):
  with tempfile.TemporaryDirectory() as tmp:
   root=Path(tmp);files=root/'files';files.mkdir();file=files/'Khaadi.xlsx';file.write_bytes(b'workbook');(root/'local-status.json').write_text(json.dumps({'status':'exported_review_required','files':[str(file)]}));calls=[]
   def api(endpoint,method=None,payload=None,check=True):
    calls.append((endpoint,method,payload));return None
   with patch.object(mobile,'DATA',root),patch.object(mobile,'api',side_effect=api),patch('builtins.print'):mobile.upload_completed()
   puts=[x for x in calls if x[1]=='PUT'];self.assertEqual(len(puts),2);self.assertTrue(puts[0][0].endswith('/Khaadi.xlsx'));self.assertTrue(puts[1][0].endswith('/latest.json'));self.assertEqual(puts[1][2]['branch'],'main')
