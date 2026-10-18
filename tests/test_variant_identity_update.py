import copy,importlib.util,sys,tempfile,time,unittest,json
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook,load_workbook
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT));sys.path.insert(0,str(ROOT/'cloud'))
spec=importlib.util.spec_from_file_location('variant_update_main',ROOT/'main.py');main=importlib.util.module_from_spec(spec);spec.loader.exec_module(main)
class IdentityTests(unittest.TestCase):
 def row(self,pid='P1',vid='1',option='S'):
  return {'combination_id':main.uid('CC','https://zellbury.com|'+vid),'product_id':pid,'source_variant_id':vid,'option_1_group_id':'G'+pid,'option_1_value':option,'base_price (PKR)':1499.95,'available (yes/no)':'no','variant_image_url':'https://example.com/image.jpg','stock_quantity (agar public ho)':'','source_url':'https://zellbury.com/products/'+pid}
 def test_unambiguous_existing_ids_and_every_field_unchanged(self):
  rows=[self.row(),self.row('P2','2')];before=copy.deepcopy(rows);result,findings=main.normalize_variant_combinations('https://zellbury.com',rows,lambda _:None);self.assertEqual(result,before);self.assertEqual(rows,before);self.assertEqual(findings,[])
 def test_shared_source_id_scoped_by_product_without_data_loss(self):
  rows=[self.row(),self.row('P2')];result,findings=main.normalize_variant_combinations('https://zellbury.com',rows,lambda _:None)
  self.assertEqual(len(result),2);self.assertEqual(len({r['combination_id'] for r in result}),2)
  for old,new in zip(rows,result):self.assertEqual({k:v for k,v in old.items() if k!='combination_id'},{k:v for k,v in new.items() if k!='combination_id'})
  reverse,_=main.normalize_variant_combinations('https://zellbury.com',list(reversed(rows)),lambda _:None);self.assertEqual({r['product_id']:r['combination_id'] for r in reverse},{r['product_id']:r['combination_id'] for r in result});self.assertTrue(findings)
 def test_exact_repeated_record_collapsed_once(self):
  row=self.row();result,findings=main.normalize_variant_combinations('https://zellbury.com',[row,dict(row)],lambda _:None);self.assertEqual(result,[row]);self.assertEqual(findings[0][0],'EXACT_DUPLICATE_COLLAPSED')
 def test_distinct_observed_options_preserved_with_review(self):
  result,findings=main.normalize_variant_combinations('https://zellbury.com',[self.row(option='S'),self.row(option='M')],lambda _:None);self.assertEqual({r['option_1_value'] for r in result},{'S','M'});self.assertEqual(len({r['combination_id'] for r in result}),2);self.assertEqual(len(findings),2)
 def test_ambiguous_price_and_availability_conflicts_are_not_guessed(self):
  for change in ({'base_price (PKR)':99},{'available (yes/no)':'yes'},{'variant_image_url':'other'}):
   with self.assertRaisesRegex(RuntimeError,'Conflicting duplicate'):main.normalize_variant_combinations('https://zellbury.com',[self.row(),dict(self.row(),**change)],lambda _:None)
 def test_actual_zellbury_reused_variant_ids_keep_each_parent_and_source_fields(self):
  fixtures=json.loads((ROOT/'tests/fixtures/zellbury_public_identity_collisions.json').read_text());rows=[];seen=set()
  for pair in fixtures:
   for record in pair.values():
    v=record['variant'];key=(record['product_id'],str(v['id']))
    if key in seen:continue
    seen.add(key);row=self.row(record['product_id'],str(v['id']),v.get('option1') or '')
    row['option_2_value']=v.get('option2') or '';row['source_url']='https://zellbury.com/products/'+record['handle'];row['available (yes/no)']='yes' if v.get('available') else 'no';rows.append(row)
  result,findings=main.normalize_variant_combinations('https://zellbury.com',rows,lambda _:None)
  self.assertEqual(len(rows),len(result));self.assertEqual(len({r['combination_id'] for r in result}),len(result));self.assertTrue(findings)
  key=lambda r:(r['product_id'],r['source_variant_id'])
  self.assertEqual({key(r):{k:v for k,v in r.items() if k!='combination_id'} for r in rows},{key(r):{k:v for k,v in r.items() if k!='combination_id'} for r in result})
 def test_missing_optional_price_api_does_not_hide_a_live_product(self):
  from types import SimpleNamespace
  error=RuntimeError('Website endpoint unavailable (404): https://example.com/products/a.js')
  for status,expected in ((200,False),(404,True),(410,True),(403,False),(429,False)):
   with patch.object(main,'S',SimpleNamespace(get=lambda *args,**kwargs:SimpleNamespace(status_code=status,url='https://example.com/products/a'))):
    self.assertEqual(main.confirmed_removed_product('https://example.com',{'handle':'a'},error,lambda _:None),expected)
  with patch.object(main,'S',SimpleNamespace(get=lambda *args,**kwargs:SimpleNamespace(status_code=404,url='https://other.example/products/a'))):
   self.assertFalse(main.confirmed_removed_product('https://example.com',{'handle':'a'},error,lambda _:None))
 def test_bulk_hyperlinks_use_append_index_not_quadratic_max_row_scan(self):
  w=Workbook();s=w.active;s.title='Variant Combinations';s.append(['combination_id','base_price (PKR)','variant_image_url']);s._chotu_export_headers=tuple(c.value for c in s[1])
  with patch.object(type(s),'max_row',new=property(lambda _: (_ for _ in ()).throw(AssertionError('max_row scan')))):
   for i in range(1500):main.write_row(s,dict(self.row(),combination_id=str(i)))
  self.assertEqual(s.max_row,1501);self.assertEqual(s.cell(1501,2).value,1499.95);self.assertEqual(s.cell(1501,3).hyperlink.target,'https://example.com/image.jpg')
  with tempfile.TemporaryDirectory() as tmp:
   p=Path(tmp)/'variants.xlsx';w.save(p);v=load_workbook(p);self.assertEqual(v.active.max_row,1501);self.assertEqual(v.active.cell(1501,3).hyperlink.target,'https://example.com/image.jpg');v.close()
 def test_telemetry_tracks_local_build_and_write_progress(self):
  from diagnostics import Collector
  with tempfile.TemporaryDirectory() as tmp:
   c=Collector(Path(tmp)/'telemetry.json');c.observe('[Store] ROW BUILD: 100/1000 products processed; 55 export rows prepared');self.assertEqual(c.stores['Store']['task_progress']['done'],100)
   c.observe('[Store] STAGE: Writing Excel rows');c.observe('[Store] WORKSHEET WRITE: 1000/25000 variants written');self.assertEqual(c.stores['Store']['task_progress']['stage'],'writing variant rows')
if __name__=='__main__':unittest.main()
