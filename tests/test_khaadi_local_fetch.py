import importlib.util,json,sys,tempfile,unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock,patch
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
spec=importlib.util.spec_from_file_location('local_fetch',ROOT/'catalog-tools/start_khaadi_local_fetch_windows.py');helper=importlib.util.module_from_spec(spec);spec.loader.exec_module(helper)
import khaadi_source_support as source
class LocalFetchTests(unittest.TestCase):
 def test_local_export_writes_template_variants_and_status_without_network(self):
  from bs4 import BeautifulSoup
  from openpyxl import load_workbook
  fixture=ROOT/'tests/fixtures/khaadi_public';soup=BeautifulSoup((fixture/'master.html').read_text(),'html.parser')
  for n in soup.select('.set-item')[0].select('[data-attr="size"] input'):
   if n['data-attr-value']=='00S':n.parent.decompose()
  url=source.BASE+'/black-cambric-kurta/1-26-320-A-C1-VG_MULTI.html'
  product=source.normalize_product(str(soup),url,lambda url:SimpleNamespace(text=(fixture/('xs.html' if '=0XS' in url else 'l.html')).read_text()))
  with tempfile.TemporaryDirectory() as tmp:
   data=Path(tmp)
   with patch.object(source.PublicCatalog,'catalog',return_value=([product],False)),patch('builtins.print'):
    self.assertEqual(helper.fetch(ROOT,data,Mock()),0)
   state=json.loads((data/'local-status.json').read_text());self.assertEqual(state['status'],'exported_review_required');self.assertFalse(state['full_catalog_verified'])
   wb=load_workbook(state['files'][0]);self.assertEqual(wb['Products'].max_row,2);self.assertEqual(wb['Variant Combinations'].max_row,3);wb.close()
   self.assertTrue((data/'pages.sqlite3').is_file());self.assertTrue((data/'fetch-log.txt').is_file())
 def test_fresh_native_curl_response_is_cache_compatible(self):
  from resume_support import ResponseCache
  url=source.BASE+'/robots.txt'
  raw=SimpleNamespace(get=Mock(return_value=SimpleNamespace(url=url,status_code=200,content=b'User-agent: *',text='User-agent: *')))
  def catalog(reader):
   response=reader.read(url,fresh=True)
   self.assertEqual(response.encoding,'utf-8')
   raise RuntimeError('Transport checked')
  with tempfile.TemporaryDirectory() as tmp:
   data=Path(tmp)
   with patch.object(source.PublicCatalog,'catalog',catalog),patch('builtins.print'):
    self.assertEqual(helper.fetch(ROOT,data,raw),2)
   state=json.loads((data/'local-status.json').read_text());self.assertEqual(state['error'],'Transport checked')
   cache=ResponseCache(data/'pages.sqlite3')
   try:self.assertEqual(cache.get(url).text,'User-agent: *')
   finally:cache.close()
 def test_failure_keeps_resume_status_and_cached_work(self):
  with tempfile.TemporaryDirectory() as tmp:
   data=Path(tmp)
   with patch.object(source.PublicCatalog,'catalog',side_effect=RuntimeError('Source unavailable')),patch('builtins.print'):
    self.assertEqual(helper.fetch(ROOT,data,Mock()),2)
   self.assertEqual(json.loads((data/'local-status.json').read_text())['status'],'failed_resume_available');self.assertTrue((data/'pages.sqlite3').exists())
 def test_prepared_sources_match_pinned_hashes(self):
  import hashlib
  for name,expected in helper.FILES.items():self.assertEqual(hashlib.sha256((ROOT/name).read_bytes()).hexdigest(),expected,name)
