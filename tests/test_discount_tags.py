from pathlib import Path
import importlib.util
import sys
import tempfile
import unittest
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
spec=importlib.util.spec_from_file_location('tags_main',ROOT/'main.py');main=importlib.util.module_from_spec(spec);spec.loader.exec_module(main)
import install_discount_tags as patch

class TagsTests(unittest.TestCase):
    def test_actual_template_export_roundtrip_and_discounted_variants(self):
        w=load_workbook(ROOT/'importer-files/Chotu_Products_Import_Template_Discount_Tags.xlsx')
        for title in ('Products','Variant Combinations'):
            for base,sale,percentage,tag in [(1490,None,None,None),(2980,2680,10.07,'10% Off'),(4470,3845,13.98,'14% Off'),(100,100,None,None),(100,110,None,None),(100,0,100,'100% Off')]:
                row={'product_id':'fixture','base_price (PKR)':base,'sale_price (PKR, agar discount ho)':sale}
                main.write_row(w[title],row)
                self.assertEqual(row,{'product_id':'fixture','base_price (PKR)':base,'sale_price (PKR, agar discount ho)':sale})
        with tempfile.TemporaryDirectory() as tmp:
            file=Path(tmp)/'export.xlsx';w.save(file);saved=load_workbook(file,data_only=True)
            for title in ('Products','Variant Combinations'):
                it=iter(saved[title].values);h=next(it);rows=[dict(zip(h,r)) for r in it]
                self.assertEqual(rows[1]['discount_percentage (%)'],10.07)
                self.assertEqual(rows[2]['discount_tag'],'14% Off')
                self.assertEqual(rows[2]['discount_tag_source'],'calculated_from_prices')
                self.assertIsNone(rows[2]['source_discount_tag'])
                self.assertEqual(rows[2]['base_price (PKR)'],4470)
                self.assertEqual(rows[2]['sale_price (PKR, agar discount ho)'],3845)
                self.assertIsNone(rows[0]['discount_tag'])
            self.assertEqual(rows[-1]['discount_tag'],'100% Off')

    def test_surgical_patch_preserves_other_updates_and_old_templates(self):
        source=(ROOT/'main.py').read_text()
        for old,new in patch.REPLACEMENTS:source=source.replace(new,old,1)
        source+='\n# saved runs and UI fix\n# adaptive_pace 0.35\n'
        updated,count=patch.transform(source.encode());expected=source
        for old,new in patch.REPLACEMENTS:expected=expected.replace(old,new,1)
        self.assertEqual(updated.decode(),expected);self.assertEqual(count,3)
        self.assertEqual(patch.transform(updated),(updated,0))
