import copy,json,tempfile,unittest
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import urlsplit,parse_qs
from unittest.mock import patch
from magento_source_support import normalize,PublicCatalog,BASE
from woocommerce_source_support import offer,variants,export

class MagentoTests(unittest.TestCase):
 def product(self):
  return dict(__typename='SimpleProduct',id=10,sku='USED-1',name='Lenovo Used Laptop',url_key='lenovo-used-laptop',url_suffix='.html',stock_status='IN_STOCK',price_range=dict(minimum_price=dict(regular_price=dict(value=59500,currency='PKR'),final_price=dict(value=58500,currency='PKR'))),description=dict(html='<p>Processor: Core i5</p><p>Memory: 16 GB</p>'),short_description=dict(html='<p>SSD: 256 GB</p>'),image=dict(url=BASE+'/media/laptop.jpg'),media_gallery=[],categories=[dict(name='Used Laptops')],options=[])
 def test_virtual_product_keeps_explicit_offer_without_physical_or_stock_inference(self):
  raw=self.product();raw.update(__typename='VirtualProduct',name='Software Licence',sku='LICENCE-1');p=normalize(raw)
  self.assertTrue(p['virtual']);self.assertEqual(p['magento_product_type'],'VirtualProduct');self.assertEqual(p['sku'],'LICENCE-1');self.assertEqual(offer(p)[:2],(59500,58500));self.assertNotIn('stock_quantity',p)
  raw['stock_status']='OUT_OF_STOCK';self.assertFalse(normalize(raw)['is_in_stock'])
  raw['stock_status']=None
  with self.assertRaises(ValueError):normalize(raw)
 def test_unknown_magento_product_kinds_still_fail_closed(self):
  for kind in ['BundleProduct','DownloadableProduct',None]:
   raw=self.product();raw['__typename']=kind
   with self.subTest(kind=kind),self.assertRaisesRegex(ValueError,'Unsupported Magento type'):normalize(raw)
 def test_missing_sku_index_uses_exact_url_key_and_keeps_offer_identity(self):
  from unittest.mock import Mock
  raw=self.product();reader=PublicCatalog({},BASE,lambda _:None);reader.byid[10]=copy.deepcopy(raw);reader.robot=SimpleNamespace(can_fetch=lambda *a:False)
  reader.query=Mock(side_effect=[{'items':[]},{'items':[raw]}])
  product=reader.read('/wp-json/wc/store/v1/products/10')
  self.assertEqual(product['id'],10);self.assertEqual(product['sku'],'USED-1');self.assertEqual(offer(product)[1],58500);self.assertIn('url_key:{eq:"lenovo-used-laptop"}',reader.query.call_args_list[1].args[0])
 def test_url_key_fallback_cannot_accept_different_sku_or_product(self):
  from unittest.mock import Mock
  for field,value in [('id',11),('sku','OTHER'),('url_key','other')]:
   raw=self.product();reader=PublicCatalog({},BASE,lambda _:None);reader.byid[10]=copy.deepcopy(raw);reader.robot=SimpleNamespace(can_fetch=lambda *a:False);wrong=dict(raw,**{field:value});reader.query=Mock(side_effect=[{'items':[]},{'items':[wrong]}])
   with self.subTest(field=field),self.assertRaisesRegex(ValueError,'identity mismatch'):reader.read('/wp-json/wc/store/v1/products/10')
 def test_exact_pkr_prices_and_used_condition_without_quantity_guess(self):
  p=normalize(self.product());self.assertEqual(offer(p)[0],59500);self.assertEqual(offer(p)[1],58500);self.assertEqual(p['condition'],'used');self.assertEqual(p['permalink'],BASE+'/lenovo-used-laptop.html');self.assertNotIn('stock_quantity',p)
 def test_invalid_currency_stock_url_and_custom_options_fail_closed(self):
  for kind in ['currency','stock','url','options']:
   raw=self.product()
   if kind=='currency':raw['price_range']['minimum_price']['final_price']['currency']='USD'
   if kind=='stock':raw['stock_status']=None
   if kind=='url':raw.pop('url_suffix')
   if kind=='options':raw['options']=[dict(title='Choose memory',required=True)]
   with self.subTest(kind=kind),self.assertRaises(ValueError):normalize(raw)
 def test_configurable_products_keep_only_actual_available_and_sold_out_selections(self):
  raw=self.product();raw.update(__typename='ConfigurableProduct',configurable_options=[dict(attribute_code='color',label='Colour',values=[dict(value_index=1,label='Black'),dict(value_index=2,label='Blue')])],variants=[])
  children={}
  for ident,value,stock in [(11,1,'IN_STOCK'),(12,2,'OUT_OF_STOCK')]:
   child=self.product();child.update(id=ident,sku=str(ident),stock_status=stock);raw['variants'].append(dict(attributes=[dict(code='color',label='Black' if value==1 else 'Blue',value_index=value)],product=child));children[ident]=normalize(child,parent=10,permalink=BASE+'/x.html')
  parent=normalize(raw);actual=variants(parent,lambda path:children[int(path.rsplit('/',1)[1])]);self.assertEqual(len(actual),2);self.assertEqual([a for _,a in actual],[{'Colour':'Black'},{'Colour':'Blue'}]);self.assertEqual([offer(v)[2] for v,_ in actual],[True,False])
  raw['variants'][0]['attributes'][0]['label']='Red'
  with self.assertRaises(ValueError):normalize(raw)
 def test_real_workbook_export_uses_verified_detail_and_preserves_specifications(self):
  from test_app import main
  from openpyxl import load_workbook
  raw=self.product();queries=[]
  def read(url):
   query=parse_qs(urlsplit(url).query)['query'][0];queries.append(query)
   if 'filter:' in query:return {'data':{'products':{'items':[copy.deepcopy(raw)]}}}
   return {'data':{'products':{'total_count':1,'items':[copy.deepcopy(raw)]}}}
  response=SimpleNamespace(url=BASE+'/robots.txt',text='User-agent: *\nAllow: /',raise_for_status=lambda:None)
  page=SimpleNamespace(status_code=200,url=BASE+'/lenovo-used-laptop.html',text='<link rel="canonical" href="'+BASE+'/lenovo-used-laptop.html"><h1>Lenovo Used Laptop</h1><table id="product-attribute-specs-table"><tr><th>Graphics</th><td>Intel UHD 620</td></tr></table>')
  session=SimpleNamespace(get=lambda url,**kwargs:response if url.endswith('/robots.txt') else page)
  with tempfile.TemporaryDirectory() as tmp:
   g=dict(vars(main),OUT=Path(tmp),S=session,get_json=read,_chotu_wait=lambda:None)
   provider=PublicCatalog(g,BASE,lambda _:None)
   export(g,BASE,'Paklap',Path(__file__).resolve().parents[1]/'importer-files/Chotu_Products_Import_Template_Discount_Tags.xlsx',lambda _:None,provider=provider)
   book=load_workbook(Path(tmp)/'Paklap - REFRESH REVIEW.xlsx',read_only=True)
   try:
    it=book['Products'].iter_rows(values_only=True);headers=next(it);row=dict(zip(headers,next(it)))
    self.assertEqual(row['name'],'Lenovo Used Laptop');self.assertEqual(row['base_price (PKR)'],59500);self.assertEqual(row['sale_price (PKR, agar discount ho)'],58500);self.assertEqual(row['condition (new/used)'],'used');self.assertIn('16 GB',row['description']);self.assertIn('256 GB',row['description']);self.assertIn('Intel UHD 620',row['description']);self.assertIn('unverified',str(row.get('website_order_status')));self.assertEqual(book['Variant Combinations'].max_row,2)
   finally:book.close()
   self.assertEqual(len(queries),2);self.assertIn('sku:{eq:"USED-1"}',queries[1])
 def test_catalog_duplicates_and_incomplete_pagination_are_not_completed(self):
  g=dict(S=None,get_json=lambda _:None,_chotu_wait=lambda:None);catalog=PublicCatalog(g,BASE,lambda _:None)
  with patch.object(catalog,'query',return_value=dict(total_count=2,items=[self.product(),self.product()])):
   with self.assertRaises(ValueError):catalog.catalog()


class CustomAddOnTests(unittest.TestCase):
 product=MagentoTests.product
 def option(self):
  return dict(__typename='CustomizableDropDownOption',option_id=7,title='Memory',required=True,value=[dict(option_type_id=71,title='8 GB',price=0,price_type='FIXED'),dict(option_type_id=72,title='16 GB',price=5000,price_type='FIXED')])
 def test_native_custom_prices_do_not_create_fake_sku_combinations(self):
  raw=self.product();raw['options']=[self.option()];p=normalize(raw)
  self.assertEqual([str(c['extra_price']) for c in p['custom_options'][0]['choices']],['0','5000'])
  self.assertEqual(len(variants(p,lambda _:None)),1)
  self.assertEqual(variants(p,lambda _:None)[0][1],{})
  self.assertEqual(p['custom_options'][0]['required'],True)
 def test_percent_and_unpriced_custom_choices_remain_unresolved(self):
  from magento_source_support import custom_options
  for changes in ({'price_type':'PERCENT'},{'price':-1}):
   option=self.option();option['value'][1].update(changes)
   with self.assertRaises(ValueError):custom_options([option])

 def test_custom_addon_workbook_retains_prices_and_one_real_sku(self):
  from test_app import main
  from openpyxl import load_workbook
  raw=self.product();raw['options']=[self.option()]
  def read(url):
   query=parse_qs(urlsplit(url).query)['query'][0]
   return {'data':{'products':{'items':[copy.deepcopy(raw)]}}} if 'filter:' in query else {'data':{'products':{'total_count':1,'items':[copy.deepcopy(raw)]}}}
  robots=SimpleNamespace(url=BASE+'/robots.txt',text='User-agent: *\nAllow: /',raise_for_status=lambda:None)
  session=SimpleNamespace(get=lambda *a,**k:robots)
  with tempfile.TemporaryDirectory() as tmp:
   g=dict(vars(main),OUT=Path(tmp),S=session,get_json=read,_chotu_wait=lambda:None)
   export(g,BASE,'Paklap',Path(__file__).resolve().parents[1]/'importer-files/Chotu_Products_Import_Template_Discount_Tags.xlsx',lambda _:None,provider=PublicCatalog(g,BASE,lambda _:None))
   file=Path(tmp)/'Paklap - REFRESH REVIEW.xlsx';book=load_workbook(file,read_only=True)
   try:
    it=book['Options'].iter_rows(values_only=True);headers=next(it);choices=[dict(zip(headers,r)) for r in it]
    self.assertEqual([r['option_name'] for r in choices],['8 GB','16 GB']);self.assertEqual([r['extra_price (PKR)'] for r in choices],[0,5000]);self.assertEqual(book['Option Groups'].max_row,2);self.assertEqual(book['Variant Combinations'].max_row,2)
   finally:book.close()
   from price_stock_gate import workbook_issues
   self.assertEqual(workbook_issues(file),{})

if __name__=='__main__':unittest.main()
