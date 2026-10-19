import importlib.util,tempfile,unittest,zipfile
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import PatternFill
spec=importlib.util.spec_from_file_location('complete_rule804_batch',Path(__file__).resolve().parents[1]/'catalog-tools/complete_rule804_batch.py');module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
class CompletionTests(unittest.TestCase):
 def test_explicit_types_do_not_turn_care_perfume_into_fragrance(self):
  label=module.extra_category({'name':'Bindiya-40 (Pearl)','description':'Keep away from perfume'})
  self.assertEqual(label[0],'Forehead Jewellery')
  self.assertIsNone(module.extra_category({'name':'Two of a Kind Deal','description':'Two of a Kind Deal.'}))
 def test_confirmed_product_type_and_stitched_set(self):
  self.assertEqual(module.extra_category({'name':'High Waisted Super Wide Leg'}, {'product_type':'JEANS'})[0],'Jeans')
  self.assertEqual(module.extra_category({'name':'CODE','brand':'Women Stitched 3 Piece'})[0],'Stitched Outfit Sets')
 def test_patches_preserve_other_xml_formulas_and_styles(self):
  with tempfile.TemporaryDirectory() as tmp:
   source=Path(tmp)/'before.xlsx';out=Path(tmp)/'after.xlsx';w=Workbook();s=w.active;s.title='Products';s.append(['shop_code','product_id','Category','description','price']);s.append(['SE1','P1',None,'Details\n10 g', '=100+50']);s['C2'].fill=PatternFill('solid',fgColor='FFEE00');other=w.create_sheet('Variant Combinations');other.append(['P1',150]);w.save(source)
   changes={'C2':'Forehead Jewellery'};part=module.patch_sheet(source,out,changes);module.verify(source,out,part,changes);result=load_workbook(out)
   self.assertEqual(result['Products']['E2'].value,'=100+50');self.assertEqual(result['Products']['C2'].fill.fgColor.rgb,'00FFEE00')
if __name__=='__main__':unittest.main()
