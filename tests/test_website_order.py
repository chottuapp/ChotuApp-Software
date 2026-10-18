import copy,sys,tempfile,unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from bs4 import BeautifulSoup
from openpyxl import Workbook,load_workbook
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
import website_order_support as order
import ranking_support
BASE='https://store.example'
HOME='<header><nav><ul><li><a href="/collections/skin">Skin</a></li><li><a href="/collections/makeup">Makeup</a><ul><li><a href="/collections/lips">Lip Colour</a></li></ul></li></ul></nav></header><footer><a href="/collections/random">Ignore footer</a></footer>'
def page(keys,next_url=''):
 return '<main><ul id="product-grid">'+''.join('<li><a href="/products/'+k+'">'+k+'</a><a href="/products/'+k+'">Image</a></li>' for k in keys)+'</ul>'+('<a rel="next" href="'+next_url+'">Next</a>' if next_url else '')+'</main>'
class WebsiteOrderTests(unittest.TestCase):
 def fact(self,cat=1,sub=1,product=1,label='Skin'):
  return {'category':label,'subcategory':'','category_order':cat,'subcategory_order':sub,'product_order':product,'source_url':BASE+'/collections/skin','checked_at':'2026-10-18T00:00:00Z'}
 def row(self,key,price):return {'product_id':key,'name':key,'category':'Old category','source_url':BASE+'/products/'+key,'base_price (PKR)':price,'sku':'SKU-'+key,'analysis_note':'Manual note'}
 def test_navigation_order_and_exact_nested_labels_exclude_footer(self):
  nav=order.navigation(BeautifulSoup(HOME,'html.parser'),BASE)
  self.assertEqual([n['url'] for n in nav],[BASE+'/collections/skin',BASE+'/collections/makeup',BASE+'/collections/lips'])
  self.assertEqual((nav[2]['category'],nav[2]['subcategory']),('Makeup','Lip Colour'));self.assertEqual(nav[0]['category_order'],1)
 def test_grid_order_ignores_prices_duplicates_recommendations_and_nav(self):
  html='<header><a href="/products/nav">No</a></header>'+page(['expensive','cheap'])+'<main><div class="related"><div class="product-grid"><a href="/products/related">No</a></div></div></main>'
  self.assertEqual(order.grid_products(BeautifulSoup(html,'html.parser'),BASE),['expensive','cheap'])
 def test_website_order_overrides_cheapest_bestseller_and_manual_priorities(self):
  rows=[self.row('cheap',1),self.row('expensive',900),self.row('later',0.5)];before=copy.deepcopy(rows)
  for row in rows:row.update(category_priority=1,product_priority=1,popularity_basis='store_bestseller')
  facts={'expensive':self.fact(product=1),'cheap':self.fact(product=2),'later':self.fact(cat=2,label='Makeup')}
  order.apply_order(rows,facts);self.assertEqual([r['product_id'] for r in rows],['expensive','cheap','later'])
  for row in rows:
   original=next(r for r in before if r['product_id']==row['product_id'])
   for key in ('base_price (PKR)','sku','source_url','product_id'):self.assertEqual(row[key],original[key])
   self.assertIn('Manual note',row['analysis_note']);self.assertIn('not verified sales',row['analysis_note'])
 def test_unknown_order_stays_stable_without_price_fallback(self):
  rows=[self.row('high',999),self.row('low',1)];ranking_support.order_rows(rows,['low'],{'low':{'url':BASE,'checked_at':'today'}})
  self.assertEqual([r['product_id'] for r in rows],['high','low']);self.assertTrue(all(r['website_order_status']=='unverified' for r in rows))
 def test_serialized_order_survives_final_ordering_without_capture(self):
  rows=[self.row('cheap',1),self.row('high',200)];order.apply_order(rows,{'high':self.fact(product=1),'cheap':self.fact(product=2)});rows.reverse();order.apply_order(rows)
  self.assertEqual([r['product_id'] for r in rows],['high','cheap']);self.assertEqual(rows[0]['analysis_note'].count('[Website listing order]'),1)
 def capture(self,pages,**options):
  calls=[]
  def get(url,**kwargs):
   calls.append(url);body=pages[url];return SimpleNamespace(status_code=200,url=url,text=body,content=body.encode())
  g={'S':SimpleNamespace(get=get),'_chotu_wait':lambda:None,'_chotu_backoff':lambda *args:None,'_chotu_tls':SimpleNamespace()}
  capture=order.Capture(g,BASE,lambda _:None,**options);return capture,calls
 def test_actual_default_pagination_preserves_first_occurrence_and_no_forced_sort(self):
  pages={BASE+'/robots.txt':'User-agent: *\nAllow: /',BASE+'/':HOME,BASE+'/collections/skin':page(['high','cheap'],'?page=2'),BASE+'/collections/skin?page=2':page(['cheap','next']),BASE+'/collections/makeup':page(['face']),BASE+'/collections/lips':page(['lip'])}
  capture,calls=self.capture(pages);facts=capture.collect();self.assertTrue(capture.complete);self.assertEqual(facts['next']['product_order'],3);self.assertEqual(facts['high']['product_order'],1);self.assertFalse(any('best-selling' in u for u in calls))
 def test_next_page_keeps_sort_only_when_website_default_control_declares_it(self):
  selected='<select name="sort_by"><option selected value="best-selling">Best selling</option></select>'
  pages={BASE+'/robots.txt':'User-agent: *\nAllow: /',BASE+'/':HOME,BASE+'/collections/skin':selected+page(['one'],'?sort_by=best-selling&page=2'),BASE+'/collections/skin?sort_by=best-selling&page=2':page(['two']),BASE+'/collections/makeup':page(['face']),BASE+'/collections/lips':page(['lip'])}
  capture,calls=self.capture(pages);facts=capture.collect();self.assertTrue(capture.complete);self.assertEqual(facts['two']['product_order'],2)
 def test_nuxt_product_cards_follow_observed_dom_order(self):
  html='<main><a href="/products/first"><div data-testid="product-card">First</div></a><a href="/products/second"><div class="product_card">Second</div></a></main>'
  self.assertEqual(order.grid_products(BeautifulSoup(html,'html.parser'),BASE),['first','second'])
 def test_request_bound_retains_only_observed_prefix(self):
  pages={BASE+'/robots.txt':'User-agent: *\nAllow: /',BASE+'/':HOME,BASE+'/collections/skin':page(['one'],'?page=2')}
  capture,calls=self.capture(pages,max_requests=3);facts=capture.collect();self.assertFalse(capture.complete);self.assertEqual(set(facts),{'one'});self.assertEqual(len(calls),3)
 def test_robots_denial_never_requests_collection(self):
  pages={BASE+'/robots.txt':'User-agent: *\nDisallow: /collections/',BASE+'/':HOME};capture,calls=self.capture(pages);self.assertEqual(capture.collect(),{});self.assertFalse(any('/collections/' in u for u in calls))
 def test_cross_host_and_changed_sort_pagination_not_followed(self):
  self.assertEqual(order.public_link(BASE,'https://other.example/collections/x'),'')
  pages={BASE+'/robots.txt':'User-agent: *\nAllow: /',BASE+'/':HOME,BASE+'/collections/skin':page(['one'],'?sort_by=best-selling&page=2'),BASE+'/collections/makeup':page(['two']),BASE+'/collections/lips':page(['three'])}
  capture,calls=self.capture(pages);capture.collect();self.assertFalse(any('best-selling' in u for u in calls));self.assertFalse(capture.complete)


class WebsiteOrderIntegrityTests(unittest.TestCase):
 def test_final_workbook_uses_observed_order_without_changing_related_rows_links_or_formulas(self):
  import main
  with tempfile.TemporaryDirectory() as tmp:
   path=Path(tmp)/'book.xlsx';book=Workbook();sheet=book.active;sheet.title='Products';sheet.append(['product_id','category','source_url','base_price (PKR)','image','formula']);sheet.append(['cheap','Old',BASE+'/products/cheap',1,'https://image/cheap','=1+1']);sheet.append(['high','Old',BASE+'/products/high',999,'https://image/high','=2+2']);sheet['E2'].hyperlink='https://image/cheap';sheet['E3'].hyperlink='https://image/high';book.create_sheet('Variant Combinations').append(['variant','high',123.45]);book.save(path);book.close()
   evidence={'high':WebsiteOrderTests().fact(product=1),'cheap':WebsiteOrderTests().fact(product=2)}
   with patch.object(main._chotu_tls,'website_order_evidence',evidence,create=True):main.finalize_catalog_order(path,Path(tmp)/'ignored.csv',lambda _:None)
   book=load_workbook(path);self.assertEqual(book['Products']['A2'].value,'high');self.assertEqual(book['Products']['D2'].value,999);self.assertEqual(book['Products']['F2'].value,'=2+2');self.assertEqual(book['Products']['E2'].hyperlink.target,'https://image/high');self.assertEqual(book['Products']['E2'].hyperlink.ref,'E2');self.assertEqual(tuple(c.value for c in book['Variant Combinations'][1]),('variant','high',123.45));book.close()
 def test_invalid_or_cross_store_positions_are_not_trusted(self):
  rows=[WebsiteOrderTests().row('one',900),WebsiteOrderTests().row('two',1)]
  rows[1].update(website_category_order=1,website_subcategory_order=1,website_product_order=1,website_order_status='observed_default_listing',website_order_source_url='https://other.example/collections/x',website_order_checked_at='2026-10-18T00:00:00Z')
  order.apply_order(rows);self.assertEqual(rows[0]['product_id'],'one');self.assertEqual(rows[1]['website_order_status'],'unverified')

if __name__=='__main__':unittest.main()
