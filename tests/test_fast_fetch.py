import copy
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import threading
import time
import unittest
from collections import Counter
from http.server import BaseHTTPRequestHandler,ThreadingHTTPServer
from unittest.mock import patch
import requests
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
import fast_fetch_support as fast
import resume_support
import install_speed_patch
import install_fast_fetch
import install_adaptive_speed

class FastFetchTests(unittest.TestCase):
    def test_parallel_and_serial_export_every_cell_and_endpoint_match(self):
        products=[]
        for i in range(6):
            products.append({'id':i+1,'title':'Fixture '+str(i),'handle':'p'+str(i),'vendor':'Fixture','product_type':'Lipstick','body_html':'<p>Source product description for this verified cosmetic.</p>',
                             'options':[{'name':'Pack'}], 'images':[{'id':100+i,'src':'https://cdn.shopify.com/s/files/1/1/1/files/p'+str(i)+'.jpg','position':1}],
                             'variants':[{'id':1000+i,'price':'1490','compare_at_price':'1990','available':True,'option1':'Pack of 1'},
                                         {'id':2000+i,'price':'2680','compare_at_price':'2980','available':False,'option1':'Pack of 2'}]})
        visits=Counter();active=0;peak=0;guard=threading.Lock()
        class Handler(BaseHTTPRequestHandler):
            def do_GET(self):
                nonlocal active,peak
                with guard:visits[self.path]+=1;active+=1;peak=max(peak,active)
                try:
                    time.sleep(.40)
                    if self.path.startswith('/products.json'):data={'products':products}
                    elif self.path.startswith('/collections.json'):data={'collections':[{'handle':'c'+str(i),'title':'Category '+str(i)} for i in range(3)]}
                    elif self.path.startswith('/collections/c'):data={'products':products}
                    elif self.path.startswith('/collections/all'):
                        html='<select><option value="best-selling" selected>Best selling</option></select><ul id="product-grid">'+''.join('<h3><a href="/products/p'+str(i)+'">Fixture</a></h3>' for i in range(5,-1,-1))+'</ul>'
                        self.send_response(200);self.end_headers();self.wfile.write(html.encode());return
                    else:
                        handle=self.path.split('/products/',1)[1].split('?')[0].removesuffix('.js');product=next(p for p in products if p['handle']==handle)
                        if self.path.endswith('.js'):data={'id':product['id'],'variants':[dict(v,price=int(v['price'])*100,compare_at_price=int(v['compare_at_price'])*100) for v in product['variants']]}
                        else:
                            html='<media-gallery><div class="product__media"><img src="'+product['images'][0]['src']+'"></div></media-gallery>'
                            self.send_response(200);self.end_headers();self.wfile.write(html.encode());return
                    self.send_response(200);self.end_headers();self.wfile.write(json.dumps(data).encode())
                finally:
                    with guard:active-=1
            def log_message(self,*args):pass
        server=ThreadingHTTPServer(('127.0.0.1',0),Handler);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
        source,_=install_speed_patch.transform((ROOT/'main.py').read_bytes());outputs=[];times=[];counts=[];peaks=[]
        original_factory=requests.Session
        try:
            with tempfile.TemporaryDirectory() as tmp:
                for parallel in (False,True):
                    folder=Path(tmp)/str(parallel);folder.mkdir();file=folder/'main.py'
                    body=source if parallel else source.replace(install_adaptive_speed.NEW_TEXT.encode(),install_adaptive_speed.OLD_TEXT.encode(),1)
                    file.write_bytes(body)
                    spec=importlib.util.spec_from_file_location('fast_fixture_'+str(parallel),file);module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
                    resume_support.enable_resume(module.__dict__)
                    if parallel:fast.enable_fast_fetch(module.__dict__)
                    module.OUT=folder/'exports';module._chotu_tls.saved_api_path=folder/'api.sqlite3'
                    module._chotu_tls.saved_response_cache=resume_support.ResponseCache(folder/'pages.sqlite3')
                    base=f'http://127.0.0.1:{server.server_port}'
                    module._chotu_open_checkpoint(base,'Fixture');logs=[]
                    visits.clear();peak=0;start=time.monotonic()
                    try:module.run_site(base,'Fixture',str(ROOT/'importer-files/Chotu_Products_Import_Template_Discount_Tags.xlsx'),logs.append)
                    finally:
                        module._chotu_close_checkpoint(False);module._chotu_tls.saved_response_cache.close()
                    times.append(time.monotonic()-start);counts.append(copy.copy(visits));peaks.append(peak)
                    book=load_workbook(module.OUT/'Fixture.xlsx',data_only=True);outputs.append({name:list(book[name].values) for name in book.sheetnames});book.close()
                    requests.Session=original_factory
                    if parallel:self.assertTrue(any('FAST FETCH:' in line for line in logs))
                self.assertEqual(outputs[0],outputs[1]);self.assertEqual(counts[0],counts[1])
                self.assertEqual(peaks[0],1);self.assertEqual(peaks[1],2);self.assertLess(times[1],times[0])
                print('Same HTTP endpoints and every workbook cell: serial %.2fs, parallel %.2fs; max concurrent requests 2.'%tuple(times))
        finally:server.shutdown();server.server_close();thread.join();requests.Session=original_factory

    def test_rate_limit_cooldown_stops_prefetch_and_respects_retry_after(self):
        logs=[];gate=fast.RequestGate(lambda:None,logs.append)
        response=requests.Response();response.status_code=429;response.headers['Retry-After']='60'
        before=time.monotonic();gate.response(response)
        self.assertTrue(gate.limited);self.assertGreaterEqual(gate.cooldown,before+60)
        self.assertGreaterEqual(gate.interval,2.5)
        with self.assertRaises(RuntimeError):gate.wait(preload=True)
        self.assertTrue(logs)

    def test_first_429_stops_retries_and_cached_data_survives_restart(self):
        original_factory=requests.Session
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);file=root/'main.py';source,_=install_speed_patch.transform((ROOT/'main.py').read_bytes());file.write_bytes(source)
            spec=importlib.util.spec_from_file_location('rate_pause_fixture',file);module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
            resume_support.enable_resume(module.__dict__);fast.enable_fast_fetch(module.__dict__)
            module._chotu_tls.saved_api_path=root/'api.sqlite3'
            module._chotu_tls.saved_response_cache=resume_support.ResponseCache(root/'pages.sqlite3')
            module._chotu_open_checkpoint('https://limited.example','Limited')
            session=module.S
            while hasattr(session,'session'):session=session.session
            ok=requests.Response();ok.status_code=200;ok._content=b'{"price":3845}'
            limited=requests.Response();limited.status_code=429;limited.headers['Retry-After']='120';limited._content=b'blocked'
            try:
                with patch.object(session,'get',side_effect=[ok,limited]) as network:
                    self.assertEqual(module.get_json('https://limited.example/saved.json')['price'],3845)
                    started=time.monotonic()
                    with self.assertRaisesRegex(module.ChotuRateLimited,'SOURCE PAUSED'):module.get_json('https://limited.example/blocked.json')
                    self.assertLess(time.monotonic()-started,2)
                    with self.assertRaises(module.ChotuRateLimited):module.get_json('https://limited.example/blocked.json')
                    self.assertEqual(module.get_json('https://limited.example/saved.json')['price'],3845)
                    self.assertEqual(network.call_count,2)
                reloaded=fast.RequestGate(lambda:None,lambda x:None,state_path=root/'source_pacing/limited.example.json')
                self.assertTrue(reloaded.limited);self.assertGreaterEqual(reloaded.interval,1.0)
                with self.assertRaises(RuntimeError):reloaded.wait()
            finally:
                module._chotu_close_checkpoint(False);module._chotu_tls.saved_response_cache.close();requests.Session=original_factory

    def test_blocked_category_preload_does_not_replace_previous_export(self):
        original_factory=requests.Session
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);file=root/'main.py';file.write_bytes((ROOT/'main.py').read_bytes())
            spec=importlib.util.spec_from_file_location('blocked_export_fixture',file);module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
            product={'id':1,'handle':'serum','variants':[{'id':10,'available':True,'price':'1490'}]}
            module.discover=lambda base,log:([product],True)
            module.collections=lambda base,log:{}
            resume_support.enable_resume(module.__dict__);fast.enable_fast_fetch(module.__dict__)
            module.OUT=root/'exports';module.OUT.mkdir();previous=module.OUT/'Fixture.xlsx';previous.write_bytes(b'previous verified file')
            module._chotu_tls.saved_api_path=root/'api.sqlite3'
            module._chotu_tls.saved_response_cache=resume_support.ResponseCache(root/'pages.sqlite3')
            module._chotu_open_checkpoint('https://blocked.example','Fixture')
            def fetch(url,tries=4):
                if '/collections.json' in url:return {'collections':[{'handle':'blocked','title':'Lipstick'}]}
                response=requests.Response();response.status_code=429
                module._chotu_tls.fast_gate.response(response)
                raise module._chotu_tls.fast_gate.blocked_error()
            module.get_json=fetch
            try:
                with self.assertRaises(module.ChotuRateLimited):module.run_site('https://blocked.example','Fixture',str(ROOT/'ChotuApp_Template.xlsx'),lambda x:None)
                self.assertEqual(previous.read_bytes(),b'previous verified file')
                self.assertFalse((module.OUT/'Fixture - REFRESH REVIEW.xlsx').exists())
            finally:
                module._chotu_close_checkpoint(False);module._chotu_tls.saved_response_cache.close();requests.Session=original_factory

    def test_installer_adds_hooks_only_and_is_idempotent(self):
        source=(ROOT/'main.py').read_text();updated=install_fast_fetch.transform(source)
        self.assertEqual(updated.replace(install_fast_fetch.RESUME_HOOK,'').replace(install_fast_fetch.FAST_HOOK,''),source)
        self.assertEqual(install_fast_fetch.transform(updated),updated)
