import copy,importlib.util,json,sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from bs4 import BeautifulSoup
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT));sys.path.insert(0,str(ROOT/'cloud'))
import laam_source_support as source
FIXTURES=json.loads((ROOT/'tests/fixtures/laam_public_products.json').read_text())
def html(raw,currency='PKR'):
 flat=[]
 def encode(value):
  i=len(flat);flat.append(None)
  if isinstance(value,dict):flat[i]={k:encode(v) for k,v in value.items()}
  elif isinstance(value,list):flat[i]=[encode(v) for v in value]
  else:flat[i]=value
  return i
 encode(raw)
 return '<script id="__NUXT_DATA__">'+json.dumps(flat)+'</script><script type="application/ld+json">'+json.dumps({'@type':'Product','offers':{'priceCurrency':currency,'price':raw['variants'][0]['price']}})+'</script>'
class LaamTests(unittest.TestCase):
 def test_real_captured_products_preserve_observed_price_stock_options_images(self):
  for fixture in FIXTURES:
   raw=fixture['product'];p=source.product_state(BeautifulSoup(html(raw),'html.parser'),raw['handle'])
   self.assertEqual(len(p['variants']),len(raw['variants']));self.assertEqual(p['body_html'],raw['description'])
   for before,after in zip(raw['variants'],p['variants']):
    self.assertEqual(after['id'],before['shopify_id']);self.assertEqual(after['price'],before['price']);self.assertEqual(after['inventory_quantity'],before['inventory_quantity']);self.assertEqual(after['option1'],before['title']);self.assertEqual(after['sku'],before['sku'])
   self.assertEqual([i['src'] for i in p['images']],[i['src'] for i in raw['media']]);self.assertEqual(p['_source_url'],fixture['url'])
 def test_actual_0085_reversed_option_array_uses_explicit_source_positions(self):
  raw=json.loads((ROOT/'tests/fixtures/laam_public_product_0085.json').read_text());before=copy.deepcopy(raw)
  product=source.product_state(BeautifulSoup(html(raw),'html.parser'),'0085')
  self.assertEqual(raw,before);self.assertEqual([d['name'] for d in product['options']],['Kids Size','Stitching Type'])
  self.assertEqual(len(product['variants']),4)
  for original,variant in zip(raw['variants'],product['variants']):
   size,stitching=original['title'].split(' / ');self.assertEqual(variant['option1'],size);self.assertEqual(variant['option2'],stitching)
   for key in ('price','compare_at_price','sku','inventory_quantity'):self.assertEqual(variant[key],original[key])
   self.assertEqual(variant['id'],original['shopify_id']);self.assertEqual(variant['available'],original['stock_status']=='limited-stock')
  self.assertEqual(sum(v['available'] for v in product['variants']),1)
 def test_incomplete_duplicate_and_non_numeric_positions_are_not_guessed(self):
  original=json.loads((ROOT/'tests/fixtures/laam_public_product_0085.json').read_text())
  for positions in ((1,1),(None,1),('2',1),(True,2),(1,3)):
   raw=copy.deepcopy(original)
   for option,position in zip(raw['options'],positions):option['position']=position
   with self.assertRaisesRegex(ValueError,'option positions'):source.normalize(raw,'0085')
 def test_parser_errors_are_not_mislabelled_as_source_access_blocks(self):
  keys=('discover','collections','refresh_variant_prices','storefront_cover','verified_top20','audit_fractional_prices','storefront_price_audit')
  for error,prefix in ((ValueError('bad option data'),'ADAPTER ERROR:'),(source.SourceAccessError('robots denied'),'SOURCE BLOCKED:')):
   def original(*args):raise error
   g={key:(lambda *args:None) for key in keys};before=dict(g);g['_chotuapp_run_site_original']=original;source.enable_laam_source(g)
   with self.assertRaisesRegex(RuntimeError,prefix):g['_chotuapp_run_site_original'](source.BASE,'LAAM',None,lambda _:None)
   for key in keys:self.assertIs(g[key],before[key])
 def test_foreign_currency_rejected(self):
  raw=FIXTURES[0]['product']
  with self.assertRaisesRegex(ValueError,'currency unverified'):source.product_state(BeautifulSoup(html(raw,'USD'),'html.parser'),raw['handle'])
 def test_unknown_stock_not_invented_and_observed_out_of_stock_kept(self):
  raw=copy.deepcopy(FIXTURES[0]['product']);raw['variants'][0]['stock_status']='out-of-stock';raw['variants'][0]['inventory_quantity']=0
  p=source.normalize(raw,raw['handle']);self.assertFalse(p['variants'][0]['available']);self.assertEqual(p['variants'][0]['inventory_quantity'],0)
  raw['variants'][0]['stock_status']=''
  with self.assertRaisesRegex(ValueError,'availability unverified'):source.normalize(raw,raw['handle'])
 def test_missing_quantity_blank_and_unknown_options_rejected(self):
  raw=copy.deepcopy(FIXTURES[0]['product']);raw['variants'][0].pop('inventory_quantity');p=source.normalize(raw,raw['handle']);self.assertNotIn('inventory_quantity',p['variants'][0]);raw['variants'][0]['title']='Not a declared size'
  with self.assertRaisesRegex(ValueError,'declared option'):source.normalize(raw,raw['handle'])
 def test_temporary_server_failures_pause_without_becoming_terminal_blocks(self):
  from types import SimpleNamespace
  for status in (500,502,504):
   response=SimpleNamespace(status_code=status,content=b'server failure')
   g={'_chotu_wait':lambda:None,'_chotu_backoff':lambda *a:None,'_chotu_debug_request':lambda *a:None,'_chotu_debug_response':lambda *a,**k:None,'S':SimpleNamespace(get=lambda *a,**k:response)}
   reader=source.PublicCatalog(g,lambda _:None)
   with self.assertRaisesRegex(RuntimeError,'SOURCE PAUSED: LAAM temporary server failure'):reader.read(source.BASE+'/products/one')
 def test_temporary_failure_recovers_in_same_runner_with_fresh_retry(self):
  from types import SimpleNamespace
  from unittest.mock import Mock
  bad=SimpleNamespace(status_code=500,content=b'failure');good=SimpleNamespace(status_code=200,content=b'product')
  session=Mock();session.get.return_value=bad;session.get_fresh.return_value=good;backoff=Mock()
  g={'_chotu_wait':lambda:None,'_chotu_backoff':backoff,'_chotu_debug_request':lambda *a:None,'_chotu_debug_response':lambda *a,**k:None,'S':session}
  self.assertIs(source.PublicCatalog(g,lambda _:None).read(source.BASE+'/products/one'),good)
  session.get_fresh.assert_called_once_with(source.BASE+'/products/one',timeout=35,allow_redirects=False)
  self.assertEqual([c.args[0] for c in backoff.call_args_list if c.args[1]=='LAAM temporary server cooldown'],[60])
 def test_retry_budget_is_bounded_and_rate_limit_is_not_retried(self):
  from types import SimpleNamespace
  from unittest.mock import Mock
  for status,count in ((500,2),(429,0)):
   session=Mock();session.get.return_value=SimpleNamespace(status_code=status,content=b'failure');session.get_fresh.return_value=session.get.return_value
   g={'_chotu_wait':lambda:None,'_chotu_backoff':lambda *a:None,'_chotu_debug_request':lambda *a:None,'_chotu_debug_response':lambda *a,**k:None,'S':session}
   with self.assertRaisesRegex(RuntimeError,'SOURCE PAUSED'):source.PublicCatalog(g,lambda _:None).read(source.BASE+'/products/one')
   self.assertEqual(session.get_fresh.call_count,count)
 def test_source_does_not_follow_other_hosts(self):
  reader=source.PublicCatalog({},lambda _:None)
  with self.assertRaisesRegex(ValueError,'changed host'):reader.read('https://other.example/products/x')
 def test_cached_public_pages_do_not_repeat_network_or_pacing_delay(self):
  from types import SimpleNamespace
  from resume_support import ResponseCache
  import requests
  with tempfile.TemporaryDirectory() as tmp:
   cache=ResponseCache(Path(tmp)/'pages.sqlite3');r=requests.Response();r.url='https://laam.com/products/test';r.status_code=200;r._content=b'cached product page';r.encoding='utf-8';cache.put(r.url,r)
   sleeps=[];g={'_chotu_tls':SimpleNamespace(saved_response_cache=cache),'_chotu_wait':lambda:None,'_chotu_backoff':lambda *args:sleeps.append(args),'_chotu_debug_request':lambda *args:None,'_chotu_debug_response':lambda *args,**kwargs:None}
   reader=source.PublicCatalog(g,lambda _:None);self.assertEqual(reader.read(r.url).content,r.content);self.assertEqual(sleeps,[]);cache.close()
 def test_bad_product_does_not_prevent_later_good_product_cache(self):
  import requests
  from types import SimpleNamespace
  raw=FIXTURES[0]['product'];good=source.BASE+'/products/'+raw['handle'];bad=source.BASE+'/products/000-broken';visited=[]
  class FixtureSource(source.PublicCatalog):
   def read(self,url,allow_missing=False):
    visited.append(url)
    if url==bad:raise source.TemporaryProductError('SOURCE PAUSED: HTTP 500')
    payload={source.BASE+'/robots.txt':'User-agent: *\nAllow: /\nSitemap: https://laam.com/sitemap.xml',source.BASE+'/sitemap.xml':'<urlset><url><loc>'+bad+'</loc></url><url><loc>'+good+'</loc></url></urlset>',good:html(raw)}[url]
    r=requests.Response();r.status_code=200;r._content=payload.encode();r.encoding='utf-8';return r
  with tempfile.TemporaryDirectory() as tmp:
   reader=FixtureSource({'_chotu_tls':SimpleNamespace(saved_api_path=Path(tmp)/'api.sqlite3')},lambda _:None)
   with self.assertRaisesRegex(source.TemporaryProductError,'catalog incomplete'):reader.catalog()
   self.assertIn(good,visited);self.assertIsNotNone(reader.cache.get(good));self.assertIsNone(reader.cache.get(bad));reader.close()
 def test_public_sitemap_images_are_not_treated_as_products(self):
  import requests
  class FixtureSource(source.PublicCatalog):
   def read(self,url,allow_missing=False):
    payload={source.BASE+'/robots.txt':'User-agent: *\nAllow: /\nSitemap: https://laam.com/sitemap.xml',source.BASE+'/sitemap.xml':'<sitemapindex><sitemap><loc>https://laam.com/part.xml</loc></sitemap></sitemapindex>',source.BASE+'/part.xml':'<urlset xmlns:i="urn:image"><url><loc>https://laam.com/products/2665a</loc><i:image><i:loc>https://cdn.example/image.jpg</i:loc></i:image></url></urlset>',source.BASE+'/products/2665a':html(FIXTURES[0]['product'])}[url]
    response=requests.Response();response.status_code=200;response._content=payload.encode();response.encoding='utf-8';return response
  products,complete=FixtureSource({},lambda _:None).catalog();self.assertTrue(complete);self.assertEqual(len(products),1);self.assertEqual(products[0]['handle'],'2665a')
 def test_parsed_cache_restarts_without_losing_any_product_or_variant_fields(self):
  product=source.normalize(FIXTURES[0]['product'],FIXTURES[0]['product']['handle']);url=product['_source_url']
  with tempfile.TemporaryDirectory() as tmp:
   path=Path(tmp)/'derived.sqlite3';cache=source.ParsedProductCache(path,'revision-one');cache.put(url,product);cache.close();cache=source.ParsedProductCache(path,'revision-one');self.assertEqual(cache.get(url),product);cache.close()
   cache=source.ParsedProductCache(path,'revision-two');self.assertIsNone(cache.get(url));cache.close()
 def test_corrupt_parsed_cache_is_rejected_instead_of_replayed(self):
  product=source.normalize(FIXTURES[0]['product'],FIXTURES[0]['product']['handle']);url=product['_source_url']
  with tempfile.TemporaryDirectory() as tmp:
   cache=source.ParsedProductCache(Path(tmp)/'derived.sqlite3','one');cache.put(url,product);cache.db.execute('UPDATE parsed_products SET payload=?',(b'corrupted',));cache.db.commit();self.assertIsNone(cache.get(url));cache.close()
 def test_resume_reuses_parsed_products_without_another_page_parse(self):
  from types import SimpleNamespace
  import requests
  raw=FIXTURES[0]['product'];url=source.BASE+'/products/'+raw['handle'];calls=[]
  class FixtureSource(source.PublicCatalog):
   def read(self,page,allow_missing=False):
    calls.append(page)
    data={source.BASE+'/robots.txt':'User-agent: *\nAllow: /\nSitemap: https://laam.com/sitemap.xml',source.BASE+'/sitemap.xml':'<urlset><url><loc>'+url+'</loc></url></urlset>',url:html(raw)}[page]
    response=requests.Response();response.status_code=200;response._content=data.encode();response.encoding='utf-8';return response
  with tempfile.TemporaryDirectory() as tmp:
   g={'_chotu_tls':SimpleNamespace(saved_api_path=Path(tmp)/'api.sqlite3')}
   first=FixtureSource(g,lambda _:None);products,complete=first.catalog();first.close();calls.clear();second=FixtureSource(g,lambda _:None)
   with patch.object(source,'product_state',side_effect=AssertionError('Successful product parsed again')):
    resumed,complete=second.catalog()
   self.assertEqual(resumed,products);self.assertEqual(second.cache_hits,1);self.assertNotIn(url,calls);second.close()
 def test_export_integration_uses_existing_template_and_preserves_fields(self):
  spec=importlib.util.spec_from_file_location('laam_test_engine',ROOT/'main.py');engine=importlib.util.module_from_spec(spec);spec.loader.exec_module(engine)
  with tempfile.TemporaryDirectory() as tmp:
   engine.OUT=Path(tmp);products=[source.normalize(f['product'],f['product']['handle']) for f in FIXTURES]
   source.enable_laam_source(vars(engine));before=engine.discover;logs=[]
   with patch.object(source.PublicCatalog,'catalog',return_value=(products,True)):
    engine._chotuapp_run_site_original(source.BASE,'LAAM',ROOT/'importer-files/Chotu_Products_Import_Template_Discount_Tags.xlsx',logs.append)
   self.assertIs(engine.discover,before)
   book=load_workbook(Path(tmp)/'LAAM - REFRESH REVIEW.xlsx');sheet=book['Products'];rows=[dict(zip([c.value for c in sheet[1]],v)) for v in sheet.iter_rows(min_row=2,values_only=True)]
   self.assertEqual(len(rows),4);self.assertEqual({r['source_url'] for r in rows},{f['url'] for f in FIXTURES})
   variants=book['Variant Combinations'];self.assertEqual(variants.max_row-1,sum(len(p['variants']) for p in products));self.assertTrue(any('REVIEW=1' in line for line in logs));self.assertFalse(any('TOP 20: Storefront best-selling' in line for line in logs));book.close()
if __name__=='__main__':unittest.main()
