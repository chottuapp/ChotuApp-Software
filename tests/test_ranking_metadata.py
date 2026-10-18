import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from types import SimpleNamespace
from openpyxl import Workbook,load_workbook
import ranking_support as ranking

class RankingMetadataTests(unittest.TestCase):
    def row(self,pid,cat,base=100,sale=None):
        return {'product_id':pid,'category':cat,'base_price (PKR)':base,'sale_price (PKR, agar discount ho)':sale,'source_url':'https://elitecosmeticspk.com/products/'+pid}
    def test_beauty_category_precedes_cheaper_automotive(self):
        rows=[self.row('car','Automotive',1),self.row('makeup','Lipstick',500),self.row('cheap','Lipstick',100)]
        ranking.order_rows(rows,policy='legacy')
        self.assertEqual([r['product_id'] for r in rows],['cheap','makeup','car'])
        self.assertTrue(all(not r.get('popularity_basis') for r in rows))
    def test_verified_regular_price_product_can_outrank_discount_without_crossing_category(self):
        rows=[self.row('cheap','Lipstick',20),self.row('best','Lipstick',200),self.row('car','Automotive',1)]
        facts={'best':{'url':'https://elitecosmeticspk.com/collections/all?sort_by=best-selling','checked_at':'2026-10-17T00:00:00+00:00'}}
        ranking.order_rows(rows,['best'],facts,policy='legacy')
        self.assertEqual([r['product_id'] for r in rows],['best','cheap','car'])
        self.assertEqual(rows[0]['popularity_basis'],'store_bestseller')
    def test_handles_without_evidence_do_not_create_popularity(self):
        rows=[self.row('expensive','Lipstick',200),self.row('cheap','Lipstick',100)]
        self.assertEqual(ranking.order_rows(rows,['expensive'],policy='legacy'),[])
        self.assertEqual(rows[0]['product_id'],'cheap')
        self.assertFalse(any(r.get('product_priority') for r in rows))
    def test_decimal_and_zero_sale_not_free_and_category_conflicts_fail(self):
        rows=[self.row('zero','Lipstick',100,0),self.row('decimal','Lipstick',100,99.25)]
        ranking.order_rows(rows,policy='legacy');self.assertEqual(rows[0]['product_id'],'decimal')
        rows[0]['category_priority']=2
        with self.assertRaises(ValueError):ranking.order_rows(rows,policy='legacy')
    def test_reordering_preserves_hyperlinks_styles_formulas_and_other_sheets(self):
        import main
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'out.xlsx';w=Workbook();s=w.active;s.title='Products'
            s.append(['product_id','category','base_price (PKR)','source_url','variant_image_url','formula'])
            s.append(['car','Automotive',1,'https://elitecosmeticspk.com/products/car','https://img/car','=1+1'])
            s.append(['lip','Lipstick',100,'https://elitecosmeticspk.com/products/lip','https://img/lip','=2+2'])
            s['E2'].hyperlink='https://img/car';s['E3'].hyperlink='https://img/lip';s['E3'].number_format='@'
            w.create_sheet('Options').append(['keep']);w.save(p);w.close()
            facts={key:{'category':cat,'subcategory':'','category_order':rank,'subcategory_order':1,'product_order':1,'source_url':'https://elitecosmeticspk.com/collections/'+key,'checked_at':'2026-10-18T00:00:00+00:00'} for key,cat,rank in (('lip','Lipstick',1),('car','Automotive',2))}
            with patch.object(main._chotu_tls,'website_order_evidence',facts,create=True):main.finalize_catalog_order(p,Path(d)/'absent.csv',lambda _:None)
            v=load_workbook(p);self.assertEqual(v['Products']['A2'].value,'lip');self.assertEqual(v['Products']['E2'].hyperlink.target,'https://img/lip');self.assertEqual(v['Products']['E2'].hyperlink.ref,'E2');self.assertEqual(v['Products']['E2'].number_format,'@');self.assertEqual(v['Products']['F2'].value,'=2+2');self.assertEqual(v['Options']['A1'].value,'keep');v.close()
    def test_verified_grid_captures_actual_evidence_and_unverified_grid_does_not(self):
        import main
        html='<select><option value="best-selling" selected>Best</option></select><div id="product-grid"><a href="/products/lip">Lip</a></div>'
        with patch.object(main.S,'get',return_value=SimpleNamespace(text=html,raise_for_status=lambda:None)):
            self.assertEqual(main.verified_top20('https://elitecosmeticspk.com',lambda _:None,{'lip'}),['lip'])
            fact=main._chotu_tls.ranking_evidence['lip'];self.assertIn('sort_by=best-selling',fact['url']);self.assertTrue(fact['checked_at'])
        with patch.object(main.S,'get',return_value=SimpleNamespace(text='<div>No sort</div>',raise_for_status=lambda:None)):
            self.assertEqual(main.verified_top20('https://elitecosmeticspk.com',lambda _:None,{'lip'}),[])
            self.assertEqual(main._chotu_tls.ranking_evidence,{})
    def test_runtime_build_contains_helper_and_resumed_columns_keep_custom_values(self):
        import sys
        sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'cloud'))
        from cloud.runtime import prepare_engine
        with tempfile.TemporaryDirectory() as d:
            runtime=prepare_engine(Path(d));self.assertTrue((Path(d)/'runtime/ranking_support.py').is_file())
        w=Workbook();s=w.active;s.append(['product_id','custom']);s.append(['p','keep']);self.assertTrue(ranking.ensure_columns(s));self.assertFalse(ranking.ensure_columns(s));self.assertEqual(s['B2'].value,'keep');w.close()
