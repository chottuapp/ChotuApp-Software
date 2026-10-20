import importlib.util,json,os,sys,tempfile,unittest,zipfile
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook,load_workbook
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT));sys.path.insert(0,str(ROOT/'cloud'))
from rule804_pipeline import process,run
import group_fetch

def book(path,price=100,orphan=False):
 w=Workbook();s=w.active;s.title='Products';s.append(['shop_code','product_id','Category','name','description','base_price (PKR)','source_url','subcategory_source','subcategory_evidence','subcategory_checked_at','formula','website_product_order']);s.append([None,'p1','Body Care','Face Wash 100ml','Face Wash 100ml. Ingredients: Water, glycerin. Use twice daily. Colour May vary',price,'https://shop.example/products/face-wash','existing','','','=F2*2',4])
 for name in ('Product Images','Option Groups','Options','Variant Combinations'):
  x=w.create_sheet(name);x.append(['product_id','source_value']);x.append(['missing' if orphan else 'p1','do not change'])
 w.save(path);w.close()
class RuleTests(unittest.TestCase):
 def test_formula_prices_stock_links_and_other_sheets_preserved(self):
  with tempfile.TemporaryDirectory() as tmp:
   src=Path(tmp)/'Miss Rose.xlsx';out=Path(tmp)/'out.xlsx';book(src)
   report=process(src,out,'Miss Rose',{'Miss Rose':'SE123'})
   self.assertEqual(report['products'],1)
   w=load_workbook(out);s=w['Products'];self.assertEqual(s['A2'].value,'SE123');self.assertNotEqual(s['C2'].value,'Body Care');self.assertEqual(s['K2'].value,'=F2*2');self.assertEqual(s['L2'].value,4);self.assertIn('Water, glycerin',s['E2'].value);self.assertIn('twice daily',s['E2'].value);self.assertEqual(s['F2'].value,100);w.close()
   with zipfile.ZipFile(src) as a,zipfile.ZipFile(out) as b:
    for part in a.namelist():
     if part!='xl/worksheets/sheet1.xml':self.assertEqual(a.read(part),b.read(part),part)
 def test_invalid_prices_or_orphan_refs_hold_workbook(self):
  with tempfile.TemporaryDirectory() as tmp:
   src=Path(tmp)/'Miss Rose.xlsx';out=Path(tmp)/'out.xlsx'
   for price,orphan in [(0,False),(100,True)]:
    book(src,price,orphan)
    with self.assertRaises(ValueError):process(src,out,'Miss Rose',{})
    self.assertFalse(out.exists())
 def test_unknown_code_blank_and_excluded_store_not_processed(self):
  with tempfile.TemporaryDirectory() as tmp:
   src=Path(tmp)/'Zero Lifestyle.xlsx';out=Path(tmp)/'out.xlsx';book(src);process(src,out,'Zero Lifestyle',{})
   w=load_workbook(out);self.assertIsNone(w['Products']['A2'].value);w.close()
   with self.assertRaisesRegex(ValueError,'Excluded'):process(src,out,'LAAM',{})
 def test_unchanged_source_is_cached_and_reports_stay_outside_workbook(self):
  with tempfile.TemporaryDirectory() as tmp:
   root=Path(tmp);src=root/'raw';src.mkdir();out=root/'checked';book(src/'Miss Rose.xlsx')
   first=run(src,out);second=run(src,out)
   self.assertEqual(first['files'],second['files']);self.assertEqual(len(list(out.glob('*.xlsx'))),1);self.assertEqual(second['held'],[])
class DispatchTests(unittest.TestCase):
 def test_fresh_26_groups_leave_excluded_active_lanes_running(self):
  commands=[]
  def gh(*args):
   commands.append(args)
   if args[:2]==('run','list'):return json.dumps([{'databaseId':50,'displayTitle':'Cloud group solo-laam','status':'in_progress'},{'databaseId':51,'displayTitle':'Cloud group solo-sapphire','status':'in_progress'}])
   return ''
  with tempfile.TemporaryDirectory() as tmp,patch.dict(os.environ,{'GITHUB_REPOSITORY':'o/r','GITHUB_RUN_ID':'999','CLOUD_STORES':'','RESUME_RUN_ID':'','CLOUD_FRESH_EXCLUDING_THREE':'true'}),patch.object(group_fetch,'gh',side_effect=gh):
   previous=Path.cwd()
   try:os.chdir(tmp);group_fetch.main();plan=json.loads(Path('fetch-group-plan.json').read_text())
   finally:os.chdir(previous)
  groups=plan['groups'];self.assertEqual([len(g['stores']) for g in groups],[3,3,3,3,3,3,3,3,2]);self.assertTrue(all(g['runner_group'].startswith('fresh26-999-') for g in groups));names={s['name'] for g in groups for s in g['stores']};self.assertEqual(len(names),26);self.assertFalse(names&{'LAAM','Khaadi','Sapphire'})
  dispatches=[c for c in commands if c[:2]==('workflow','run')];self.assertEqual(len(dispatches),9);self.assertTrue(all('resume_run_id=' in c for c in dispatches));self.assertFalse(any('cloud-stop.yml' in c for c in commands))
 def test_unrelated_active_lane_blocks_duplicate_fresh_work(self):
  with patch.dict(os.environ,{'GITHUB_REPOSITORY':'o/r','GITHUB_RUN_ID':'999','CLOUD_FRESH_EXCLUDING_THREE':'true'}),patch.object(group_fetch,'gh',return_value=json.dumps([{'databaseId':5,'displayTitle':'Cloud group other','status':'in_progress'}])):
   with self.assertRaisesRegex(ValueError,'Active run IDs'):group_fetch.main()
