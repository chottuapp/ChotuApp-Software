import unittest,tempfile
from pathlib import Path
from openpyxl import Workbook
from price_stock_gate import check,workbook_issues
class PriceStockTests(unittest.TestCase):
 def row(self,**extra):return dict({'product_id':'P1','base_price (PKR)':100,'sale_price (PKR, agar discount ho)':80,'available (yes/no)':'yes'},**extra)
 def test_exact_decimal_prices_and_unknown_quantity_are_preserved(self):
  row=self.row(**{'base_price (PKR)':'100.25','sale_price (PKR, agar discount ho)':'80.10'});before=row.copy();self.assertEqual(check(row),[]);self.assertEqual(row,before)
 def test_invalid_prices_and_discounts_fail(self):
  for change in [{'base_price (PKR)':0},{'base_price (PKR)':'NaN'},{'sale_price (PKR, agar discount ho)':101},{'sale_price (PKR, agar discount ho)':-1},{'sale_price (PKR, only if discounted)':70}]:self.assertTrue(check(self.row(**change)))
 def test_unknown_and_conflicting_stock_fail_without_becoming_in_stock(self):
  for change in [{'available (yes/no)':''},{'stock_quantity (agar public ho)':0},{'stock_quantity (agar public ho)':-1}]:self.assertTrue(check(self.row(**change)))
  self.assertEqual(check(self.row(**{'available (yes/no)':'no','stock_quantity (agar public ho)':0})),[])
 def test_mixed_variant_stock_product_remains_available(self):
  with tempfile.TemporaryDirectory() as tmp:
   path=Path(tmp)/'x.xlsx';book=Workbook();sheet=book.active;sheet.title='Products';headers=list(self.row());sheet.append(headers);sheet.append(list(self.row().values()));variants=book.create_sheet('Variant Combinations');variants.append(headers);variants.append(list(self.row(**{'available (yes/no)':'no'}).values()));variants.append(list(self.row().values()));book.save(path)
   self.assertEqual(workbook_issues(path),{})
   sheet.cell(2,headers.index('available (yes/no)')+1).value='no';book.save(path);self.assertIn('product and variant availability disagree',' '.join(workbook_issues(path)['P1']))
