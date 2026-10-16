import importlib.util,tempfile,unittest
from pathlib import Path
from openpyxl import load_workbook,Workbook
ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('stock_main',ROOT/'main.py');main=importlib.util.module_from_spec(spec);spec.loader.exec_module(main)
class StockImageTests(unittest.TestCase):
 def test_stock_total_requires_all_counts(self):
  self.assertEqual(main.product_stock_quantity({},[{'inventory_quantity':4},{'inventory_quantity':'6'}]),10)
  for variant in [{},{'inventory_quantity':True},{'inventory_quantity':-1},{'inventory_quantity':4,'inventory_management':None}]:
   self.assertEqual(main.product_stock_quantity({},[{'inventory_quantity':4},variant]),'')
  self.assertEqual(main.product_stock_quantity({},[{'inventory_quantity':0}]),0)
 def test_exact_image_association_and_ambiguous_options(self):
  p={'images':[{'id':8,'src':'https://example.com/red.jpg','variant_ids':[1]},{'id':9,'src':'https://example.com/blue.jpg'}]}
  a={'id':1,'option1':'Red'};b={'id':2,'image_id':9,'option1':'Red'}
  self.assertEqual(main.variant_image_url(p,a),'https://example.com/red.jpg')
  self.assertEqual(main.variant_image_url(p,b),'https://example.com/blue.jpg')
  self.assertEqual(main.variant_image_url(p,{'id':3}),'')
  self.assertEqual(main.option_image_url(p,[a,b],1,'Red'),'')
  self.assertEqual(main.option_image_url(p,[a],1,'Red'),'https://example.com/red.jpg')
 def test_templates_and_hyperlink_roundtrip(self):
  for file in ['ChotuApp_Template.xlsx','importer-files/Chotu_Products_Import_Template_Discount_Tags.xlsx','importer-files/Chotu_Products_Import_Template_Updated.xlsx']:
   w=load_workbook(ROOT/file);self.assertIn('stock_quantity (kitni quantity bechni hai)',[c.value for c in w['Products'][1]]);w.close()
  w=Workbook();ws=w.active;ws.title='Variant Combinations';ws.append(['variant_image_url','stock_quantity (agar public ho)'])
  main.write_row(ws,{'variant_image_url':'https://example.com/red.jpg'})
  with tempfile.TemporaryDirectory() as tmp:
   p=Path(tmp)/'test.xlsx';w.save(p);saved=load_workbook(p);self.assertEqual(saved.active['A2'].hyperlink.target,'https://example.com/red.jpg');self.assertIsNone(saved.active['B2'].value);saved.close()
if __name__=='__main__':unittest.main()
