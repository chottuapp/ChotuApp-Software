import importlib.util
import json
import os
from pathlib import Path
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
from http.server import BaseHTTPRequestHandler, HTTPServer
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('chotu', ROOT / 'main.py')
main = importlib.util.module_from_spec(spec)
spec.loader.exec_module(main)

class GuiTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.config = patch.object(main, 'CONFIG', Path(self.tmp.name)/'sites.json')
        self.config.start()
        self.root = main.tk.Tk()
        self.errors = []
        self.root.report_callback_exception = lambda *args: self.errors.append(args)
        self.app = main.App(self.root)
        self.app.sites = [{'name':f'Store {i}', 'url':f'https://fixture{i}.example'} for i in range(40)]
        self.app.redraw()
        self.root.update()

    def tearDown(self):
        main._chotu_pause_gate.set()
        # Cancel scheduled callbacks before destroying interpreter widgets.
        for callback in self.root.tk.splitlist(self.root.tk.call('after', 'info')):
            self.root.tk.call('after', 'cancel', callback)
        self.root.destroy()
        self.config.stop()
        self.tmp.cleanup()
        self.assertFalse(self.errors)

    def wheel(self, widget, sequence='<Button-5>', **kwargs):
        widget.event_generate(sequence, **kwargs)
        self.root.update()

    def test_store_scroll_independent_after_redraw(self):
        a=self.app
        page=a._wheel_workspace._parent_canvas
        page.yview_moveto(.25); self.root.update()
        before=page.yview()
        store=a.store_list._parent_canvas
        before_store=store.yview()
        child=a.store_list.winfo_children()[0].winfo_children()[0]
        self.wheel(child)
        self.assertGreater(store.yview()[0],before_store[0])
        self.assertEqual(page.yview(),before)
        store.yview_moveto(1);self.root.update()
        self.wheel(child)
        self.assertEqual(page.yview(),before)
        a.redraw();self.root.update()
        child=a.store_list.winfo_children()[0].winfo_children()[0]
        self.assertEqual(child.bindtags()[0],a._wheel_tag)

    def test_logs_scroll_without_scrolling_page(self):
        a=self.app;page=a._wheel_workspace._parent_canvas
        page.yview_moveto(1);self.root.update()
        a.logbox.insert('end','line\n'*120);self.root.update()
        a.logbox.yview_moveto(0);before=page.yview()
        self.wheel(a.logbox)
        self.assertGreater(a.logbox.yview()[0],0)
        self.assertEqual(page.yview(),before)

    def test_windows_wheel_event_and_page_scroll(self):
        page=self.app._wheel_workspace._parent_canvas
        page.yview_moveto(0)
        self.wheel(page,'<MouseWheel>',delta=-120)
        self.assertGreater(page.yview()[0],0)

    def test_selection_and_idle_pause(self):
        a=self.app
        self.assertEqual(a.refresh_selected.cget('state'),'normal')
        a.store_checks[a.sites[5]['url']].set(True)
        a._selection_changed()
        self.assertEqual(a._selected_indices(),[5])
        self.assertEqual(a.refresh_selected.cget('state'),'normal')
        a.toggle_pause();self.assertTrue(main._chotu_pause_gate.is_set())
        a.running=True;a.toggle_pause();self.assertFalse(main._chotu_pause_gate.is_set())
        completed=threading.Event()
        t=threading.Thread(target=lambda:(main._chotu_wait(),completed.set()),daemon=True);t.start()
        self.assertFalse(completed.wait(.05))
        a.toggle_pause();self.assertTrue(completed.wait(1));t.join(1)

    def pump_until_done(self):
        deadline=time.monotonic()+3
        while self.app.running and time.monotonic()<deadline:
            self.root.update();time.sleep(.01)
        self.assertFalse(self.app.running)

    def test_worker_failure_restores_controls(self):
        a=self.app
        with patch.object(a,'_run_checked',side_effect=RuntimeError('fixture failure')):
            a._launch_jobs([a.sites[0]],'fixture.xlsx')
            self.pump_until_done()
        self.assertTrue(a.job_failed)
        self.assertEqual(a.refresh.cget('state'),'normal')
        self.assertEqual(a.pause_button.cget('state'),'disabled')
        self.assertTrue(main._chotu_pause_gate.is_set())
        self.assertIn('ERROR',a.live_status.cget('text'))

    def test_review_is_distinct_from_completed(self):
        a=self.app
        def job(*args):
            a.events.put('SAVED REVIEW EXCEL: fixture.xlsx')
            a.events.put(('__JOB_RESULT__',True,'Fixture'))
        with patch.object(a,'_run_checked',side_effect=job):
            a._launch_jobs([a.sites[0]],'fixture.xlsx');self.pump_until_done()
        self.assertIn('REVIEW REQUIRED',a.live_status.cget('text'))
        self.assertEqual(a.progress_value.get(),100)

    def test_touchpad_drag_scrolls_store_without_toggling(self):
        from types import SimpleNamespace
        a = self.app
        checkbox = a.store_list.winfo_children()[0].winfo_children()[0]
        canvas = a.store_list._parent_canvas
        canvas.yview_moveto(0); self.root.update()
        page_before = a._wheel_workspace._parent_canvas.yview()
        a._pan_start(SimpleNamespace(widget=checkbox, y_root=400))
        a._pan_move(SimpleNamespace(y_root=300))
        a._pan_end(SimpleNamespace(y_root=300))
        self.assertGreater(canvas.yview()[0], 0)
        self.assertEqual(a._selected_indices(), [])
        self.assertEqual(a._wheel_workspace._parent_canvas.yview(), page_before)

    def test_pan_bindings_and_click_selection(self):
        a = self.app
        checkbox = a.store_list.winfo_children()[0].winfo_children()[0]
        from types import SimpleNamespace
        self.assertTrue(self.root.bind_class(a._wheel_tag, '<B1-Motion>'))
        a._pan_start(SimpleNamespace(widget=checkbox, y_root=300))
        a._pan_end(SimpleNamespace(y_root=300))
        self.assertEqual(a._selected_indices(), [0])
        a._pan_start(SimpleNamespace(widget=a.bulk, y_root=300))
        self.assertIsNone(a._pan_state)

    def test_touchpad_drag_main_page(self):
        from types import SimpleNamespace
        a = self.app
        canvas = a._wheel_workspace._parent_canvas
        canvas.yview_moveto(0); self.root.update()
        a._pan_start(SimpleNamespace(widget=canvas, y_root=400))
        a._pan_move(SimpleNamespace(y_root=280))
        a._pan_end(SimpleNamespace(y_root=280))
        self.assertGreater(canvas.yview()[0], 0)

    def test_focus_root_routes_to_hovered_store(self):
        a = self.app
        page = a._wheel_workspace._parent_canvas
        page.yview_moveto(.35); self.root.update()
        canvas = a.store_list._parent_canvas
        canvas.yview_moveto(0); self.root.update()
        before_page = page.yview()
        # Windows wheel messages can name the focused root, not the hovered pane.
        x = canvas.winfo_rootx() + 20
        y = canvas.winfo_rooty() + 20
        from types import SimpleNamespace
        event = SimpleNamespace(widget=self.root, x_root=x, y_root=y, delta=-120, num=None, state=0)
        self.assertEqual(a._route_wheel(event), 'break')
        self.assertGreater(canvas.yview()[0], 0)
        self.assertEqual(page.yview(), before_page)

    def test_fetch_button_visible_at_small_window(self):
        self.root.geometry('900x640'); self.root.update()
        button = self.app.refresh_selected
        self.assertTrue(button.winfo_ismapped())
        self.assertLess(button.winfo_rooty() + button.winfo_height(),
                        self.root.winfo_rooty() + self.root.winfo_height())

    def test_selected_launch_runs_only_ticked_store(self):
        a=self.app
        a.store_checks[a.sites[7]['url']].set(True);a._selection_changed()
        template=Path(self.tmp.name)/'template.xlsx';template.touch();a.template.set(str(template))
        calls=[]
        def job(site,template):
            calls.append(site['url']);a.events.put(('__JOB_RESULT__',True,site['name']))
        with patch.object(a,'_run_checked',side_effect=job):
            a.start_selected();self.pump_until_done()
        self.assertEqual(calls,[a.sites[7]['url']])

class ExportTests(unittest.TestCase):
    def export(self,fractional,all_unavailable=False,eyebrow=False,body_image=False,with_gift=False,ranked=False,technical=False):
        price='999.88' if fractional else '999'
        product={'id':1,'title':'Fixture Cotton Shirt','handle':'fixture-shirt','vendor':'Fixture','product_type':'Shirts','body_html':'<p>Cotton shirt for daily wear.</p>','images':[{'id':1,'src':'https://example.com/shirt.jpg','position':1}],'options':[{'name':'Size'}],'variants':[{'id':11,'price':'699.44' if fractional else '699','available':False,'option1':'Medium'},{'id':12,'price':price,'available':True,'option1':'Big'}]}
        if technical:
            product['title']='Laptop Model X'
            product['product_type']='Laptops'
            product['body_html']='<table>'+''.join(f'<tr><th>Specification {i}</th><td>Verified source value {i}</td></tr>' for i in range(25))+'<tr><th>Warranty</th><td>One year</td></tr></table>'
        if body_image:
            product['body_html'] += '<img src="https://images.example/instructions.jpg">'
        if eyebrow:
            product['title'] = 'Miss Rose Fashion Eyebrow'
            product['variants'] = [
                {'id': vid, 'price':'499', 'compare_at_price':'899',
                 'available':shade != '01', 'option1':shade}
                for vid,shade in ((32830548279399,'03'),(39858917736551,'06'),
                                  (39474715197543,'02'),(32830548246631,'01'))]
        if all_unavailable:
            for variant in product['variants']: variant['available'] = False
        catalog=[product]
        if with_gift:
            import copy
            gift=copy.deepcopy(product)
            gift.update(id=2,title='Miss Rose Com Pk Gift (Free for Premium Orders Rs4999+)',handle='free-gift-orders-above-rs-4999')
            for v in gift['variants']:v['id']+=100
            catalog.append(gift)
        class Handler(BaseHTTPRequestHandler):
            def do_GET(self):
                if self.path.startswith('/products.json'): data={'products':catalog}
                elif self.path.startswith('/products/fixture-shirt.js'):
                    data={'id':product['id'],'variants':[dict(v,price=int(main.Decimal(v['price'])*100),compare_at_price=int(main.Decimal(v['compare_at_price'])*100) if v.get('compare_at_price') is not None else None) for v in product['variants']]}
                elif self.path.startswith('/products/fixture-shirt'):
                    self.send_response(200);self.end_headers();self.wfile.write(f'<span class="price-item">PKR {price}</span>'.encode());return
                else:data={'collections':[]}
                self.send_response(200);self.send_header('Content-Type','application/json');self.end_headers();self.wfile.write(json.dumps(data).encode())
            def log_message(self,*args):pass
        server=HTTPServer(('127.0.0.1',0),Handler)
        thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
        try:
            with tempfile.TemporaryDirectory() as tmp:
                folder=Path(tmp);template=folder/'template.xlsx'
                wb=Workbook();wb.remove(wb.active)
                for name,headers in json.loads((ROOT/'tests/schema.json').read_text()).items():wb.create_sheet(name).append(headers)
                wb.save(template);wb.close()
                out=folder/'exports';out.mkdir()
                final=out/'Fixture - FINAL.xlsx'
                final.write_bytes(template.read_bytes());previous_final=final.read_bytes()
                logs=[]
                with patch.object(main,'OUT',out),patch.object(main,'__file__',str(folder/'main.py')):
                    if all_unavailable:
                        with self.assertRaisesRegex(RuntimeError, 'No public, priced products'):
                            main.run_site(f'http://127.0.0.1:{server.server_port}','Fixture',str(template),logs.append)
                        self.assertEqual(final.read_bytes(), previous_final)
                        self.assertFalse((out/'Fixture.xlsx').exists())
                        return
                    if ranked:
                        main._chotu_tls.ranking_evidence={product['handle']:{'url':f'http://127.0.0.1:{server.server_port}/collections/all?sort_by=best-selling','checked_at':'2026-10-17T00:00:00+00:00'}}
                        with patch.object(main,'verified_top20',return_value=[product['handle']]) as forced_sort:
                            main.run_site(f'http://127.0.0.1:{server.server_port}','Fixture',str(template),logs.append)
                            forced_sort.assert_not_called()
                    else:
                        main.run_site(f'http://127.0.0.1:{server.server_port}','Fixture',str(template),logs.append)
                if ranked:
                    import csv
                    report=folder/'audit_reports/Fixture/Fixture - BEST SELLERS.csv'
                    with report.open(encoding='utf-8-sig') as stream:ranks=list(csv.DictReader(stream))
                    self.assertEqual(ranks,[])
                    self.assertTrue(any('FILE ORDER:' in line for line in logs))
                self.assertEqual(final.read_bytes(),previous_final)
                wb=load_workbook(out/'Fixture.xlsx')
                def records(name):
                    rows=list(wb[name].values);return [dict(zip(rows[0],r)) for r in rows[1:]]
                self.assertEqual(len(records('Products')),1)
                self.assertIn('[Website listing order]',records('Products')[0]['analysis_note'])
                notes_file=folder/'audit_reports/Fixture/Fixture - ANALYSIS NOTES.csv'
                import csv
                with notes_file.open(encoding='utf-8-sig') as stream:analysis_rows=list(csv.DictReader(stream))
                self.assertEqual(len(analysis_rows),1)
                self.assertEqual(analysis_rows[0]['product_id'],records('Products')[0]['product_id'])
                self.assertEqual(analysis_rows[0]['analysis_note'],records('Products')[0]['analysis_note'])
                if technical:
                    description=records('Products')[0]['description']
                    self.assertGreater(len(description),300)
                    self.assertIn('Specification 24: Verified source value 24',description)
                    self.assertIn('Warranty: One year',description)
                    quality=folder/'audit_reports/Fixture/Fixture - QUALITY REPORT.csv'
                    self.assertNotIn('LONG_APP_DESCRIPTION',quality.read_text(encoding='utf-8-sig'))
                if body_image:
                    image_urls={r['scraped_image_url (auto)'] for r in records('Product Images')}
                    self.assertEqual(image_urls, {'https://example.com/shirt.jpg','https://images.example/instructions.jpg'})
                variants={str(v['source_variant_id']):v for v in records('Variant Combinations')}
                if eyebrow:
                    self.assertEqual(len(variants), 4)
                    self.assertEqual({v['option_1_value'] for v in variants.values()}, {'01','02','03','06'})
                    self.assertEqual(variants['32830548246631']['available (yes/no)'], 'no')
                    self.assertEqual(variants['32830548279399']['sale_price (PKR, agar discount ho)'],499)
                    self.assertEqual(variants['32830548279399']['base_price (PKR)'],899)
                    options={o['option_name']:o for o in records('Options')}
                    self.assertEqual(options['01']['available (yes/no)'],'no')
                    self.assertEqual(records('Products')[0]['sale_price (PKR, agar discount ho)'],499)
                else:
                    self.assertEqual(len(variants),2)
                    self.assertEqual(variants['11']['available (yes/no)'],'no')
                    self.assertEqual(variants['12']['available (yes/no)'],'no' if all_unavailable else 'yes')
                    self.assertEqual(variants['11']['base_price (PKR)'],699.44 if fractional else 699)
                    self.assertEqual(variants['12']['base_price (PKR)'],999.88 if fractional else 999)
                    self.assertEqual(records('Products')[0]['available (yes/no)'], 'no' if all_unavailable else 'yes')
                wb.close()
                if fractional:
                    self.assertTrue(any('REVIEW REQUIRED:' in line for line in logs))
                    audit=(folder/'audit_reports/Fixture/Fixture - PRICE AUDIT.csv').read_text(encoding='utf-8-sig')
                    self.assertIn('API_PRICES_MATCH',audit)
        finally:
            server.shutdown();server.server_close();thread.join()
    def test_long_technical_description_exported_to_excel(self):self.export(False,technical=True)
    def test_best_seller_order_report_exported(self):self.export(False,ranked=True)
    def test_free_gift_not_exported_with_regular_product(self):self.export(False,with_gift=True)
    def test_description_image_exported_with_gallery(self):self.export(False,body_image=True)
    def test_all_sold_out_product_excluded(self):self.export(False,all_unavailable=True)
    def test_eyebrow_four_shades_sale_and_unavailable_option(self):self.export(False,eyebrow=True)
    def test_integer_prices_and_unavailable_variants(self):self.export(False)
    def test_fractional_prices_preserved_and_previous_final_protected(self):self.export(True)

class PriceVerificationTests(unittest.TestCase):
    def test_live_variant_price_replaces_catalog_price_and_preserves_discount(self):
        product={'id':1,'title':'Suit','handle':'suit','variants':[{'id':11,'price':'4250','compare_at_price':'5000','available':True,'option1':'Small'}]}
        live={'id':1,'variants':[{'id':11,'price':212500,'compare_at_price':425000,'available':True,'option1':'Small'}]}
        with patch.object(main,'get_json',return_value=live):
            variants=main.refresh_variant_prices('https://store.example',product,lambda x:None)
        self.assertEqual(main.money(variants[0]['price']),main.Decimal('2125'))
        self.assertEqual(main.money(variants[0]['compare_at_price']),main.Decimal('4250'))

    def test_rate_limit_stops_additional_live_price_requests(self):
        p={'id':1,'title':'Suit','handle':'suit','variants':[]}
        main._chotu_tls.price_throttled=False
        try:
            with patch.object(main,'get_json',side_effect=main.ChotuRateLimited('rate limited')) as get:
                with self.assertRaises(main.ChotuRateLimited):main.refresh_variant_prices('https://store.example',p,lambda x:None)
                with self.assertRaises(main.ChotuRateLimited):main.refresh_variant_prices('https://store.example',p,lambda x:None)
                self.assertEqual(get.call_count,1)
        finally:main._chotu_tls.price_throttled=False

    def test_wrong_identity_does_not_change_catalog(self):
        p={'id':1,'title':'Suit','handle':'suit','variants':[{'id':11,'price':'2500'}]}
        with patch.object(main,'get_json',return_value={'id':2,'variants':[]}):
            with self.assertRaises(ValueError):main.refresh_variant_prices('https://store.example',p,lambda x:None)
        self.assertEqual(p['variants'][0]['price'],'2500')
        for value in ('NaN','Infinity','-1','499.5'):
            with self.assertRaises(ValueError):main.shopify_minor_price(value)

class FreshFetchTests(unittest.TestCase):
    def tearDown(self):
        main._chotu_close_checkpoint()

    def test_new_fetch_ignores_failed_run_checkpoint_and_reads_changed_price(self):
        import requests
        def response(price):
            r=requests.Response();r.status_code=200
            r._content=json.dumps({'price':price}).encode()
            return r
        with tempfile.TemporaryDirectory() as tmp, patch.object(main,'__file__',str(Path(tmp)/'main.py')), patch.object(main,'_chotu_backoff'):
            with patch.object(main.S,'get',side_effect=[response(499),response(399)]) as get:
                first=main._chotu_open_checkpoint('https://store.example','Fixture')
                self.assertEqual(main.get_json('https://store.example/product.js')['price'],499)
                self.assertEqual(main.get_json('https://store.example/product.js')['price'],499)
                self.assertEqual(get.call_count,1)
                main._chotu_close_checkpoint(success=False)
                second=main._chotu_open_checkpoint('https://store.example','Fixture')
                self.assertNotEqual(first,second)
                self.assertTrue(first.exists())
                self.assertEqual(main.get_json('https://store.example/product.js')['price'],399)
                self.assertEqual(get.call_count,2)
                main._chotu_close_checkpoint(success=True)
                self.assertFalse(second.exists())

    def test_uploaded_structure_template_contains_no_old_catalog_records(self):
        wb=load_workbook(ROOT/'ChotuApp_Template.xlsx',read_only=True)
        self.assertIn('Ranking Guide',wb.sheetnames)
        self.assertTrue(set(('category_priority','product_priority','popularity_basis','popularity_source_url','popularity_checked_at')).issubset({c.value for c in wb['Products'][1]}))
        for sheet in ('Products','Product Images','Option Groups','Options','Variant Combinations'):
            self.assertEqual(wb[sheet].max_row,1)
        self.assertTrue(all(row[1].value is None for row in wb['Brand Info'].iter_rows(min_row=2)))
        wb.close()

class CatalogQualityTests(unittest.TestCase):
    def test_free_gift_filter_preserves_paid_gift_sets(self):
        self.assertTrue(main.is_free_gift({'title':'Miss Rose Com Pk Gift (Free for Premium Orders Rs4999+)'}))
        self.assertTrue(main.is_free_gift({'title':'Promo','handle':'free-gift-orders-above-rs-4999'}))
        self.assertFalse(main.is_free_gift({'title':'Gift Set of 4 Body Mists','handle':'gift-set'}))
        self.assertFalse(main.is_free_gift({'title':'Oil Free Foundation','body_html':'Free gift on orders over Rs4999'}))

    def test_electronics_specs_not_cut_at_300_characters(self):
        specs=[('Processor','Intel Core i7'),('RAM','16 GB DDR5'),('Storage','1 TB SSD'),('Display','15.6 inch 1920 x 1080'),('Battery','56 Wh'),('Connectivity','Wi-Fi 6, Bluetooth 5.2'),('Ports','USB-C, HDMI, 2 x USB-A'),('Operating System','Windows 11'),('Weight','1.7 kg'),('Warranty','One year manufacturer warranty')]
        body='<table>'+''.join(f'<tr><th>{k}</th><td>{v}</td></tr>' for k,v in specs)+'</table>'
        body+='<p>'+'Source technical detail. '*30+'</p>'
        result=main.friendly_description({'title':'Laptop Model X','body_html':body})
        self.assertGreater(len(result),300)
        for key,value in specs:
            self.assertIn(key+': '+value,result)
        self.assertNotIn('32 GB',result)

    def test_specification_heading_enables_detailed_description(self):
        text='Materials: source cotton. Measurements: source sizes. '*15
        result=main.friendly_description({'title':'Shirt','body_html':'<h2>Specifications</h2><p>'+text+'</p>'})
        self.assertIn(text.strip(),result)
        self.assertGreater(len(result),300)

    def test_long_electronics_paragraph_and_definition_list_preserved(self):
        text='Verified source details. '*25
        result=main.friendly_description({'title':'Smartphone X','body_html':'<dl><dt>Battery</dt><dd>5000 mAh</dd></dl><p>'+text+'</p><script>ignore()</script>'})
        self.assertIn('Battery: 5000 mAh',result)
        self.assertIn(text.strip(),result)
        self.assertNotIn('ignore()',result)
        ordinary=main.friendly_description({'title':'Lipstick','body_html':'<p>Source cosmetic details.</p>'})
        self.assertLessEqual(len(ordinary),300)

    def test_categories_capitalized_consistently(self):
        self.assertEqual(main.normalize_category('eyebrow'),'Eyebrow')
        self.assertEqual(main.normalize_category('lip gloss'),'Lip Gloss')
        self.assertEqual(main.normalize_category('bb cream'),'BB Cream')
        self.assertEqual(main.normalize_category('COMPACT POWDER'),'Compact Powder')

    def test_description_images_preserved_separately_from_video(self):
        product={'body_html':'<img src="https://images.example/instructions.jpg"><img src="https://images.example/instructions.jpg"><img src="data:image/png;base64,abc"><video src="video.mp4"></video>'}
        images=main.description_images(product,'https://store.example')
        self.assertEqual(images,[{'src':'https://images.example/instructions.jpg'}])

    def test_canonical_beauty_categories(self):
        for title,expected in [('MISS ROSE Eyeliner Feather Pen','Eyeliner'),
                               ('Miss Rose Matte Oil-Free Foundation','Foundation'),
                               ('Miss Rose Shiny Gloss','Lip Gloss'),
                               ('Cotton embroidered kurta','Kurtas & Kurtis'),
                               ('Blush embroidered kurta','Kurtas & Kurtis')]:
            p={'title':title,'product_type':'Miss Rose Product Categories'}
            self.assertEqual(main.best_category(p,[]),expected)
        self.assertEqual(main.best_category({'title':'Item','product_type':'lipgloss'},[]),'Lip Gloss')

    def test_description_drops_heading_and_preserves_facts(self):
        p={'title':'Cotton Shirt','body_html':'<h2>What is Cotton Shirt</h2><p>Made from cotton fabric. Includes one shirt.</p><p>Buy now!</p><script>bad()</script>'}
        result=main.friendly_description(p)
        self.assertIn('cotton fabric',result)
        self.assertIn('Includes one shirt',result)
        self.assertNotIn('What is',result)
        self.assertNotIn('Buy now',result)
        self.assertNotIn('bad()',result)
        self.assertEqual(main.friendly_description({'title':'No details','body_html':''}),'No details')

    def test_front_image_matches_product_hero(self):
        from types import SimpleNamespace
        main._chotu_tls.image_audit=[]
        main._chotu_tls.image_throttled=False
        main._chotu_tls.adaptive_last_request=0
        pictures=[{'src':'https://cdn.shopify.com/a.jpg'},{'src':'https://cdn.shopify.com/b.jpg'}]
        html='<media-gallery><div class="product__media"><img src="https://cdn.shopify.com/b_800x.jpg?width=800"></div></media-gallery>'
        html += '<product-info><table><tr><th>RAM</th><td>16 GB</td></tr></table></product-info>'
        response=SimpleNamespace(status_code=200,text=html,raise_for_status=lambda:None)
        product={'handle':'item','title':'Laptop Model X'}
        with patch.object(main.S,'get',return_value=response):
            cover=main.storefront_cover('https://store.example',product,pictures,lambda x:None)
        self.assertIn('RAM: 16 GB',main.friendly_description(product))
        self.assertEqual(cover,pictures[1]['src'])
        self.assertEqual(main._chotu_tls.image_audit[0][1],'VERIFIED')
        self.assertEqual(main.image_key('https://missrose.com.pk/cdn/shop/files/front.png?width=800'),
                         main.image_key('https://cdn.shopify.com/s/files/1/0072/9101/1175/files/front.png?v=1'))

    def test_unrelated_image_does_not_replace_gallery(self):
        from types import SimpleNamespace
        main._chotu_tls.image_audit=[];main._chotu_tls.image_throttled=False;main._chotu_tls.adaptive_last_request=0
        response=SimpleNamespace(status_code=200,text='<meta property="og:image" content="https://cdn.shopify.com/other.jpg">',raise_for_status=lambda:None)
        with patch.object(main.S,'get',return_value=response):
            cover=main.storefront_cover('https://store.example',{'handle':'item'},[{'src':'https://cdn.shopify.com/a.jpg'}],lambda x:None)
        self.assertIsNone(cover)
        self.assertEqual(main._chotu_tls.image_audit[0][1],'REVIEW')

    def test_top20_preserves_categories_and_handles_url_scheme(self):
        rows=[{'name':f'Product {i}', 'category':('Foundation','Lipstick','Eyebrow')[i%3],
               'source_url':f'http://store.example/products/p{i}/?variant=123'} for i in range(25)]
        rows.sort(key=lambda r:r['category'])
        original=[r['name'] for r in rows]
        handles=[f'p{i}' for i in range(19,-1,-1)]
        main._chotu_tls.ranking_evidence={h:{'url':'https://store.example/collections/all?sort_by=best-selling','checked_at':'2026-10-17T00:00:00+00:00'} for h in handles}
        ranked=main.prioritize_top_sellers(rows,handles)
        self.assertEqual(ranked,[])
        self.assertEqual([r['name'] for r in rows],original)
        self.assertTrue(all(r['website_order_status']=='unverified' for r in rows))

    def test_top20_paginates_past_unavailable_products(self):
        from types import SimpleNamespace
        def response(handles):
            html='<select><option value="best-selling" selected>Best selling</option></select><ul id="product-grid">'
            html+=''.join(f'<a href="/products/{h}">{h}</a>' for h in handles)+'</ul>'
            return SimpleNamespace(text=html,raise_for_status=lambda:None)
        responses=[response(['soldout']+[f'p{i}' for i in range(10)]),response([f'p{i}' for i in range(10,20)])]
        with patch.object(main.S,'get',side_effect=responses),patch.object(main,'_chotu_backoff'):
            result=main.verified_top20('https://store.example',lambda x:None,{f'p{i}' for i in range(20)})
        self.assertEqual(result,[f'p{i}' for i in range(20)])

    def test_top20_uses_grid_and_excludes_navigation(self):
        from types import SimpleNamespace
        html='<a href="/products/advertisement">Ad</a><select><option value="best-selling" selected>Best selling</option></select><ul id="product-grid">'
        html+=''.join(f'<li><a href="/products/p{i}">Product</a></li>' for i in range(20))+'</ul>'
        response=SimpleNamespace(text=html,raise_for_status=lambda:None)
        with patch.object(main.S,'get',return_value=response):
            ranked=main.verified_top20('https://store.example',lambda x:None,set(f'p{i}' for i in range(20)))
        self.assertEqual(ranked,[f'p{i}' for i in range(20)])
        response.text=html.replace(' selected','')
        with patch.object(main.S,'get',return_value=response):
            self.assertEqual(main.verified_top20('https://store.example',lambda x:None),[])

if __name__=='__main__':unittest.main()
