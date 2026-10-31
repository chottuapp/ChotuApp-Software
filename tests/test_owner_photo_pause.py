import json,tempfile,unittest,copy
from pathlib import Path
from openpyxl import Workbook,load_workbook
from locked_store_rules import has_source_photo,owner_paused
from runner_identity import paused_sources
from owner_product_exclusions import apply

class OwnerPolicyTests(unittest.TestCase):
 def test_newly_disabled_sources_remain_paused_even_if_registry_enabled(self):
  for name,url in [('Gift4U','https://gift4u.pk'),('Xcessories Hub','https://xcessorieshub.com')]:
   self.assertTrue(owner_paused(name,url));self.assertEqual(paused_sources([dict(name=name,url=url)]),{name})
  self.assertFalse(owner_paused('House of Tools','https://houseoftools.com.pk'))
  from unittest.mock import patch
  from requested_stores import probe
  with patch('requested_stores.requests.Session') as session:
   self.assertEqual(probe(dict(name='Gift4U',url='https://gift4u.pk'))['status'],'owner_paused');session.assert_not_called()
 def test_real_variant_or_gallery_photos_count_but_placeholders_do_not(self):
  self.assertFalse(has_source_photo({'images':[{'src':'https://shop.test/woocommerce-placeholder.png'}]}))
  self.assertFalse(has_source_photo({'images':[]}))
  self.assertTrue(has_source_photo({'variants':[{'featured_image':{'src':'https://shop.test/item.jpg'}}]}))
  self.assertTrue(has_source_photo({'scraped_image_url (auto)':'https://shop.test/item.jpg'}))
 def test_missing_photo_is_sale_blocked_with_linked_rows_removed_and_other_data_preserved(self):
  with tempfile.TemporaryDirectory() as tmp:
   root=Path(tmp);source=root/'source.xlsx';out=root/'out.xlsx';b=Workbook();p=b.active;p.title='Products'
   p.append(['shop_code','product_id','Category','name','description','source_url','scraped_image_url (cover, auto)','available (yes/no)','base_price (PKR)'])
   p.append(['S1','missing','Tools','Saw','Exact saw','https://shop.test/saw',None,'yes',1234.50]);p.append(['S1','good','Tools','Drill','Exact drill','https://shop.test/drill',None,'yes',5678.25])
   g=b.create_sheet('Product Images');g.append(['product_id','scraped_image_url (auto)']);g.append(['good','https://shop.test/drill.jpg'])
   v=b.create_sheet('Variant Combinations');v.append(['product_id','available (yes/no)']);v.append(['missing','yes']);v.append(['good','yes'])
   b.save(source);b.close();before=source.read_bytes();result,changes=apply(source,out,'House of Tools');self.assertEqual(source.read_bytes(),before)
   self.assertEqual(changes[0]['product_id'],'missing');self.assertIn('source stock unchanged',changes[0]['reason'])
   b=load_workbook(result);self.assertEqual([r[1] for r in b['Products'].iter_rows(min_row=2,values_only=True) if r[1]],['good']);self.assertEqual(b['Products']['I3'].value,5678.25);self.assertEqual([r[0] for r in b['Variant Combinations'].iter_rows(min_row=2,values_only=True) if r[0]],['good']);b.close()
 def test_dispatch_and_direct_fetch_do_not_start_paused_source(self):
  from fresh804_dispatch import selected_sites
  result=selected_sites([dict(name='Gift4U'),dict(name='Xcessories Hub'),dict(name='House of Tools')],{})
  self.assertEqual([s['name'] for s in result],['House of Tools'])
  from woocommerce_source_support import export
  with self.assertRaisesRegex(ValueError,'OWNER PAUSED'):export({},'https://xcessorieshub.com','Xcessories Hub','unused',lambda _:None)

 def test_export_continues_after_missing_photo_without_changing_source_stock(self):
  from types import SimpleNamespace
  from woocommerce_source_support import export
  from test_app import main
  base='https://houseoftools.com.pk'
  missing=dict(id=1,type='simple',name='Missing photo saw',permalink=base+'/product/saw/',is_in_stock=True,is_purchasable=True,is_on_backorder=False,prices=dict(currency_code='PKR',currency_minor_unit=0,price='100',regular_price='100'),images=[],description='Exact saw specifications.',categories=[])
  good=dict(copy.deepcopy(missing),id=2,name='Drill',permalink=base+'/product/drill/',images=[{'src':base+'/drill.jpg'}])
  before=copy.deepcopy(missing);products={1:missing,2:good};logs=[]
  provider=SimpleNamespace(label='fixture',supports=lambda _:True,catalog=lambda:[missing,good],read=lambda path:products[int(path.rsplit('/',1)[1])])
  with tempfile.TemporaryDirectory() as tmp:
   g=dict(vars(main),OUT=Path(tmp),_chotu_wait=lambda:None)
   export(g,base,'House of Tools',Path(__file__).resolve().parents[1]/'importer-files/Chotu_Products_Import_Template_Discount_Tags.xlsx',logs.append,provider=provider)
   b=load_workbook(Path(tmp)/'House of Tools - REFRESH REVIEW.xlsx',read_only=True);it=b['Products'].iter_rows(values_only=True);headers=next(it);rows=[dict(zip(headers,r)) for r in it];b.close()
   self.assertEqual([r['name'] for r in rows],['Drill']);self.assertEqual(rows[0]['base_price (PKR)'],100);self.assertEqual(missing,before);self.assertTrue(any('OWNER SALE BLOCK' in line for line in logs))
   import csv
   with (Path(tmp)/'House of Tools - OWNER PHOTO HOLDS.csv').open(encoding='utf-8-sig') as f:hold=next(csv.DictReader(f))
   self.assertEqual(hold['catalog_available'],'no');self.assertEqual(hold['source_is_in_stock'],'True')
