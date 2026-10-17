from pathlib import Path
import csv
import importlib.util
import sys
import subprocess
import tempfile
import unittest
from openpyxl import Workbook,load_workbook
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
spec=importlib.util.spec_from_file_location('final_order_main',ROOT/'main.py');main=importlib.util.module_from_spec(spec);spec.loader.exec_module(main)
import install_final_order

class FinalOrderTests(unittest.TestCase):
    def row(self,name,sale=None,category='A'):
        return {'product_id':name,'name':name,'category':category,'base_price (PKR)':100,'sale_price (PKR, agar discount ho)':sale,'source_url':'https://shop/products/'+name,'description':'keep '+name}
    def test_verified_rank_stays_inside_its_category(self):
        rows=[self.row('full'),self.row('other-sale',80),self.row('best-sale',90,'Z'),self.row('equal',100)]
        main._chotu_tls.ranking_evidence={'best-sale':{'url':'https://shop/collections/all?sort_by=best-selling','checked_at':'2026-10-17T00:00:00+00:00'}}
        ranked=main.prioritize_top_sellers(rows,['best-sale'])
        self.assertEqual([r['name'] for r in rows],['other-sale','full','equal','best-sale'])
        self.assertEqual([r['name'] for r in ranked],['best-sale'])
    def test_unverified_handles_never_assign_rank(self):
        rows=[self.row('full'),self.row('sale',80)]
        main._chotu_tls.ranking_evidence={}
        self.assertEqual(main.prioritize_top_sellers(rows,['full']),[])
        self.assertEqual([r['name'] for r in rows],['sale','full'])
    def test_final_saved_file_preserves_every_product_field_and_other_sheet(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);file=root/'export.xlsx';receipt=root/'ranking.csv'
            rows=[self.row('full'),self.row('sale',80),self.row('bundle')]
            book=Workbook();sheet=book.active;sheet.title='Products';headers=list(rows[0]);sheet.append(headers)
            for row in rows:sheet.append([row[h] for h in headers])
            variants=book.create_sheet('Variant Combinations');variants.append(['product_id','base_price (PKR)','sale_price (PKR, agar discount ho)','available (yes/no)','variant_image_url'])
            variants.append(['bundle',300,250,'yes','https://images/bundle.jpg'])
            variants.append(['full',100,10,'no','https://images/unavailable.jpg'])
            images=book.create_sheet('Product Images');images.append(['product_id','url']);images.append(['bundle','https://images/bundle.jpg']);book.save(file)
            with receipt.open('w',newline='') as stream:
                writer=csv.DictWriter(stream,fieldnames=['source_url']);writer.writeheader();writer.writerow({'source_url':'https://shop/products/bundle'})
            main.finalize_catalog_order(file,receipt,lambda x:None)
            result=load_workbook(file);it=iter(result['Products'].values);h=next(it);actual=[dict(zip(h,r)) for r in it]
            self.assertEqual([r['name'] for r in actual],['sale','full','bundle'])
            self.assertEqual({r['product_id']:{k:r[k] for k in rows[0]} for r in actual},{r['product_id']:r for r in rows})
            self.assertEqual(list(result['Variant Combinations'].values),list(variants.values))
            self.assertEqual(list(result['Product Images'].values),list(images.values))
    def test_installer_only_changes_ordering_and_preserves_fast_manual_hooks(self):
        source=subprocess.check_output(['git','-C',str(ROOT),'show','3293422:main.py']).decode()
        for old,new in install_final_order.REPLACEMENTS:source=source.replace(new,old,1)
        source+='\n# manual + fast hooks kept\n# adaptive_pace 0.35\n'
        changed,count=install_final_order.transform(source.encode());expected=source
        for old,new in install_final_order.REPLACEMENTS:expected=expected.replace(old,new,1)
        self.assertEqual(changed.decode(),expected);self.assertEqual(count,5)
        self.assertEqual(install_final_order.transform(changed),(changed,0))
