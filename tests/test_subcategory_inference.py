import copy,tempfile,unittest
from pathlib import Path
from openpyxl import Workbook,load_workbook
from subcategory_inference_support import infer,apply,FIELDS
from website_order_support import apply_order
from website_category_support import fact
class SubcategoryInferenceTests(unittest.TestCase):
 def test_descriptor_in_prose_and_in_title_are_attributed_separately(self):
  a=infer('Certified Delulu','Unisex. A soft regular fit t-shirt with a round neck. Cotton blend polyester.');self.assertEqual(a[:2],('T-Shirts','inferred_description'));self.assertIn('t-shirt',a[2])
  self.assertEqual(infer('Cheddar Cheese 200g','A dairy product.')[0:2],('Cheddar Cheese','inferred_product_name'))
 def test_complete_outfit_not_its_dupattas_and_trousers(self):
  self.assertEqual(infer('Aurelia','An unstitched 3-piece ensemble. Embroidered shirt, dupatta and cotton trousers.')[0],'Unstitched Suits')
  self.assertEqual(infer('Rosaline','Unstitched & RTW - 3PC ensemble. Shirt, dupatta and trousers.')[0],'Suits & Ensembles')
 def test_bundles_and_ingredients_do_not_become_first_component(self):
  self.assertEqual(infer('Night Bundle','A skin care routine with face wash, serum and cream.')[0],'Skin Care Bundles')
  self.assertEqual(infer('Breakfast Combo','Butter and cheese slices.')[0],'Dairy Bundles')
  self.assertEqual(infer('Pack of 12 Karahi Gosht','FARDA Karahi Masala is a spice blend. Ingredients include red chilli powder.')[0],'Spice Blends')
 def test_colour_words_and_component_milk_do_not_cross_product_domains(self):
  self.assertEqual(infer('Rosine','A blush pink cotton kurti. Garment Details: RTW - 1 PC. Shirt: Lawn.')[0],'Kurtas & Kurtis')
  self.assertEqual(infer('Night Bundle','A face-care routine containing cleansing milk and moisturising cream.')[0],'Skin Care Bundles')
  self.assertEqual(infer('Hair Growth Serum','A serum for hair growth.')[0],'Hair Serums')
 def test_conflicting_or_insufficient_source_stays_unverified(self):
  self.assertIsNone(infer('Mystery','A wonderful new arrival.'))
  self.assertIsNone(infer('Unknown','Face serum and face wash in equal amounts.'))
  self.assertIsNone(infer('Hydra Lift','Product Clarification: The description refers to cocoa butter while this listing is for another lotion.'))
 def test_no_false_black_tea_inference(self):
  self.assertEqual(infer('Islamabad Tea Super Dana','Tea imported from Kenya, blended for its aroma.')[0],'Tea Blends')
  self.assertNotEqual(infer('Tea','Tea with a delicate aroma.')[0],'Black Tea')
 def test_product_type_refines_broad_labels_without_falsifying_website_evidence(self):
  row=dict(fact(['SKIN','CLEANSERS'],'https://shop.example/collections/face','verified_navigation_membership'),source_url='https://shop.example/products/x',name='Face Serum',description='A serum.',analysis_note='Owner note')
  apply_order([row]);self.assertEqual(row['subcategory'],'Face Serums');self.assertEqual(row['subcategory_source'],'inferred_product_name');self.assertEqual(row['website_category_status'],'verified_navigation_membership')
  existing={'subcategory':'Owner label','subcategory_source':'manual','name':'A serum'};apply(existing);self.assertEqual(existing['subcategory'],'Owner label')
 def test_single_level_source_label_and_inferred_fallback_repeat_safely(self):
  row=dict(fact(['WOMEN'],'https://shop.example/collections/women','verified_navigation_membership'),source_url='https://shop.example/products/a',name='Dress',description='A dress.')
  apply_order([row]);self.assertEqual(row['subcategory'],'Dresses');first=copy.deepcopy(row);apply_order([row]);self.assertEqual(row,first)
  row={'product_id':'p','name':'Abstract','description':'Soft oversized t-shirt with cotton material.','source_url':'https://shop.example/products/p','base_price (PKR)':100,'analysis_note':'Owner'}
  apply_order([row]);first=copy.deepcopy(row);apply_order([row]);self.assertEqual(row,first);self.assertEqual(row['category'],'');self.assertEqual(row['website_category_status'],'unverified');self.assertEqual(row['subcategory_source'],'inferred_description');self.assertEqual(row['website_category_path'],'')
 def test_purpose_and_application_instructions_are_not_confused(self):
  cases=[('Baby Clothes Cleansing Wipes','Removes stains from clothing.','Laundry Stain-Removal Wipes'),('White Cleansing Shampoo','A detergent for white garments.','Fabric Detergents'),('Baby Shampoo','Gentle cleansing for baby hair.','Baby Shampoos'),('Cat Baby Powder Mist','A deodorising mist for cats.','Pet Deodorisers & Mists'),('Oud Hand Wash','An oud-scented hand wash.','Hand Washes'),('Rice Toner','A face toner.\nHow to Use\nApply using cotton pads.','Face Toners'),('Beard Colour','Colour for your beard.\nHow to Use\nRemove cap and mix.','Beard Colours'),('Turmeric Luffa Soap','A soap containing turmeric powder.','Bath Soaps'),('Cotton T-Shirt','Cotton clothing.\nCare Instructions\nWash with fabric detergent.','T-Shirts'),('Anti Pigmentation Kit','Targets pigmentation and dark spots.','Skin Care Bundles')]
  for title,body,label in cases:
   with self.subTest(title=title):self.assertEqual(infer(title,body)[0],label)
 def test_fragrance_domain_is_not_a_food_or_ingredient_claim(self):
  self.assertEqual(infer('Strawberry Milkshake','A gourmand scent.','Fragrances > Gourmand')[0],'Perfumes & Fragrances')
  self.assertIsNone(infer('Mystery','A new arrival.\nIngredients\nFragrance, water.'))
 def test_export_finalization_keeps_prices_formulas_links_and_provenance(self):
  import main
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'test.xlsx';w=Workbook();s=w.active;s.title='Products';s.append(['shop_code','product_id','name','description','subcategory','source_url','base_price (PKR)','photo_image_url','formula']);s.append(['SE1','p','Abstract','A cotton t-shirt.','','https://shop.example/products/p',100,'https://img.example/p','=1+1']);s['H2'].hyperlink='https://img.example/p';w.create_sheet('Options').append(['keep']);w.save(p)
   main.finalize_catalog_order(p,Path(d)/'unused.csv',lambda _:None);v=load_workbook(p);h=[c.value for c in v['Products'][1]];r=dict(zip(h,next(v['Products'].iter_rows(min_row=2,values_only=True))));self.assertEqual(r['subcategory'],'T-Shirts');self.assertEqual(r['subcategory_source'],'inferred_description');self.assertEqual(r['base_price (PKR)'],100);self.assertEqual(r['formula'],'=1+1');self.assertEqual(v['Products']['H2'].hyperlink.target,'https://img.example/p');self.assertEqual(v['Options']['A1'].value,'keep');self.assertTrue(set(FIELDS)<=set(h));v.close()
if __name__=='__main__':unittest.main()
