import base64,io,json,os,socket,subprocess,sys,tempfile,threading,time,unittest,urllib.error,urllib.request,zipfile
from pathlib import Path
from unittest.mock import patch
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'cloud'));sys.path.insert(0,str(ROOT/'tests'))
from runtime import CloudService,prepare_engine,validate_site
from server import make_server
from cli import parse_stores,checkpoint,restore,write_summary,print_logs
from test_parallel_stores import ENGINE
from openpyxl import Workbook,load_workbook


class CloudTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name);self.engine=self.root/'main.py';self.engine.write_text(ENGINE)
        self.service=CloudService(self.root/'data',engine=self.engine,resolve_sites=False)
    def tearDown(self):self.service.close();self.temp.cleanup()
    def wait_for(self,fn):
        deadline=time.monotonic()+10
        while time.monotonic()<deadline:
            if fn():return
            time.sleep(.025)
        self.fail('Timed out')
    def test_cloud_runs_two_stores_and_download_state_is_per_store(self):
        self.service.start([{'name':'First','url':'https://first.example','steps':12},{'name':'Second','url':'https://second.example'}])
        self.wait_for(lambda:not self.service.running);state=self.service.snapshot()
        self.assertEqual(state['status'],'complete');self.assertEqual(len(state['files']),2)
        self.assertTrue(all(s['software_revision'] for s in state['sites']))
    def test_stop_restart_continue_uses_same_checkpoint_and_updated_source(self):
        self.service.start([{'name':'First','url':'https://first.example'}])
        state=self.service.store.load();folder=self.service.store.root/state['id']/'0'
        self.wait_for(lambda:(folder/'count.txt').exists());self.service.control('pause');time.sleep(.15)
        old=(folder/'count.txt').read_text();time.sleep(.15);self.assertEqual((folder/'count.txt').read_text(),old)
        self.service.close();self.engine.write_text(ENGINE.replace("meta['end']=time.time()","meta['updated']=True;meta['end']=time.time()"))
        template=Path(state['template']);book=load_workbook(template);sheet=book['Products']
        for cell in list(sheet[1]):
            if cell.value=='stock_quantity (kitni quantity bechni hai)':sheet.delete_cols(cell.column)
        sheet.cell(1,sheet.max_column+1,'custom_preserved');book.save(template);book.close()
        self.service=CloudService(self.root/'data',engine=self.engine,resolve_sites=False)
        self.assertFalse(self.service.pause.is_set());self.service.start(resume=True);self.wait_for(lambda:not self.service.running)
        book=load_workbook(template);headers=[cell.value for cell in book['Products'][1]];book.close()
        self.assertIn('custom_preserved',headers);self.assertIn('stock_quantity (kitni quantity bechni hai)',headers)
        self.assertEqual(self.service.store.load()['id'],state['id']);self.assertEqual(self.service.store.load()['status'],'complete')
        self.assertTrue(json.loads((folder/'finished.json').read_text())['updated'])
    def test_saved_checkpoint_roundtrip_keeps_completed_files_and_snapshot(self):
        self.service.start([{'name':'First','url':'https://first.example'}]);self.wait_for(lambda:not self.service.running);self.service.close()
        archive=self.root/'checkpoint.zip';checkpoint(self.root/'data',archive)
        restored=self.root/'restore';restore(archive,restored)
        restored_service=CloudService(restored,engine=self.engine,resolve_sites=False)
        try:
            state=restored_service.store.load();self.assertEqual(restored_service.store.pending(state),[])
            self.assertTrue(Path(state['template']).is_file());self.assertEqual(len(restored_service.snapshot()['files']),1)
        finally:restored_service.close()
    def test_restore_refuses_active_server(self):
        with self.assertRaisesRegex(RuntimeError,'Another ChotuApp'):restore(self.root/'unused.zip',self.service.data)
    def test_checkpoint_archive_rejects_traversal_and_code_files(self):
        for entry in ('../main.py','runtime/main.py','saved_runs/../../escaped'):
            archive=self.root/'bad.zip'
            with zipfile.ZipFile(archive,'w') as z:z.writestr(entry,'bad')
            with self.assertRaises(ValueError):restore(archive,self.root/'reject')
        self.assertFalse((self.root/'escaped').exists())
    def test_distinct_names_and_valid_public_urls_required(self):
        for url in ('file:///etc/passwd','https://localhost','https://user:secret@store.com','https://store.com:8000','https://store.com/product?secret=123'):
            with self.assertRaises(ValueError):validate_site({'name':'Store','url':url},False)
        with patch('runtime.socket.getaddrinfo',return_value=[(None,None,None,None,('127.0.0.1',443))]):
            with self.assertRaises(ValueError):validate_site({'name':'Store','url':'https://store.example'})
        with self.assertRaises(ValueError):self.service.start([{'name':'Same','url':'https://a.example'},{'name':'same','url':'https://b.example'}])
    def test_separate_server_cannot_write_same_volume(self):
        with self.assertRaisesRegex(RuntimeError,'Another ChotuApp'):CloudService(self.service.data,engine=self.engine)
    def test_final_summary_lists_done_and_blocked_stores(self):
        import contextlib
        self.service.start([{'name':'First','url':'https://first.example'}]);self.wait_for(lambda:not self.service.running)
        self.service.store.update(lambda state:state['sites'].append({'name':'Seran','status':'blocked','detail':'SOURCE BLOCKED: HTTP 403'}))
        with contextlib.redirect_stdout(io.StringIO()) as captured:
            summary=write_summary(self.service,'saved')
        self.assertIn('First: DONE',captured.getvalue());self.assertIn('Seran: BLOCKED',captured.getvalue())
        self.assertEqual(summary['stores'][0]['products'],1)
        self.assertEqual(json.loads((self.service.output/'run-summary.json').read_text())['checkpoint'],'saved')
    def test_final_log_drain_includes_last_saved_confirmation(self):
        import contextlib
        self.service.event('FILE SAVED NOW: Example.xlsx')
        with contextlib.redirect_stdout(io.StringIO()) as captured:seen=print_logs(self.service,0)
        self.assertIn('FILE SAVED NOW:',captured.getvalue());self.assertEqual(seen,1)
    def test_store_input_never_runs_shell_text(self):
        parsed=parse_stores('One | https://one.example; Two | https://two.example')
        self.assertEqual([x['name'] for x in parsed],['One','Two'])
        with self.assertRaises(ValueError):parse_stores('invalid')
    def test_real_engine_loads_headlessly_without_changing_desktop_source(self):
        before=(ROOT/'main.py').read_bytes();engine=prepare_engine(self.root/'real')
        script="import runpy; g=runpy.run_path(%r); assert g['_chotu_czone_source']; assert g['_chotu_farda_source']; assert g['_chotu_seran_source']; print('Headless engine ready')"%str(engine)
        result=subprocess.run([sys.executable,'-c',script],capture_output=True,text=True,env=dict(os.environ,CHOTUAPP_DATA_DIR=str(self.root/'home')))
        self.assertEqual(result.returncode,0,result.stderr);self.assertIn('Headless engine ready',result.stdout);self.assertEqual((ROOT/'main.py').read_bytes(),before)


class DashboardTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();root=Path(self.temp.name);engine=root/'main.py';engine.write_text(ENGINE)
        self.service=CloudService(root/'data',engine=engine,resolve_sites=False);self.password='test-password-long-enough'
        self.server=make_server(self.service,self.password,port=0);self.thread=threading.Thread(target=self.server.serve_forever);self.thread.start()
        self.base='http://127.0.0.1:'+str(self.server.server_port)
    def tearDown(self):self.server.shutdown();self.server.server_close();self.thread.join();self.service.close();self.temp.cleanup()
    def request(self,path,body=None,auth=True,csrf=True,origin=None,raw=False):
        headers={}
        if auth:headers['Authorization']='Basic '+base64.b64encode(('admin:'+self.password).encode()).decode()
        if csrf:headers['X-ChotuApp-Request']='1'
        if origin:headers['Origin']=origin
        payload=body if raw else json.dumps(body).encode() if body is not None else None
        request=urllib.request.Request(self.base+path,data=payload,headers=headers)
        try:
            with urllib.request.urlopen(request,timeout=5) as response:return response.status,response.read(),response.headers
        except urllib.error.HTTPError as response:return response.code,response.read(),response.headers
    def test_login_required_on_dashboard_status_and_downloads(self):
        for path in ('/','/api/status','/download?name=test.xlsx'):
            status,_,headers=self.request(path,auth=False);self.assertEqual(status,401);self.assertIn('Basic',headers['WWW-Authenticate'])
        status,body,headers=self.request('/');self.assertEqual(status,200);self.assertIn(b'New Fresh Run',body);self.assertIn('no-store',headers['Cache-Control'])
    def test_cross_origin_control_is_refused(self):
        self.assertEqual(self.request('/api/control',{'command':'pause'},csrf=False)[0],403)
        self.assertEqual(self.request('/api/control',{'command':'pause'},origin='https://evil.example')[0],403)
    def test_fetch_status_and_workbook_download(self):
        self.assertEqual(self.request('/api/start',{'sites':[{'name':'Test','url':'https://test.example'}]})[0],200)
        end=time.monotonic()+10
        while self.service.running and time.monotonic()<end:time.sleep(.03)
        self.assertFalse(self.service.running)
        status,body,_=self.request('/api/status');state=json.loads(body);self.assertEqual(state['status'],'complete')
        status,body,headers=self.request('/download?name=Test.xlsx');self.assertEqual(status,200);self.assertTrue(zipfile.is_zipfile(io.BytesIO(body)))
        self.assertEqual(self.request('/download?name=..%2Ffetch.log')[0],404)
    def test_template_upload_and_validation(self):
        book=Workbook();book.active.title='Products';stream=io.BytesIO();book.save(stream)
        self.assertEqual(self.request('/api/template',stream.getvalue(),raw=True)[0],200)
        self.assertEqual((self.service.data/'template.xlsx').read_bytes(),stream.getvalue())
        self.assertNotEqual(self.request('/api/template',b'invalid',raw=True)[0],200)
    def test_short_password_is_refused(self):
        with self.assertRaises(ValueError):make_server(self.service,'weak')

if __name__=='__main__':unittest.main()
