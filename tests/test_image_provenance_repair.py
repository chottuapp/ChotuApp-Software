import hashlib,tempfile,unittest
from pathlib import Path
from openpyxl import Workbook,load_workbook
from image_provenance_repair import repair
from price_stock_gate import image_metadata_issues
from hold804_repairs import category
from audit804_semantics import review
class ImageRepairTests(unittest.TestCase):
 def fixture(self,path,*,multi=False,wrong=False):
  b=Workbook();p=b.active;p.title='Products';pid='CP'+hashlib.sha256(b'https://www.czone.com.pk|123').hexdigest()[:12]
  product={'product_id':pid,'source_url':'https://www.czone.com.pk/one','scraped_image_url (cover, auto)':'https://cdn.example/one.jpg','option_groups':'G1' if multi else None}
  p.append(list(product));p.append(list(product.values()));v=b.create_sheet('Variant Combinations')
  row={'product_id':pid,'source_variant_id':'123','source_url':product['source_url'],'variant_image_url':'https://cdn.example/other.jpg' if wrong else product['scraped_image_url (cover, auto)'],'variant_image_status':None,'fallback_image_url':None,'variant_image_note':None,'base_price (PKR)':100,'available (yes/no)':'yes'}
  v.append(list(row));v.append(list(row.values()));b.save(path);return product,row
 def test_backfill_changes_only_missing_metadata_for_verified_single_sku(self):
  with tempfile.TemporaryDirectory() as tmp:
   p=Path(tmp)/'source.xlsx';out=Path(tmp)/'out.xlsx';product,original=self.fixture(p);_,n=repair(p,out,'Computer Zone');self.assertEqual(n,1)
   b=load_workbook(out,read_only=True);it=iter(b['Variant Combinations'].values);h=next(it);row=dict(zip(h,next(it)));b.close()
   for k,v in original.items():
    if k not in ('variant_image_status','variant_image_note'):self.assertEqual(row[k],v)
   self.assertEqual(image_metadata_issues(row,product),[])
 def test_multivariant_or_wrong_product_image_is_not_auto_approved(self):
  with tempfile.TemporaryDirectory() as tmp:
   for extra in [{'multi':True},{'wrong':True}]:
    p=Path(tmp)/'x.xlsx';self.fixture(p,**extra);self.assertEqual(repair(p,Path(tmp)/'out.xlsx','Computer Zone'),(p,0))
 def test_serum_capsules_not_dietary_supplements(self):
  row={'name':'Facial Capsules','description':'A lightweight serum for your face and neck.'};label=category(row,'Rivaj UK')[0];self.assertEqual(label,'Face Serums');self.assertEqual(review(row['name'],row['description'],label),[])
