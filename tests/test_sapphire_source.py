import copy,importlib.util,json,sys,tempfile,unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock,patch
from urllib.parse import urlsplit
from bs4 import BeautifulSoup
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT));sys.path.insert(0,str(ROOT/'cloud'))
import sapphire_source_support as source
FIX=ROOT/'tests/fixtures/sapphire_public'
URL=source.BASE+'/collections/ready-to-wear/products/2PCSTCASV996_999.html'
UNS=source.BASE+'/collections/unstitched/products/U3PELXS26V91.html'
def fixture(name):return (FIX/name).read_text()
def response(text,url):return SimpleNamespace(text=text,content=text.encode(),status_code=200,url=url)
class SapphireTests(unittest.TestCase):
 def test_actual_single_regular_fit_preserved_in_public_size_selectors(self):
  url=source.BASE+'/collections/kurtas/products/MSTEKR26V610.html'
  page=source.inspect_page(fixture('regular_fit.html'),url)
  self.assertEqual(page['fixed_options'],{'dwvar_MSTEKR26V610_itemType':'RGL'})
  self.assertEqual([s['label'] for s in page['sizes']],['S','M','L','XL'])
  for size in page['sizes']:
   selected=source.variant_url(url,size)
   self.assertIn('dwvar_MSTEKR26V610_itemType=RGL',selected)
   self.assertIn('dwvar_MSTEKR26V610_size='+size['value'],selected)
   self.assertNotIn('/on/demandware.store/',selected)
 def test_multiple_or_unselected_fit_options_require_verification(self):
  url=source.BASE+'/collections/kurtas/products/MSTEKR26V610.html'
  for multiple in (True,False):
   soup=BeautifulSoup(fixture('regular_fit.html'),'html.parser');node=soup.select_one('[data-attr="itemType"] input')
   if multiple:node.parent.append(copy.copy(node))
   else:del node['checked']
   with self.assertRaisesRegex(ValueError,'Additional product option'):source.inspect_page(str(soup),url)
 def test_size_control_cannot_silently_change_fixed_fit(self):
  url=source.BASE+'/collections/kurtas/products/MSTEKR26V610.html'
  soup=BeautifulSoup(fixture('regular_fit.html'),'html.parser');node=soup.select_one('[data-attr="size"] input');node['value']=node['value'].replace('itemType=RGL','itemType=SLM')
  with self.assertRaisesRegex(ValueError,'preserve fixed'):source.inspect_page(str(soup),url)
 def test_actual_selected_xs_and_l_pages_have_distinct_ids_and_exact_prices(self):
  for filename,value,pid,label in [('rtw_xs.html','XSM','2PCSTCASV996-XSM-999','XS'),('rtw_l.html','LRG','2PCSTCASV996-LRG-999','L')]:
   with self.subTest(filename=filename):
    p=source.inspect_page(fixture(filename),URL);self.assertEqual(p['price'],'12990.00');self.assertEqual(p['id'],pid);self.assertEqual(p['selected_size']['value'],value);self.assertEqual(p['selected_size']['label'],label);self.assertTrue(p['available'])
 def test_unselected_disabled_button_is_not_an_out_of_stock_size(self):
  p=source.inspect_page(fixture('product_ready_to_wear.html'),URL)
  self.assertIsNone(p['selected_size']);self.assertEqual([s['label'] for s in p['sizes']],['XS','S','M','L','XL'])
 def test_actual_unstitched_keeps_complete_contents_images_and_unknown_quantity(self):
  p=source.normalize_product(fixture('product_unstitched.html'),UNS,Mock(side_effect=AssertionError('Unstitched needs no variant request')))
  self.assertEqual(p['variants'][0]['price'],'14990.00');self.assertTrue(p['variants'][0]['available']);self.assertEqual(p['options'],[])
  self.assertIn('1.98m',p['body_html']);self.assertIn('2.5m',p['body_html']);self.assertNotIn('inventory_quantity',p['variants'][0]);self.assertGreater(len(p['images']),2)
 def test_selected_public_pages_verified_not_cartesian_price_copies(self):
  soup=BeautifulSoup(fixture('product_ready_to_wear.html'),'html.parser')
  for node in soup.select('input.options-select'):
   if node.get('data-attr-value') not in ('XSM','LRG'):node.find_parent(class_='size-item').decompose()
  urls=[]
  def read(url):
   urls.append(url);return response(fixture('rtw_xs.html' if '=XSM' in url else 'rtw_l.html'),url)
  product=source.normalize_product(str(soup),URL,read)
  self.assertEqual([v['id'] for v in product['variants']],['2PCSTCASV996-XSM-999','2PCSTCASV996-LRG-999'])
  self.assertTrue(all('/on/demandware.store/' not in url for url in urls));self.assertEqual([v['option1'] for v in product['variants']],['XS','L'])
 def test_master_replayed_for_size_selection_is_rejected(self):
  with self.assertRaisesRegex(ValueError,'did not select'):source.normalize_product(fixture('product_ready_to_wear.html'),URL,lambda url:response(fixture('product_ready_to_wear.html'),url))
 def test_foreign_currency_wrong_product_and_disagreeing_prices_rejected(self):
  original=fixture('rtw_xs.html')
  for html in [original.replace('PKR','USD'),original.replace('2PCSTCASV996-XSM-999.html','OTHER.html'),original.replace('content="12990.00"','content="1.00"')]:
   with self.assertRaises(ValueError):source.inspect_page(html,URL)
 def test_disallowed_api_and_pagination_never_requested(self):
  session=Mock();g={'S':session};reader=source.PublicCatalog(g,lambda _:None)
  for url in [source.BASE+'/on/demandware.store/Sites-Sapphire-Site/default/Product-Variation',source.BASE+'/collections/test?start=24',source.BASE+'/collections/test?sz=999']:
   with self.assertRaises(source.SourceUnavailable):reader.read(url)
  session.get.assert_not_called()
 def test_public_cache_hits_do_not_pay_network_delay(self):
  r=response('cached',URL);cache=Mock();cache.get.return_value=r;session=Mock();backoff=Mock();g={'S':session,'_chotu_tls':SimpleNamespace(saved_response_cache=cache),'_chotu_wait':lambda:None,'_chotu_backoff':backoff,'_chotu_debug_request':lambda *a:None,'_chotu_debug_response':lambda *a,**k:None}
  self.assertIs(source.PublicCatalog(g,lambda _:None).read(URL),r);backoff.assert_not_called();session.get.assert_not_called()
 def test_collection_and_product_discovery_is_explicitly_incomplete(self):
  robot='User-agent: *\nDisallow: /on/demandware.store/\nSitemap: '+source.BASE+'/sitemap_1.xml'
  sitemap='<urlset><url><loc>'+source.BASE+'/collections/unstitched</loc></url></urlset>'
  pages={source.BASE+'/robots.txt':robot,source.BASE+'/sitemap_1.xml':sitemap,source.BASE+'/collections/unstitched':'<div class="product-tile"><a href="'+UNS+'?source=nonplp">Product</a></div>',UNS:fixture('product_unstitched.html')}
  class Reader(source.PublicCatalog):
   def read(self,url,**kwargs):return response(pages[url],url)
  logs=[];products,complete=Reader({'_chotu_wait':lambda:None},logs.append).catalog()
  self.assertEqual(len(products),1);self.assertFalse(complete);self.assertTrue(any('REVIEW REQUIRED' in x for x in logs))
 def test_cloud_runtime_ships_and_enables_the_adapter(self):
  from runtime import prepare_engine
  with tempfile.TemporaryDirectory() as tmp:
   engine=prepare_engine(Path(tmp));self.assertTrue((engine.parent/'sapphire_source_support.py').is_file());self.assertIn('enable_sapphire_source(globals())',engine.read_text())
 def test_real_unstitched_export_preserves_source_and_is_review_only(self):
  from openpyxl import load_workbook
  spec=importlib.util.spec_from_file_location('sapphire_test_engine',ROOT/'main.py');engine=importlib.util.module_from_spec(spec);spec.loader.exec_module(engine)
  product=source.normalize_product(fixture('product_unstitched.html'),UNS,Mock())
  with tempfile.TemporaryDirectory() as tmp:
   engine.OUT=Path(tmp);source.enable_sapphire_source(vars(engine));before=engine.discover;logs=[]
   with patch.object(source.PublicCatalog,'catalog',return_value=([product],False)):
    engine._chotuapp_run_site_original(source.BASE,'Sapphire',ROOT/'importer-files/Chotu_Products_Import_Template_Discount_Tags.xlsx',logs.append)
   self.assertIs(engine.discover,before)
   book=load_workbook(Path(tmp)/'Sapphire - REFRESH REVIEW.xlsx');sheet=book['Products'];headers=[x.value for x in sheet[1]];row=dict(zip(headers,next(sheet.iter_rows(min_row=2,values_only=True))))
   self.assertIn('partial public-page catalog coverage',row['analysis_note']);self.assertEqual(row['source_url'],UNS);self.assertEqual(row['base_price (PKR)'],14990);self.assertIsNone(row['shop_code']);self.assertIn('2.5m',row['description']);self.assertEqual(book['Variant Combinations'].max_row-1,1)
   self.assertTrue(any('Catalog completeness NOT PROVEN' in line for line in logs));book.close()

 def test_selected_disabled_variant_keeps_out_of_stock_without_quantity_guess(self):
  soup=BeautifulSoup(fixture('rtw_l.html'),'html.parser');soup.select_one('button.add-to-cart')['disabled']=''
  p=source.inspect_page(str(soup),URL);self.assertFalse(p['available']);self.assertEqual(p['selected_size']['label'],'L')
 def test_actual_pdp_recommendation_prices_do_not_contaminate_main_price(self):
  url=source.BASE+'/collections/unstitched/products/U1SDDS26V91T.html';html=fixture('price_with_recommendations.html')
  p=source.inspect_page(html,url);self.assertEqual(p['id'],'U1SDDS26V91T');self.assertEqual(p['price'],'1490.00')
  soup=BeautifulSoup(html,'html.parser');main=soup.select_one('.product-wrapper')
  self.assertEqual(len(main.select('.price .sales .cc-price[content]')),3)
  self.assertEqual(len(source.primary_nodes(main,'.price .sales .cc-price[content]')),1)
  # A recommendation still matching analytics must never hide a real primary-price mismatch.
  source.primary_nodes(main,'.price .sales .cc-price[content]')[0]['content']='1491.00'
  with self.assertRaisesRegex(ValueError,'primary product prices disagree'):source.inspect_page(str(soup),url)
 def test_recommendation_sizes_and_images_do_not_become_main_product_fields(self):
  soup=BeautifulSoup(fixture('product_unstitched.html'),'html.parser');main=soup.select_one('.product-wrapper')
  tile=BeautifulSoup('<div class="product-tile"><div data-attr="size"><input class="options-select" data-pid="OTHER" data-attr-value="XXL"></div><img class="pdp-image" data-src="https://pk.sapphireonline.pk/foreign.jpg"></div>','html.parser')
  main.append(tile);p=source.inspect_page(str(soup),UNS)
  self.assertEqual(p['sizes'],[]);self.assertNotIn(source.BASE+'/foreign.jpg',p['images'])
