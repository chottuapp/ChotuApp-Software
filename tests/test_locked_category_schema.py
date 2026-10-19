import copy,tempfile,unittest
from pathlib import Path
from openpyxl import Workbook,load_workbook
from website_order_support import apply_order
from website_category_support import fact
from subcategory_inference_support import export_values
class LockedCategorySchemaTests(unittest.TestCase):
 def test_public_header_uses_leaf_not_parent_and_confirmed_shop_code(self):
  r={'category':'Parent','subcategory':'Leaf','source_url':'https://www.adamsestore.com/products/x','base_price (PKR)':5}
  values=export_values(r);self.assertEqual(values['Category'],'Leaf');self.assertEqual(values['shop_code'],'SE10003');self.assertNotIn('Category',r);self.assertEqual(export_values({'Category':'Owner group','subcategory':''})['Category'],'Owner group')
  self.assertEqual(export_values(dict(r,source_url='https://new.example/products/x'))['shop_code'],'')
  self.assertEqual(export_values(dict(r,shop_code='OWNER'))['shop_code'],'OWNER')
 def test_canonical_website_and_existing_category_survive_repeated_finalization(self):
  r=dict(fact(['SKIN','Face Wash'],'https://shop.example/collections/face','verified_navigation_membership'),source_url='https://shop.example/products/x',Category='Face Wash',subcategory_source='website',name='Unrelated serum',description='A serum.')
  r.pop('category');r.pop('subcategory');apply_order([r]);first=copy.deepcopy(r);apply_order([r]);self.assertEqual(r,first);self.assertEqual(r['Category'],'Face Wash');self.assertEqual(r['website_category_status'],'verified_navigation_membership')
  r={'Category':'Owner group','name':'Mystery','source_url':'https://shop.example/products/y'};apply_order([r]);first=copy.deepcopy(r);apply_order([r]);self.assertEqual(r,first);self.assertEqual(r['Category'],'Owner group')
 def test_existing_category_with_foreign_provenance_is_not_claimed_as_website_verified(self):
  r=dict(fact(['Other'],'https://other.example/collections/x','verified_navigation_membership'),source_url='https://shop.example/products/x',Category='Owner group',subcategory_source='website');r.pop('category');r.pop('subcategory');apply_order([r]);self.assertEqual(r['Category'],'Owner group');self.assertEqual(r['website_category_status'],'unverified');self.assertEqual(r['subcategory_source'],'existing')
 def test_canonical_export_and_finalize_preserve_formula_links_and_other_sheets(self):
  import main
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'out.xlsx';w=Workbook();s=w.active;s.title='Products';s.append(['shop_code','product_id','Category','name','description','source_url','base_price (PKR)','image_url','formula']);main.write_row(s,{'product_id':'p','subcategory':'Face Wash','name':'Cleanser','description':'A facial cleanser.','source_url':'https://vincecare.com/products/x','base_price (PKR)':100,'image_url':'https://img.example/x','formula':'=2+3'});w.create_sheet('Options').append(['keep']);w.save(p)
   main.finalize_catalog_order(p,Path(d)/'unused.csv',lambda _:None);v=load_workbook(p);s=v['Products'];self.assertEqual(s['A2'].value,'SE10008');self.assertEqual(s['C2'].value,'Face Wash');self.assertEqual(s['I2'].value,'=2+3');self.assertEqual(s['H2'].hyperlink.target,'https://img.example/x');self.assertEqual(v['Options']['A1'].value,'keep');self.assertNotIn('subcategory',[c.value for c in s[1]]);self.assertNotIn('category',[c.value for c in s[1]]);v.close()
 def test_default_template_contract(self):
  w=load_workbook(Path(__file__).resolve().parents[1]/'importer-files/Chotu_Products_Import_Template_Discount_Tags.xlsx');h=[c.value for c in w['Products'][1]];self.assertEqual(h[:3],['shop_code','product_id','Category']);self.assertEqual(h.count('Category'),1);self.assertNotIn('category',h);self.assertNotIn('subcategory',h);w.close()
if __name__=='__main__':unittest.main()
