import tempfile,unittest,zipfile,sys
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from test_mobile_khaadi import MobileTests
from legacy_original_schema import normalize,formula
class LegacyTests(unittest.TestCase):
 def source(self,root):
  source=MobileTests().workbook(root,url='https://pk.sapphireonline.pk/products/P1.html');book=load_workbook(source);sheet=book['Products'];sheet['A1']='product_id';sheet['B1']='category';sheet['C1']='subcategory';sheet['A2']='P1';sheet['B2']='Legacy broad';sheet['C2']='Kurta';sheet.auto_filter.ref=sheet.dimensions
  sheet['A2'].number_format='@';sheet.column_dimensions['A'].width=22;sheet.column_dimensions['B'].width=17
  dv=DataValidation(type='list',formula1='"A,B"');sheet.add_data_validation(dv);dv.add('B2:B20');book['Variant Combinations']['C2']='=Products!A2';book.save(source);book.close();return source
 def test_exact_legacy_conversion_preserves_formulas_prices_styles(self):
  with tempfile.TemporaryDirectory() as tmp:
   root=Path(tmp);source=self.source(root);out,changed=normalize(source,root/'new.xlsx');self.assertTrue(changed);book=load_workbook(out)
   self.assertEqual([c.value for c in book['Products'][1]][:3],['shop_code','product_id','Category']);self.assertIsNone(book['Products']['A2'].value);self.assertEqual(book['Products']['B2'].value,'P1');self.assertEqual(book['Products']['C2'].value,'Kurta');self.assertEqual(book['Products']['K2'].value,'=100+50');self.assertEqual(book['Variant Combinations']['C2'].value,'=Products!B2');self.assertEqual(book['Products']['B2'].number_format,'@');self.assertEqual(len(book['Products'].data_validations.dataValidation),0);book.close()
   with zipfile.ZipFile(source) as old,zipfile.ZipFile(out) as new:self.assertEqual(old.read('xl/styles.xml'),new.read('xl/styles.xml'))
 def test_deleted_category_formula_not_guessed(self):
  self.assertEqual(formula('Products!$A$2'),'Products!$B$2')
  self.assertEqual(formula('SUM(C2:D2)',True),'SUM(C2:D2)')
  with self.assertRaises(ValueError):formula('Products!B2')
 def test_locked_schema_passes_through(self):
  with tempfile.TemporaryDirectory() as tmp:
   root=Path(tmp);source=MobileTests().workbook(root);out,changed=normalize(source,root/'new.xlsx');self.assertFalse(changed);self.assertEqual(source,out)
