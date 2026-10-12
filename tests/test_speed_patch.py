"""Compare the exact installed R10 with the speed-only patch on fixed HTTP data."""
import contextlib
import copy
import importlib.util
import json
from pathlib import Path
import subprocess
import tempfile
import threading
import time
import unittest
from http.server import BaseHTTPRequestHandler, HTTPServer
from openpyxl import load_workbook
import sys
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import install_speed_patch

class SpeedPatchTests(unittest.TestCase):
    def test_original_and_speed_patch_export_identical_cells(self):
        source=subprocess.check_output(['git','-C',str(ROOT),'show','28f80c2:main.py'])
        updated,count=install_speed_patch.transform(source)
        self.assertEqual(count,5)
        products=[]
        for i in range(3):
            products.append({'id':i+1,'title':['Lipstick A','Foundation B','Eyebrow C'][i],
              'handle':f'p{i}','vendor':'Fixture','product_type':'Makeup',
              'body_html':'<p>Verified source product details with a smooth texture for daily use.</p>',
              'images':[{'id':100+i,'src':f'https://cdn.shopify.com/s/files/1/1/1/files/p{i}-first.jpg','position':1},
                        {'id':200+i,'src':f'https://cdn.shopify.com/s/files/1/1/1/files/p{i}-hero.jpg','position':2}],
              'options':[{'name':'Shade'}], 'variants':[
                {'id':1000+i*10,'price':'499','compare_at_price':'899','available':True,'option1':'03'},
                {'id':1001+i*10,'price':'599','compare_at_price':'999','available':False,'option1':'01'}]})
        soldout=copy.deepcopy(products[0]);soldout.update(id=99,handle='soldout',title='Soldout Product')
        for v in soldout['variants']:v['available']=False
        gift=copy.deepcopy(products[0]);gift.update(id=100,handle='free-gift-orders-above-rs-4999',title='Free Gift')
        catalog=products+[soldout,gift]
        class Handler(BaseHTTPRequestHandler):
            def do_GET(self):
                if self.path.startswith('/products.json'):data={'products':catalog}
                elif self.path.startswith('/collections.json'):data={'collections':[]}
                elif self.path.startswith('/collections/all'):
                    html='<select><option value="best-selling" selected>Best selling</option></select><ul id="product-grid">'
                    html+=''.join(f'<li><h3><a href="/products/p{i}">Product</a></h3></li>' for i in (2,0,1))+'</ul>'
                    self.send_response(200);self.end_headers();self.wfile.write(html.encode());return
                elif self.path.startswith('/products/'):
                    handle=self.path.split('/products/',1)[1].split('?')[0].removesuffix('.js')
                    p=next(p for p in catalog if p['handle']==handle)
                    if self.path.split('?')[0].endswith('.js'):
                        data={'id':p['id'],'variants':[dict(v,price=int(float(v['price'])*100),compare_at_price=int(float(v['compare_at_price'])*100)) for v in p['variants']]}
                    else:
                        html='<media-gallery><div class="product__media"><img src="'+p['images'][1]['src']+'"></div></media-gallery>'
                        self.send_response(200);self.end_headers();self.wfile.write(html.encode());return
                else:data={}
                self.send_response(200);self.end_headers();self.wfile.write(json.dumps(data).encode())
            def log_message(self,*args):pass
        server=HTTPServer(('127.0.0.1',0),Handler)
        t=threading.Thread(target=server.serve_forever,daemon=True);t.start()
        base=f'http://127.0.0.1:{server.server_port}'
        outputs=[];elapsed=[]
        try:
            with tempfile.TemporaryDirectory() as tmp:
                for label,body in [('original',source),('speed',updated)]:
                    folder=Path(tmp)/label;folder.mkdir();file=folder/'main.py';file.write_bytes(body)
                    spec=importlib.util.spec_from_file_location('speed_fixture_'+label,file)
                    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
                    module.OUT=folder/'exports'
                    module._chotu_open_checkpoint(base,'Fixture')
                    start=time.monotonic()
                    try:module.run_site(base,'Fixture',str(ROOT/'ChotuApp_Template.xlsx'),lambda x:None)
                    finally:module._chotu_close_checkpoint(success=True)
                    elapsed.append(time.monotonic()-start)
                    wb=load_workbook(module.OUT/'Fixture.xlsx',read_only=True,data_only=True)
                    snapshot={n:list(wb[n].values) for n in wb.sheetnames};wb.close();outputs.append(snapshot)
                self.assertEqual(outputs[0],outputs[1])
                self.assertLess(elapsed[1],elapsed[0])
                print('Same fixed data: every workbook cell matches. Original %.2fs; speed patch %.2fs.'%tuple(elapsed))
        finally:server.shutdown();server.server_close();t.join()

    def test_retry_after_and_rate_limit_backoff_preserved(self):
        from unittest.mock import patch
        import requests
        source=subprocess.check_output(['git','-C',str(ROOT),'show','28f80c2:main.py'])
        updated,_=install_speed_patch.transform(source)
        with tempfile.TemporaryDirectory() as tmp:
            file=Path(tmp)/'main.py';file.write_bytes(updated)
            spec=importlib.util.spec_from_file_location('speed_rate_fixture',file)
            module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
            limited=requests.Response();limited.status_code=429;limited.headers['Retry-After']='30';limited._content=b'{}'
            ok=requests.Response();ok.status_code=200;ok._content=b'{"verified":true}'
            waits=[]
            with patch.object(module.S,'get',side_effect=[limited,ok]),patch.object(module,'_chotu_backoff',side_effect=lambda seconds,reason:waits.append((seconds,reason))):
                self.assertTrue(module.get_json('https://store.example/test',tries=2)['verified'])
            self.assertTrue(any(delay>=30 and reason=='Rate limit' for delay,reason in waits))
            self.assertGreater(module._chotu_tls.adaptive_pace,.35)

    def test_unknown_pacing_source_rejected(self):
        with self.assertRaises(RuntimeError):install_speed_patch.transform(b'print("unrelated app")')

if __name__=='__main__':unittest.main()
